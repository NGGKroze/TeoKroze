"""Лаборатория: инспекция на файлове за подобряване на парсването (PDF с координати/OCR, Excel, CSV)."""
import base64
import csv
import io
import json
import os
import re
import time
import zipfile
from pathlib import Path

from . import ocr, pdf

MAX_ROWS, MAX_COLS = 400, 60


# ------------------------- PDF -------------------------
def group_lines(words, y_tol=3.0):
    """words: [[x0,y0,x1,y1,текст]] -> редове [{y, words:[{x,text}], text}] отгоре надолу, отляво надясно."""
    ws = sorted(words, key=lambda w: ((w[1] + w[3]) / 2, w[0]))
    lines = []
    for w in ws:
        yc = (w[1] + w[3]) / 2
        if lines and abs(lines[-1]["y"] - yc) <= y_tol:
            line = lines[-1]
            line["words"].append({"x": w[0], "x1": w[2], "text": w[4]})
            line["y"] = (line["y"] * (len(line["words"]) - 1) + yc) / len(line["words"])
        else:
            lines.append({"y": yc, "words": [{"x": w[0], "x1": w[2], "text": w[4]}]})
    for line in lines:
        line["words"].sort(key=lambda a: a["x"])
        line["text"] = " ".join(a["text"] for a in line["words"])
        line["y"] = round(line["y"], 1)
    return lines


def layout_text(lines, page_width, cols=140):
    """Текст "по колони" като pdftotext -layout: думите се слагат на позиция, пропорционална на x."""
    out = []
    scale = cols / max(page_width, 1)
    for line in lines:
        buf = []
        for w in line["words"]:
            pos = int(w["x"] * scale)
            cur = sum(len(s) for s in buf)
            if pos > cur:
                buf.append(" " * (pos - cur))
            elif buf:
                buf.append(" ")
            buf.append(w["text"])
        out.append("".join(buf).rstrip())
    return "\n".join(out)


def analyze_pdf(data: bytes, ocr_mode="auto", lang="eng", dpi=300):
    res = pdf.extract_pdf(data, ocr_mode, lang, dpi)
    for p in res["pages"]:
        lines = group_lines(p["words"])
        p["lines"] = lines
        p["layout"] = layout_text(lines, p["width"])
    return res


def render_page(data: bytes, page: int, dpi=80) -> bytes:
    try:
        import fitz
    except ImportError:
        import pymupdf as fitz
    doc = fitz.open(stream=data, filetype="pdf")
    page = max(1, min(page, len(doc)))
    return doc[page - 1].get_pixmap(dpi=dpi).tobytes("png")


# ------------------------- Excel / CSV -------------------------
def _clean(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return int(v)
    if hasattr(v, "isoformat"):
        return v.isoformat(sep=" ")
    return v if isinstance(v, (int, float, str, bool)) else str(v)


def analyze_table(data: bytes, filename: str):
    name = filename.lower()
    sheets = []
    if name.endswith((".xlsx", ".xlsm")):
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
        for ws in wb.worksheets:
            rows = []
            for r in ws.iter_rows(min_row=1, max_row=min(ws.max_row, MAX_ROWS), max_col=min(ws.max_column, MAX_COLS), values_only=True):
                rows.append([_clean(c) for c in r])
            sheets.append({"name": ws.title, "rows_total": ws.max_row, "cols_total": ws.max_column,
                           "merged": [str(m) for m in list(ws.merged_cells.ranges)[:200]], "state": ws.sheet_state, "rows": rows})
    elif name.endswith(".xls"):
        try:
            import xlrd
        except ImportError as exc:
            raise RuntimeError("За .xls е нужен пакетът xlrd (или запишете файла като .xlsx).") from exc
        wb = xlrd.open_workbook(file_contents=data)
        for ws in wb.sheets():
            rows = [[_clean(ws.cell_value(r, c)) for c in range(min(ws.ncols, MAX_COLS))] for r in range(min(ws.nrows, MAX_ROWS))]
            sheets.append({"name": ws.name, "rows_total": ws.nrows, "cols_total": ws.ncols, "merged": [], "state": "visible", "rows": rows})
    else:  # csv / txt
        text = data.decode("utf-8-sig", errors="replace")
        try:
            dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
        except csv.Error:
            dialect = csv.excel
        rows = [[_clean(c) for c in row[:MAX_COLS]] for _, row in zip(range(MAX_ROWS), csv.reader(io.StringIO(text), dialect))]
        sheets.append({"name": "CSV", "rows_total": len(rows), "cols_total": max((len(r) for r in rows), default=0), "merged": [], "state": "visible", "rows": rows})
    for s in sheets:
        s["header_candidates"] = header_candidates(s["rows"])
    return {"sheets": sheets}


HEADER_WORDS = re.compile(r"(carton|ctn|box|colis|sku|style|color|colour|coloris|size|qty|quantity|quantit|total|ref|ean|weight|poids|order|po\b)", re.I)


def header_candidates(rows):
    """Редове, които приличат на заглавие на таблица (много ключови думи) - подсказка за парсера."""
    out = []
    for i, r in enumerate(rows):
        cells = [str(c) for c in r if c != ""]
        hits = sum(1 for c in cells if HEADER_WORDS.search(c))
        if hits >= 3:
            out.append({"row": i + 1, "hits": hits, "cells": cells[:16]})
    return sorted(out, key=lambda x: -x["hits"])[:5]


# ------------------------- Пакет за анализ -------------------------
def save_package(out_dir: str, name: str, report: dict, file_b64: str = "", file_name: str = "") -> str:
    """Записва ZIP (оригиналният файл + report.json) в <out_dir>\\Лаборатория и връща пътя."""
    safe = re.sub(r"[^\w.-]+", "_", Path(name or "file").stem)[:60] or "file"
    folder = Path(out_dir) / "Лаборатория"
    folder.mkdir(parents=True, exist_ok=True)
    target = folder / f"{time.strftime('%Y%m%d_%H%M%S')}_{safe}.zip"
    with zipfile.ZipFile(target, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("report.json", json.dumps(report, ensure_ascii=False, indent=1))
        if file_b64:
            z.writestr("original/" + (re.sub(r"[\\/:*?\"<>|]+", "_", file_name) or "file"), base64.b64decode(file_b64))
    return str(target)
