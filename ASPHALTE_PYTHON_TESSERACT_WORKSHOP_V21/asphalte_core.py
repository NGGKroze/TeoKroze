from __future__ import annotations

import io
import json
import math
import os
import re
import shutil
import tempfile
import uuid
import zipfile
from copy import copy
from dataclasses import dataclass, asdict
from datetime import datetime, timedelta
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional, Tuple

import fitz  # PyMuPDF
from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.page import PageMargins
from openpyxl.worksheet.pagebreak import Break
from openpyxl.drawing.image import Image as XLImage
from openpyxl.drawing.spreadsheet_drawing import OneCellAnchor, AnchorMarker
from openpyxl.drawing.xdr import XDRPositiveSize2D
from openpyxl.utils.units import pixels_to_EMU
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.units import mm
from reportlab.pdfgen import canvas
from reportlab.graphics.barcode import eanbc
from reportlab.graphics.shapes import Drawing
from reportlab.graphics import renderPDF

try:
    from barcode import EAN13 as BarcodeEAN13, Code128 as BarcodeCode128
    from barcode.writer import ImageWriter as BarcodeImageWriter
except Exception:
    BarcodeEAN13 = None
    BarcodeCode128 = None
    BarcodeImageWriter = None

try:
    from PIL import Image
    import pytesseract
except Exception:  # Tesseract is optional until OCR is requested
    Image = None
    pytesseract = None

DEFAULT_SHIPPER = "PELINTEX BULGARIA LTD\n5 HRISTO BOTEV BLVD.\n7000 RUSE\nBULGARIA"
DEFAULT_RECEIVER = "LOGTEX\nZI du Cormier 5\n150 rue Pierre-Gilles de Gennes\n49300 CHOLET, FRANCE"
SIZE_ORDER = ["XXS", "XS", "S", "M", "L", "XL", "XXL", "3XL", "4XL", "5XL"]
SIZE_RE = re.compile(r"(?<![A-Z0-9])(3XL|4XL|5XL|XXXL|XXL|XXS|XS|XL|S|M|L)(?![A-Z0-9])", re.I)
EAN_RE = re.compile(r"\b(\d{8}|\d{12,14})\b")

COLOR_ALIASES = {
    "beig": "Beige",
    "beige": "Beige",
    "corteccia": "Beige",
    "camel": "Camel",
    "came": "Camel",
    "bleu marine": "Bleu marine",
    "blue marine": "Bleu marine",
    "marine": "Bleu marine",
    "navy": "Bleu marine",
    "black navy": "Bleu marine",
    "bluemarine": "Bleu marine",
    "bleumarine": "Bleu marine",
    "blon": "Bleu marine",
    "noir": "Noir",
    "black": "Noir",
    "white": "Blanc",
    "blanc": "Blanc",
    "ecru": "Ecru",
    "écru": "Ecru",
    "green": "Vert",
    "vert": "Vert",
    "vemi": "Vert militaire",
    "vert militaire": "Vert militaire",
}

HEADER_SKIP_WORDS = {
    "order", "placement", "raw", "fabric", "asphalte", "color", "supplier", "total", "export",
    "date", "factory", "quality", "control", "plan", "breakdown", "final", "size", "qc", "bureau",
    "global", "quantities", "price", "amount", "launch", "end", "season", "name", "main", "certification",
    "reference", "delivery", "address", "optional", "batch", "buffer", "confirmed", "received", "sent", "trims",
}


def clean_cell(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%d/%m/%Y")
    return str(value).strip()


def compact_spaces(value: str) -> str:
    return re.sub(r"\s+", " ", str(value or "")).strip()


def normalize_key(value: str) -> str:
    value = str(value or "").lower().strip()
    value = value.replace("®", "").replace("é", "e").replace("è", "e").replace("ê", "e")
    value = re.sub(r"[^a-z0-9]+", "", value)
    return value


def normalize_color(value: str) -> str:
    raw = compact_spaces(str(value or ""))
    if not raw:
        return ""
    low = raw.lower().replace("®", "").replace("é", "e").replace("è", "e")
    low_clean = re.sub(r"[^a-z0-9]+", " ", low).strip()
    key_clean = normalize_key(low_clean)
    # Preserve compound visible color names such as "Bleu marine rayé" instead of collapsing
    # them to "Bleu marine". Exact aliases and SKU shorthand codes still canonicalize.
    for alias, canonical in COLOR_ALIASES.items():
        a = alias.lower().replace("é", "e")
        ak = normalize_key(a)
        if key_clean == ak or low_clean == a:
            return canonical
        if ak in {"came", "vemi", "blon"} and ak in key_clean:
            return canonical
    return raw[:1].upper() + raw[1:]


def parse_date(value: str) -> Optional[datetime]:
    text = compact_spaces(value)
    if not text:
        return None
    # Normalize ordinal suffixes and commas for older 2021 order forms,
    # e.g. "June 4, 2021" / "4th June 2021".
    norm = re.sub(r"\b(\d{1,2})(st|nd|rd|th)\b", r"\1", text, flags=re.I).replace(",", "")
    for fmt in (
        "%d/%m/%Y", "%d.%m.%Y", "%d-%m-%Y", "%Y-%m-%d", "%d/%m/%y", "%d.%m.%y",
        "%B %d %Y", "%b %d %Y", "%d %B %Y", "%d %b %Y",
        "%B %d %y", "%b %d %y", "%d %B %y", "%d %b %y",
    ):
        try:
            return datetime.strptime(norm, fmt)
        except Exception:
            pass
    m = re.search(r"(\d{1,2})[\./-](\d{1,2})[\./-](\d{2,4})", text)
    if m:
        d, mo, y = m.groups()
        y = int(y)
        if y < 100:
            y += 2000
        try:
            return datetime(y, int(mo), int(d))
        except Exception:
            return None
    return None


def fmt_date(value: Optional[datetime]) -> str:
    if not value:
        return ""
    return value.strftime("%d/%m/%Y")


def identification_from_shipment_date(date_value: Optional[datetime]) -> str:
    if not date_value:
        return ""
    return "189" + date_value.strftime("%d%m%Y")


def find_tesseract() -> Optional[str]:
    if pytesseract is None:
        return None
    common = [
        os.environ.get("TESSERACT_CMD", ""),
        r"C:\Program Files\Tesseract-OCR\tesseract.exe",
        r"C:\Program Files (x86)\Tesseract-OCR\tesseract.exe",
        r"C:\Users\%USERNAME%\AppData\Local\Programs\Tesseract-OCR\tesseract.exe",
    ]
    for p in common:
        if not p:
            continue
        p = os.path.expandvars(p)
        if os.path.exists(p):
            pytesseract.pytesseract.tesseract_cmd = p
            return p
    found = shutil.which("tesseract")
    if found:
        pytesseract.pytesseract.tesseract_cmd = found
    return found


def extract_pdf_text(pdf_path: str, ocr_mode: str = "auto", dpi: int = 220) -> Tuple[str, List[str]]:
    """Extract text with PyMuPDF first, then OCR pages if needed/forced."""
    warnings: List[str] = []
    doc = fitz.open(pdf_path)
    text_parts: List[str] = []
    per_page_text: List[str] = []
    for page in doc:
        page_text = page.get_text("text") or ""
        per_page_text.append(page_text)
        text_parts.append(page_text)
    base_text = "\n".join(text_parts).strip()

    should_ocr = ocr_mode == "always" or (ocr_mode == "auto" and len(base_text) < 250)
    if not should_ocr:
        return base_text, warnings

    if pytesseract is None or Image is None:
        warnings.append("OCR requested, but pytesseract/Pillow is not installed.")
        return base_text, warnings
    if not find_tesseract():
        warnings.append("Tesseract OCR executable was not found. Text layer extraction was used only.")
        return base_text, warnings

    ocr_parts: List[str] = []
    for i, page in enumerate(doc):
        if ocr_mode == "auto" and len(per_page_text[i].strip()) > 200:
            ocr_parts.append(per_page_text[i])
            continue
        try:
            pix = page.get_pixmap(matrix=fitz.Matrix(dpi / 72, dpi / 72), alpha=False)
            img = Image.open(io.BytesIO(pix.tobytes("png")))
            ocr_text = pytesseract.image_to_string(img, lang="eng+fra")
            ocr_parts.append(ocr_text)
        except Exception as exc:
            warnings.append(f"OCR failed on page {i + 1}: {exc}")
            ocr_parts.append(per_page_text[i])
    final_text = "\n".join(ocr_parts).strip() or base_text
    return final_text, warnings


# ---------------------------------------------------------------------------
# Layout-aware order reading
#
# Some Asphalte PDFs (e.g. OF-2188) draw the supplier colour text so that it
# overlaps the first size columns. A plain text extraction sorts characters by
# position and mixes them ("N1atur Ecr3u 000" = "Natur Ecru 000" + the XS/S
# quantities 1 and 3), which shifts every quantity. In the PDF content stream
# the cells are separate runs, so we rebuild cells from the stream order
# (a new cell starts on a backward jump, a big gap or a new line) and emit one
# cell per line - exactly what the table parser expects.
# ---------------------------------------------------------------------------
def _chars_to_cells(chars: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
    cells: List[Dict[str, Any]] = []
    cur: List[Dict[str, Any]] = []

    def flush():
        if cur:
            txt = re.sub(r"\s+", " ", "".join(ch["c"] for ch in cur)).strip()
            if txt:
                cells.append({
                    "text": txt,
                    "x0": min(ch["x0"] for ch in cur),
                    "x1": max(ch["x1"] for ch in cur),
                    "y": sum(ch["y"] for ch in cur) / len(cur),
                    "size": max(ch["size"] for ch in cur),
                })
        cur.clear()

    prev = None
    for ch in chars:
        if prev is not None:
            size = max(ch["size"], prev["size"], 1.0)
            gap = ch["x0"] - prev["x1"]
            if abs(ch["y"] - prev["y"]) > 0.45 * size or gap < -0.6 * size or gap > 1.0 * size:
                flush()
            elif gap > 0.18 * size and not prev["c"].isspace() and not ch["c"].isspace():
                cur.append({"c": " ", "x0": prev["x1"], "x1": ch["x0"], "y": prev["y"], "size": prev["size"]})
        cur.append(ch)
        prev = ch
    flush()
    return cells


def _cells_to_text(cells: List[Dict[str, Any]]) -> str:
    rows: List[Dict[str, Any]] = []
    for cell in sorted(cells, key=lambda c: (c["y"], c["x0"])):
        if rows and abs(cell["y"] - rows[-1]["y"]) <= 0.5 * max(cell["size"], 3.0):
            rows[-1]["cells"].append(cell)
        else:
            rows.append({"y": cell["y"], "cells": [cell]})
    lines: List[str] = []
    for row in rows:
        for cell in sorted(row["cells"], key=lambda c: c["x0"]):
            lines.append(cell["text"])
    return "\n".join(lines)


def extract_pdf_layout_text(pdf_path: str, per_run: bool = True) -> str:
    """Text layer rebuilt from content-stream cells, one table cell per line.

    per_run=True treats every text-show operation of the PDF as its own cell (an Excel/web export
    draws each table cell separately), so a supplier colour that ends right next to the first
    quantity ("...NAVY" + "3") can no longer be glued into one cell. per_run=False only uses the
    gap/jump heuristics of _chars_to_cells.
    """
    doc = fitz.open(pdf_path)
    parts: List[str] = []
    for page in doc:
        cells: List[Dict[str, Any]] = []
        if per_run:
            try:
                for span in page.get_texttrace():
                    size = float(span.get("size") or 8.0)
                    chars = []
                    for c in span.get("chars", []):
                        x0, y0, x1, y1 = c[3]
                        chars.append({"c": chr(c[0]), "x0": x0, "x1": x1, "y": (y0 + y1) / 2.0, "size": size})
                    cells.extend(_chars_to_cells(chars))
            except Exception:
                cells = []
        if not cells:
            chars = []
            data = page.get_text("rawdict") or {}
            for block in data.get("blocks", []):
                if block.get("type", 0) != 0:
                    continue
                for line in block.get("lines", []):
                    for span in line.get("spans", []):
                        size = float(span.get("size") or 8.0)
                        for ch in span.get("chars", []):
                            x0, y0, x1, y1 = ch["bbox"]
                            chars.append({"c": ch.get("c", ""), "x0": x0, "x1": x1, "y": (y0 + y1) / 2.0, "size": size})
            cells = _chars_to_cells(chars)
        parts.append(_cells_to_text(cells))
    return "\n".join(parts)


def _declared_total(text: str) -> int:
    """TOTAL printed under GLOBAL ORDER (used as a cross-check for the detected quantities)."""
    lines = [compact_spaces(l) for l in text.splitlines() if compact_spaces(l)]
    start = next((i for i, l in enumerate(lines) if l.upper() == "GLOBAL ORDER"), -1)
    if start < 0:
        return 0
    for i in range(start + 1, len(lines)):
        m = re.match(r"^TOTAL\s+(\d+)\b", lines[i], re.I)
        if m:
            return int(m.group(1))
        if lines[i].upper() == "TOTAL" and i + 1 < len(lines) and is_numeric_token(lines[i + 1]):
            return int(lines[i + 1])
    return 0


def _order_quality(parsed: Dict[str, Any], declared: int) -> Tuple[int, int, int]:
    lines = parsed.get("order_lines", [])
    good = bad = 0
    for l in lines:
        if l.get("size_breakdown_missing"):
            continue
        if sum(int(v or 0) for v in (l.get("sizes") or {}).values()) == int(l.get("total", 0) or 0):
            good += 1
        else:
            bad += 1
    total = int(parsed.get("total_pieces", 0) or 0)
    return (1 if declared and total == declared else 0, good - 2 * bad, total)


def _is_line_warning(w: str) -> bool:
    return bool(re.search(r"order line|Buffer/Batch|Packing/label|color totals", w, re.I))


def parse_order_pdf(pdf_path: str, ocr_mode: str = "auto") -> Tuple[Dict[str, Any], List[str]]:
    """Read an order PDF; use the layout-aware reading when it recovers the table better."""
    text, warnings = extract_pdf_text(pdf_path, ocr_mode=ocr_mode)
    parsed = parse_asphalte_order_text(text)
    declared = _declared_total(text)
    for per_run in (True, False):
        try:
            layout_text = extract_pdf_layout_text(pdf_path, per_run=per_run)
        except Exception:
            layout_text = ""
        if not layout_text.strip():
            continue
        parsed_l = parse_asphalte_order_text(layout_text)
        declared = declared or _declared_total(layout_text)
        if _order_quality(parsed_l, declared) > _order_quality(parsed, declared):
            merged = dict(parsed)
            meta = dict(parsed.get("metadata", {}))
            for k, v in parsed_l.get("metadata", {}).items():
                if not meta.get(k):
                    meta[k] = v
            meta["export_date"] = parsed_l.get("metadata", {}).get("export_date", "") or meta.get("export_date", "")
            layout_name = parsed_l.get("metadata", {}).get("product_name", "")
            if layout_name and not looks_like_product_name(meta.get("product_name", "")) and looks_like_product_name(layout_name):
                meta["product_name"] = layout_name
            merged["metadata"] = meta
            for k in ("order_lines", "sizes", "total_pieces"):
                merged[k] = parsed_l.get(k)
            merged["warnings"] = [w for w in parsed.get("warnings", []) if not _is_line_warning(w)] + [w for w in parsed_l.get("warnings", []) if _is_line_warning(w)]
            merged["warnings"].append("Order table was read with layout-aware column detection (the PDF text layer mixes supplier colour text with the first quantities).")
            parsed = merged
    total = int(parsed.get("total_pieces", 0) or 0)
    if declared and total != declared:
        parsed.setdefault("warnings", []).append(f"Detected {total} pcs, but the PDF GLOBAL ORDER total says {declared}. Check the order lines before exporting.")
    return parsed, warnings


def extract_between_lines(lines: List[str], start_label: str, end_label: str) -> List[str]:
    low = [x.lower() for x in lines]
    start_idx = next((i for i, x in enumerate(low) if start_label.lower() in x), -1)
    if start_idx == -1:
        return []
    end_idx = next((i for i in range(start_idx + 1, len(lines)) if end_label.lower() in low[i]), len(lines))
    return [l for l in lines[start_idx + 1:end_idx] if l.strip()]


def parse_label_value(lines: List[str], label: str) -> str:
    """Read Label: value pairs without accidentally matching words like 'vendredi' for 'End'."""
    label_clean = re.escape(label.strip())
    pattern = re.compile(rf"^{label_clean}\b\s*:??\s*(.*)$", re.I)
    bad_next = {"color", "colour", "asphalte color", "supplier color", "xs", "s", "m", "l", "xl", "xxl", "3xl", "total"}
    for i, line in enumerate(lines):
        line = compact_spaces(line)
        m = pattern.match(line)
        if not m:
            continue
        after = compact_spaces(m.group(1)).lstrip(": ").strip()
        if after and not after.lower().startswith(("mail", "vta", "contact", "delivery address", "supplier")) and after.lower() not in bad_next:
            return after
        for j in range(i + 1, min(i + 5, len(lines))):
            nxt = compact_spaces(lines[j])
            if not nxt:
                continue
            nxt_low = nxt.lower()
            if nxt_low in bad_next:
                continue
            if nxt and not re.search(r"^(mail|vta|contact|delivery address|supplier address|supplier|order details)\b", nxt, re.I):
                return nxt
    return ""


def is_numeric_token(v: str) -> bool:
    return bool(re.fullmatch(r"\d+", str(v or "").strip()))


def is_headerish(v: str) -> bool:
    low = str(v or "").lower().strip()
    if not low:
        return True
    words = re.findall(r"[a-z]+", low)
    return any(w in HEADER_SKIP_WORDS for w in words)


PRODUCT_WORD_RE = re.compile(
    r"\b(veste|blouson|parka|manteau|caban|trench|imper|jacket|gilet|pantalon|pant|slack|varsity|coat|balmacaan|woolpure|ventile|moleskine|gabardine)\b",
    re.I,
)
SEASON_TOKEN_RE = re.compile(r"^(PE|AH|SS|FW|W)\s*\d{2,4}$", re.I)


def looks_like_product_name(value: str) -> bool:
    v = compact_spaces(value)
    if not v:
        return False
    low = v.lower()
    if is_numeric_token(v) or parse_date(v) or normalize_size(v) in SIZE_ORDER:
        return False
    return bool(PRODUCT_WORD_RE.search(low) or re.search(r"\bv\d+\b", low))


def is_batch_noise_line(value: str) -> bool:
    low = compact_spaces(value).lower()
    if not low:
        return True
    if low in {"reference", "color", "colour", "asphalte color", "supplier color", "season", "total", "quantities", "price (€)", "amount (€)"}:
        return True
    if SEASON_TOKEN_RE.match(compact_spaces(value)):
        return True
    if re.search(r"^(delivery date|export date|ex factory|asphalte warehouse|quality control|main fabric|final size|color breakdown|order placement)\b", low):
        return True
    if re.search(r"^(repartition|répartition|sent|received|to be confirmed|tbc|option)\b", low):
        return True
    return False


def is_color_candidate(value: str) -> bool:
    v = compact_spaces(value)
    if not v:
        return False
    if is_numeric_token(v) or parse_date(v) or normalize_size(v) in SIZE_ORDER:
        return False
    if is_batch_noise_line(v):
        return False
    # Product/reference tokens can appear in the first column of old 2021 PDFs; they are not colors.
    if looks_like_product_name(v):
        return False
    if len(v) > 80:
        return False
    return True


def find_export_date_after(tokens: List[str], start: int) -> str:
    for look in tokens[start:start + 8]:
        dt = parse_date(look)
        if dt:
            return fmt_date(dt)
    return ""


def parse_qty_block(tokens: List[str], start: int, sizes: List[str]) -> Optional[Tuple[Dict[str, int], int, int, str, bool]]:
    """Return sizes, total, next-index, export_date, missing_breakdown."""
    n = len(sizes)
    if start >= len(tokens):
        return None
    # Normal table: one number per size followed by total.
    qty_tokens = tokens[start:start + n]
    total_token = tokens[start + n] if start + n < len(tokens) else ""
    if len(qty_tokens) == n and all(is_numeric_token(q) for q in qty_tokens):
        quantities = {sizes[idx]: int(qty_tokens[idx]) for idx in range(n)}
        sum_qty = sum(quantities.values())
        if is_numeric_token(total_token):
            export_date = find_export_date_after(tokens, start + n + 1)
            return quantities, int(total_token), start + n + 1, export_date, False
        if parse_date(total_token) or re.search(r"to be confirmed|qc|fin de chaine|début|debut|tbc", str(total_token), re.I):
            # Missing explicit total due PDF column collapse. Sometimes last size was actually the total.
            total = sum_qty
            if sizes and sizes[-1] in {"3XL", "4XL", "5XL", "XXL"}:
                prev_sum = sum(int(q) for q in qty_tokens[:-1])
                last_val = int(qty_tokens[-1]) if qty_tokens else 0
                if last_val == prev_sum and last_val > 0:
                    quantities[sizes[-1]] = 0
                    total = last_val
            export_date = find_export_date_after(tokens, start + n)
            return quantities, total, start + n, export_date, False
    # Old 2021 placeholders: only the color total is present, no size breakdown.
    # Example: sable / 546 / June 4, 2021.
    if is_numeric_token(tokens[start]):
        maybe_date = find_export_date_after(tokens, start + 1)
        # Keep total-only rows, but mark that they cannot be cartonized until size distribution is filled in.
        if maybe_date or (start + 1 < len(tokens) and not is_numeric_token(tokens[start + 1])):
            return {}, int(tokens[start]), start + 1, maybe_date, True
    return None


def is_color_code_token(v: str) -> bool:
    """Numeric colour references like 18225 / 17749 (supplier or Asphalte colour codes).

    Quantities are short numbers, colour codes are 4+ digits, so a 4-6 digit token can be a
    colour - but only if it is followed by a quantity block that adds up (checked by caller).
    """
    return bool(re.fullmatch(r"\d{4,6}", str(v or "").strip()))


def _strict_qty_block(tokens: List[str], start: int, sizes: List[str]):
    """parse_qty_block, but only accept it when the size quantities add up to the explicit total.

    Used to decide between ambiguous layouts (numeric colour code vs. quantity) safely.
    """
    parsed = parse_qty_block(tokens, start, sizes)
    if not parsed:
        return None
    quantities, total, _next_idx, _date, missing = parsed
    if missing:
        return None
    if start + len(sizes) >= len(tokens) or not is_numeric_token(tokens[start + len(sizes)]):
        return None
    if sum(quantities.values()) != total:
        return None
    return parsed


def parse_batch_section(batch_name: str, section_lines: List[str]) -> List[Dict[str, Any]]:
    sizes = extract_sizes_from_section(section_lines)
    n = len(sizes)
    has_supplier_color = any(compact_spaces(x).lower() == "supplier color" for x in section_lines[:max(60, n + 8)])
    results: List[Dict[str, Any]] = []

    i = 0
    while i < len(section_lines):
        token = compact_spaces(section_lines[i])
        if not token or is_batch_noise_line(token):
            i += 1
            continue
        if is_numeric_token(token):
            # A numeric colour code (e.g. 18225 instead of "beige") can start a row, but only
            # when a complete, self-consistent quantity row follows it.
            if not (is_color_code_token(token) and (_strict_qty_block(section_lines, i + 1, sizes) or (has_supplier_color and _strict_qty_block(section_lines, i + 2, sizes)))):
                i += 1
                continue

        if has_supplier_color:
            # Rows where the Supplier color (or the Asphalte color) is a numeric code, e.g.
            # "Beige / 18225 / 38 38 9 5 0 / 90", or where the supplier column is empty.
            # Pick the layout whose size quantities add up to the row total.
            nxt = compact_spaces(section_lines[i + 1]) if i + 1 < len(section_lines) else ""
            chosen = None
            if nxt and not is_batch_noise_line(nxt) and (not is_numeric_token(nxt) or is_color_code_token(nxt)):
                strict = _strict_qty_block(section_lines, i + 2, sizes)
                if strict:
                    chosen = (nxt, strict)
            if chosen is None and (is_numeric_token(nxt) and not is_color_code_token(nxt) or not nxt):
                strict = _strict_qty_block(section_lines, i + 1, sizes)
                if strict:
                    chosen = ("", strict)
            if chosen is None and is_color_code_token(nxt):
                # Numeric supplier code whose row total is missing/collapsed - still accept it.
                parsed = parse_qty_block(section_lines, i + 2, sizes)
                if parsed:
                    chosen = (nxt, parsed)
            if chosen is not None:
                supplier, (quantities, total, next_idx, export_date, missing) = chosen
                results.append({
                    "batch": batch_name,
                    "asphalte_color": normalize_color(token),
                    "supplier_color": supplier,
                    "sizes": quantities,
                    "total": total,
                    "export_date": export_date,
                    "size_breakdown_missing": missing,
                })
                i = max(next_idx, i + 1)
                continue
            if is_numeric_token(token):
                i += 1
                continue
            # Modern layouts: Asphalte color / Supplier color / size quantities / total / date.
            color = token
            supplier = compact_spaces(section_lines[i + 1]) if i + 1 < len(section_lines) else ""
            if not supplier or is_numeric_token(supplier) or is_batch_noise_line(supplier):
                i += 1
                continue
            parsed = parse_qty_block(section_lines, i + 2, sizes)
            if not parsed:
                i += 1
                continue
            quantities, total, next_idx, export_date, missing = parsed
            results.append({
                "batch": batch_name,
                "asphalte_color": normalize_color(color),
                "supplier_color": supplier,
                "sizes": quantities,
                "total": total,
                "export_date": export_date,
                "size_breakdown_missing": missing,
            })
            i = max(next_idx, i + 1)
            continue

        # Old one-color layouts: Reference column, Color column, then quantities.
        # Scan for a color token followed immediately by a quantity block.
        if is_color_candidate(token) or (is_color_code_token(token) and _strict_qty_block(section_lines, i + 1, sizes)):
            nxt = compact_spaces(section_lines[i + 1]) if i + 1 < len(section_lines) else ""
            code_row = _strict_qty_block(section_lines, i + 2, sizes) if is_color_code_token(nxt) else None
            if code_row and not _strict_qty_block(section_lines, i + 1, sizes):
                # "Beige / 18225 / 38 38 9 5 0 / 90": numeric colour code between colour and quantities.
                quantities, total, next_idx, export_date, missing = code_row
                results.append({
                    "batch": batch_name,
                    "asphalte_color": normalize_color(token),
                    "supplier_color": nxt,
                    "sizes": quantities,
                    "total": total,
                    "export_date": export_date,
                    "size_breakdown_missing": missing,
                })
                i = max(next_idx, i + 1)
                continue
            parsed = parse_qty_block(section_lines, i + 1, sizes)
            if parsed:
                quantities, total, next_idx, export_date, missing = parsed
                results.append({
                    "batch": batch_name,
                    "asphalte_color": normalize_color(token),
                    "supplier_color": "",
                    "sizes": quantities,
                    "total": total,
                    "export_date": export_date,
                    "size_breakdown_missing": missing,
                })
                i = max(next_idx, i + 1)
                continue
        i += 1
    return results


def _clean_order_address_lines(raw_lines: List[str], kind: str = "") -> List[str]:
    """Clean/split address blocks extracted from Asphalte order PDFs.

    The 40+ order training set contains several PDF text-layer styles:
    - modern two-column blocks: Delivery address + Supplier address side by side;
    - sequential blocks: Supplier first, Delivery below;
    - adjacent headers followed by only one visual block;
    - old PDFs where labels are read near the top but address values appear after GLOBAL ORDER.
    """
    out: List[str] = []
    skip_re = re.compile(
        r"^(billing address|date|vta|contact|mail|horaires|order details|launch|end|product information|"
        r"order placement|quality control|color breakdown|final size|supplier address|delivery address|supplier|"
        r"fabric and trims|to be confirmed|reference\s*:|asphalte color|supplier color|global order|buffer\s+\d+|batch\s+\d+)\b",
        re.I,
    )
    stop_re = re.compile(r"^(ORDER DETAILS|GLOBAL ORDER|Buffer\s+\d+|Batch\s+\d+)", re.I)
    for raw in raw_lines:
        line = compact_spaces(raw)
        if not line or stop_re.search(line):
            continue
        # Remove labels accidentally attached at the start of the line.
        line = re.sub(r"^(Delivery address|Supplier address|Supplier)\s+", "", line, flags=re.I).strip()
        if not line or skip_re.search(line):
            continue
        # Split common PDF-collapsed address rows: "ZI ... - 150 rue ...".
        pieces = [p.strip(" -") for p in re.split(r"\s+-\s+", line) if p.strip(" -")]
        for piece in pieces:
            # Split lines where the street and postal city were collapsed together.
            m = re.match(r"(.+?)(\b\d{5}\s+[A-ZÀ-Ÿ][A-ZÀ-Ÿ\s,'-]+)$", piece)
            if m and len(m.group(1).strip()) > 8:
                out.append(compact_spaces(m.group(1)))
                out.append(compact_spaces(m.group(2)))
            else:
                out.append(piece)
    # Remove repeated/empty rows while preserving order.
    dedup: List[str] = []
    seen = set()
    for line in out:
        key = normalize_key(line)
        if key and key not in seen:
            seen.add(key)
            dedup.append(line)
    if dedup and normalize_key(dedup[0]) == "logtex" and not any("france" in x.lower() for x in dedup):
        if len(dedup) >= 2:
            dedup[-1] = dedup[-1].rstrip(" ,") + ", FRANCE"
    if dedup and normalize_key(dedup[0]).startswith("pelintex"):
        dedup = [x.replace("Bulgarie", "BULGARIA").replace("bulgarie", "BULGARIA") for x in dedup]
    return dedup[:8]


def _is_delivery_start(line: str) -> bool:
    k = normalize_key(line)
    return k.startswith("logtex") or k.startswith("sasdat") or "sympl" in k or "zonedustrielle" in k or "moimont" in k


def _is_supplier_start(line: str) -> bool:
    k = normalize_key(line)
    return k.startswith("pelintex")


def _looks_bad_delivery(lines: List[str]) -> bool:
    if not lines:
        return True
    first = normalize_key(lines[0])
    joined = " ".join(lines).lower()
    return first.startswith("pelintex") or "reference" in joined or re.search(r"\b(batch|buffer)\s+\d+", joined, re.I) is not None


def _looks_bad_supplier(lines: List[str]) -> bool:
    if not lines:
        return True
    first = normalize_key(lines[0])
    joined = " ".join(lines).lower()
    return first.startswith("logtex") or first.startswith("sasdat") or "sympl" in joined or "fabric and trims" in joined or re.search(r"\b(batch|buffer)\s+\d+", joined, re.I) is not None


def _split_mixed_address_lines(block: List[str]) -> Tuple[List[str], List[str]]:
    """Split a mixed address block into delivery and supplier by known markers."""
    cleaned = [compact_spaces(x) for x in block if compact_spaces(x)]
    if not cleaned:
        return [], []
    pel_positions = [i for i, l in enumerate(cleaned) if _is_supplier_start(l)]
    del_positions = [i for i, l in enumerate(cleaned) if _is_delivery_start(l)]
    if pel_positions and del_positions:
        pel = pel_positions[-1]
        first_del = del_positions[0]
        if first_del < pel:
            return cleaned[first_del:pel], cleaned[pel:]
        if pel < first_del:
            return cleaned[first_del:], cleaned[pel:first_del]
    if pel_positions:
        return [], cleaned[pel_positions[-1]:]
    if del_positions:
        return cleaned[del_positions[0]:], []
    return [], []


def _collect_delivery_anywhere(lines: List[str]) -> List[str]:
    starts = [i for i, l in enumerate(lines) if _is_delivery_start(l)]
    if not starts:
        return []
    # Prefer LOGTEX/SAS blocks that are not inside a supplier address.
    for start in starts:
        out = []
        for l in lines[start:start + 10]:
            if out and (_is_supplier_start(l) or re.match(r"^(ORDER DETAILS|GLOBAL ORDER|PRODUCTION ORDER|Buffer\s+\d+|Batch\s+\d+)", l, re.I)):
                break
            out.append(l)
        clean = _clean_order_address_lines(out, "delivery")
        if clean and not _looks_bad_delivery(clean):
            return clean
    return []


def _collect_supplier_anywhere(lines: List[str]) -> List[str]:
    starts = [i for i, l in enumerate(lines) if _is_supplier_start(l)]
    if not starts:
        return []
    # Prefer a PELINTEX block with actual address lines; otherwise return PELINTEX only and let defaults complete it.
    best: List[str] = []
    for start in starts:
        out = []
        for l in lines[start:start + 8]:
            if out and (_is_delivery_start(l) or re.match(r"^(ORDER DETAILS|GLOBAL ORDER|PRODUCTION ORDER|Buffer\s+\d+|Batch\s+\d+)", l, re.I)):
                break
            out.append(l)
        clean = _clean_order_address_lines(out, "supplier")
        if clean and not _looks_bad_supplier(clean):
            if len(clean) > len(best):
                best = clean
    return best


def _complete_supplier(lines: List[str]) -> List[str]:
    if not lines:
        return []
    if normalize_key(lines[0]).startswith("pelintex") and len(lines) <= 1:
        return DEFAULT_SHIPPER.split("\n")
    return lines


def _complete_delivery(lines: List[str]) -> List[str]:
    if not lines:
        return []
    return lines


def _extract_legacy_parallel_addresses(lines: List[str], order_idx: int) -> Tuple[List[str], List[str]]:
    """Older SS21 PDFs read supplier and delivery columns as interleaved lines.

    Visual layout: Supplier header on the left, Delivery address header on the right;
    the text layer then reads PELINTEX, one supplier line, SAS DAT/LOGTEX, Bulgaria,
    then delivery lines. This rule separates those common old Asphalte orders.
    """
    pre = [compact_spaces(x) for x in lines[:order_idx] if compact_spaces(x)]
    supplier_header = next((i for i, l in enumerate(pre) if re.fullmatch(r"Supplier", l, re.I)), -1)
    delivery_header = next((i for i, l in enumerate(pre) if re.fullmatch(r"Delivery address", l, re.I)), -1)
    if supplier_header == -1 or delivery_header == -1:
        return [], []
    block = pre[min(supplier_header, delivery_header) + 1:]
    # Trim before order details if present.
    stop = next((i for i, l in enumerate(block) if re.search(r"^order details\b", l, re.I)), len(block))
    block = block[:stop]
    if not block:
        return [], []

    supplier: List[str] = []
    delivery: List[str] = []
    delivery_started = False
    for l in block:
        low = l.lower()
        nk = normalize_key(l)
        if nk in {"supplier", "deliveryaddress"}:
            continue
        if _is_delivery_start(l):
            delivery_started = True
            delivery.append(l)
            continue
        if nk.startswith("pelintex"):
            supplier.append(l)
            continue
        if any(x in low for x in ["hristo", "botev", "ruse", "bulgaria", "bulgarie"]):
            supplier.append(l.replace("BULGARIE", "BULGARIA").replace("Bulgarie", "BULGARIA"))
            continue
        if delivery_started or re.search(r"\b\d{5}\b|zone industrielle|rue jean|marly|cholet|moimont|sympl", low, re.I):
            delivery.append(l)
            continue
    supplier = _clean_order_address_lines(supplier, "supplier")
    delivery = _clean_order_address_lines(delivery, "delivery")
    return delivery, supplier

def extract_order_address_blocks(lines: List[str]) -> Tuple[str, str, List[str]]:
    warnings: List[str] = []
    joined = "\n".join(lines)

    # Layout A: headers appear together, values flow below them.
    m = re.search(r"Delivery address\s+Supplier address\s+(.*?)\s+ORDER DETAILS", joined, re.I | re.S)
    delivery_clean: List[str] = []
    supplier_clean: List[str] = []
    if m:
        d, s = _split_mixed_address_lines(m.group(1).split("\n"))
        delivery_clean = _clean_order_address_lines(d, "delivery")
        supplier_clean = _clean_order_address_lines(s, "supplier")

    order_idx = next((i for i, l in enumerate(lines) if re.search(r"^ORDER DETAILS\b", l, re.I)), len(lines))
    delivery_idx = next((i for i, l in enumerate(lines[:order_idx]) if re.search(r"\bDelivery address\b", l, re.I)), -1)
    supplier_idx = next((i for i, l in enumerate(lines[:order_idx]) if re.search(r"\bSupplier(?: address)?\b", l, re.I)), -1)

    if not delivery_clean and not supplier_clean:
        delivery_lines: List[str] = []
        supplier_lines: List[str] = []
        if delivery_idx != -1 and supplier_idx != -1:
            if abs(delivery_idx - supplier_idx) <= 2:
                d, s = _split_mixed_address_lines(lines[max(delivery_idx, supplier_idx) + 1:order_idx])
                delivery_lines, supplier_lines = d, s
            elif supplier_idx < delivery_idx:
                supplier_lines = lines[supplier_idx + 1:delivery_idx]
                delivery_lines = lines[delivery_idx + 1:order_idx]
            else:
                delivery_lines = lines[delivery_idx + 1:supplier_idx]
                supplier_lines = lines[supplier_idx + 1:order_idx]
        elif delivery_idx != -1:
            delivery_lines = lines[delivery_idx + 1:order_idx]
        elif supplier_idx != -1:
            supplier_lines = lines[supplier_idx + 1:order_idx]
        delivery_clean = _clean_order_address_lines(delivery_lines, "delivery")
        supplier_clean = _clean_order_address_lines(supplier_lines, "supplier")
        # If an adjacent-header layout put both address blocks in one side, split it.
        if _looks_bad_supplier(supplier_clean) or _looks_bad_delivery(delivery_clean):
            d, s = _split_mixed_address_lines((delivery_lines or []) + (supplier_lines or []))
            if d:
                delivery_clean = _clean_order_address_lines(d, "delivery")
            if s:
                supplier_clean = _clean_order_address_lines(s, "supplier")

    # Targeted scans over the whole page for older PDFs where address text is read at the bottom.
    if _looks_bad_delivery(delivery_clean):
        scanned_delivery = _collect_delivery_anywhere(lines)
        if scanned_delivery:
            delivery_clean = scanned_delivery
    if _looks_bad_supplier(supplier_clean):
        scanned_supplier = _collect_supplier_anywhere(lines)
        if scanned_supplier:
            supplier_clean = scanned_supplier

    # Old SS21 orders have supplier/delivery columns extracted as interleaved rows.
    if _looks_bad_delivery(delivery_clean) or _looks_bad_supplier(supplier_clean):
        legacy_delivery, legacy_supplier = _extract_legacy_parallel_addresses(lines, order_idx)
        if legacy_delivery and not _looks_bad_delivery(legacy_delivery):
            delivery_clean = legacy_delivery
        if legacy_supplier and not _looks_bad_supplier(legacy_supplier):
            supplier_clean = legacy_supplier

    supplier_clean = _complete_supplier(supplier_clean)
    delivery_clean = _complete_delivery(delivery_clean)

    if not delivery_clean or not supplier_clean:
        warnings.append("PDF address block used fallback extraction/defaults. Please check shipper/destination once.")
    return "\n".join(delivery_clean) if delivery_clean else "", "\n".join(supplier_clean) if supplier_clean else "", warnings


def infer_product_name_from_lines(lines: List[str]) -> str:
    # Prefer a real product/reference value from the first order table. Old 2021 PDFs often have no
    # Product information block, and a naive "Reference" lookup returns the next header: "color".
    for i, line in enumerate(lines):
        if re.fullmatch(r"Reference\s*:??", compact_spaces(line), re.I) or compact_spaces(line).lower() == "reference":
            for nxt in lines[i + 1:i + 12]:
                val = compact_spaces(nxt)
                if looks_like_product_name(val):
                    return val
    # Fallback: any strong product-looking line before GLOBAL ORDER.
    stop = next((i for i, l in enumerate(lines) if re.search(r"^GLOBAL ORDER\b", l, re.I)), len(lines))
    for val in lines[:stop]:
        if looks_like_product_name(val) and not re.search(r"billing address|asphalte|contact|supplier|delivery", val, re.I):
            return compact_spaces(val)
    return ""

def parse_asphalte_order_text(text: str) -> Dict[str, Any]:
    text = text.replace("\u2013", "-").replace("\u2014", "-")
    lines = [compact_spaces(l) for l in text.splitlines() if compact_spaces(l)]
    warnings: List[str] = []

    production_order = ""
    for line in lines[:20]:
        m = re.search(r"PRODUCTION\s+ORDER\s+([A-Z]{1,4}-?\d+)", line, re.I)
        if m:
            production_order = m.group(1).upper()
            break
    order_date = parse_label_value(lines, "Date")
    launch_date = parse_label_value(lines, "Launch")
    end_date = parse_label_value(lines, "End")
    season = parse_label_value(lines, "Season")
    product_name = parse_label_value(lines, "Name") or parse_label_value(lines, "Reference")
    if (not product_name) or product_name.lower() in {"color", "colour", "asphalte color", "supplier color"} or not looks_like_product_name(product_name):
        inferred_product = infer_product_name_from_lines(lines)
        if inferred_product:
            product_name = inferred_product
    # Some planning PDFs put MEN/WOMEN on the next line after the product name.
    if product_name:
        try:
            p_idx = next((i for i, l in enumerate(lines) if compact_spaces(l) == product_name), -1)
            if p_idx != -1 and p_idx + 1 < len(lines) and re.fullmatch(r"(MEN|MAN|WOMEN|WOMAN|FEMME|HOMME)", lines[p_idx + 1], re.I):
                product_name = compact_spaces(product_name + " " + lines[p_idx + 1].upper())
        except Exception:
            pass
    main_fabric = parse_label_value(lines, "Main fabric")
    certifications = parse_label_value(lines, "Certification")
    if certifications.lower() in {"delivery address", "supplier address", "order details"}:
        certifications = ""

    # Address blocks can be read as parallel columns or sequential blocks depending on the PDF.
    delivery_address, supplier_address, address_warnings = extract_order_address_blocks(lines)
    warnings.extend(address_warnings)
    if not delivery_address:
        delivery_address = DEFAULT_RECEIVER
    if not supplier_address:
        supplier_address = DEFAULT_SHIPPER

    # Batch sections
    batch_positions: List[Tuple[int, str]] = []
    for idx, line in enumerate(lines):
        clean = compact_spaces(line)
        if re.match(r"^(OPTION|OPTIONAL BATCH)\b", clean, re.I):
            batch_positions.append((idx, "Option"))
            continue
        # Modern: Buffer 0, Batch 1, Batch 100 RETAIL / Batch 101 - STOCK.
        # Old 2021: Buffer without a number, Batch 1 (ferme), Batch 2 (option).
        m_buffer = re.match(r"^Buffer(?:\s+(\d+))?(?:\s*(?:[-–—]\s*)?(.*?))?$", clean, re.I)
        m_batch = re.match(r"^Batch\s+(?:(RETAIL|STOCK)\s*)?(\d+(?:\.\d+)?)(?:\s*(?:[-–—]\s*)?(.*?))?$", clean, re.I)
        if m_buffer:
            num = m_buffer.group(1) or "0"
            rest = compact_spaces(m_buffer.group(2) or "")
            if len(rest) <= 35 and not re.search(r"\b(Reference|Asphalte color|Supplier color|XS|total|delivery date)\b", rest, re.I):
                batch_positions.append((idx, f"Buffer {num}" + (f" {rest}" if rest else "")))
            continue
        if m_batch:
            kind = compact_spaces(m_batch.group(1) or "")
            num = m_batch.group(2)
            rest = compact_spaces(m_batch.group(3) or "")
            # Always register the batch number even when a long planning note follows
            # (e.g. "Batch 2 - You will receive the fabric by Mid October...").
            label_rest = ""
            if len(rest) <= 45 and not re.search(r"\b(Reference|Asphalte color|Supplier color|XS|total|delivery date)\b", rest, re.I):
                label_rest = compact_spaces((kind + " " + rest).strip(" -"))
            elif kind:
                label_rest = kind
            batch_positions.append((idx, f"Batch {num}" + (f" {label_rest}" if label_rest else "")))
    order_lines: List[Dict[str, Any]] = []
    for pos_idx, (start, name) in enumerate(batch_positions):
        end = batch_positions[pos_idx + 1][0] if pos_idx + 1 < len(batch_positions) else next((i for i in range(start + 1, len(lines)) if lines[i].upper() == "GLOBAL ORDER"), len(lines))
        section = lines[start + 1:end]
        order_lines.extend(parse_batch_section(name, section))
    if not order_lines:
        warnings.append("No Buffer/Batch size lines were detected. OCR/manual check may be needed.")
    zero_lines = [r for r in order_lines if int(r.get("total", 0) or 0) == 0]
    if zero_lines and len(zero_lines) == len(order_lines):
        warnings.append("All detected order lines have zero/to-be-confirmed quantities. Packing/label exports will contain no cartons until quantities are edited.")
    elif zero_lines:
        warnings.append(f"{len(zero_lines)} detected order line(s) have zero quantity and will not create cartons unless edited.")
    missing_breakdown = [r for r in order_lines if int(r.get("total", 0) or 0) > 0 and r.get("size_breakdown_missing")]
    if missing_breakdown:
        warnings.append(f"{len(missing_breakdown)} positive order line(s) only have color totals in the PDF. Open the edit popup and fill size quantities before exporting real carton labels.")

    all_sizes = sort_sizes({s for row in order_lines for s in row.get("sizes", {}).keys()})
    total_pieces = sum(int(r.get("total", 0) or 0) for r in order_lines)
    export_dates = [parse_date(r.get("export_date", "")) for r in order_lines if parse_date(r.get("export_date", ""))]
    export_date = max(export_dates).strftime("%d/%m/%Y") if export_dates else ""

    return {
        "metadata": {
            "production_order": production_order,
            "order_date": order_date,
            "launch_date": launch_date,
            "end_date": end_date,
            "season": season,
            "product_name": product_name,
            "main_fabric": main_fabric,
            "certifications": certifications,
            "shipper": supplier_address,
            "receiver": delivery_address,
            "export_date": export_date,
        },
        "order_lines": order_lines,
        "sizes": all_sizes,
        "total_pieces": total_pieces,
        "warnings": warnings,
    }


def detect_color_from_sku_or_text(sku: str, text: str = "") -> str:
    combined = f"{sku} {text}"
    return normalize_color(combined)


def lookup_ean(ean_map: Dict[str, str], color: str, size: str) -> str:
    size = normalize_size(size)
    color_key = normalize_key(color)
    candidates = [
        f"{color_key}|{size}",
        f"{normalize_key(normalize_color(color))}|{size}",
    ]
    # Also try canonical aliases for color-specific EAN files while keeping the visible color text unchanged.
    for alias, canonical in COLOR_ALIASES.items():
        ak = normalize_key(alias)
        if ak and ak in color_key:
            candidates.append(f"{normalize_key(canonical)}|{size}")
    candidates.append(f"*|{size}")
    seen = set()
    for k in candidates:
        if k in seen:
            continue
        seen.add(k)
        if ean_map.get(k):
            return str(ean_map[k])
    return ""


def merge_ean_maps(*maps: Dict[str, Any]) -> Dict[str, Any]:
    final_map: Dict[str, str] = {}
    rows: List[Dict[str, str]] = []
    warnings: List[str] = []
    for m in maps:
        if not m:
            continue
        final_map.update(m.get("map", {}))
        rows.extend(m.get("rows", []))
        warnings.extend(m.get("warnings", []))
    return {"map": final_map, "rows": rows, "warnings": warnings}


def generate_cartons_for_line(line: Dict[str, Any], ean_map: Dict[str, str], max_pcs: int = 15, mode: str = "sequential", tare_kg: float = 1.2, unit_kg: float = 1.3, start_box: int = 1) -> List[Dict[str, Any]]:
    color = line.get("asphalte_color") or line.get("supplier_color") or ""
    sizes = {normalize_size(k): int(v or 0) for k, v in line.get("sizes", {}).items() if int(v or 0) > 0}
    ordered_sizes = sort_sizes(sizes.keys())
    cartons: List[Dict[str, Any]] = []
    box_no = start_box

    if mode == "single_size_then_mixed":
        remainder: Dict[str, int] = {}
        for size in ordered_sizes:
            qty = sizes[size]
            full = qty // max_pcs
            rem = qty % max_pcs
            for _ in range(full):
                row_sizes = {s: 0 for s in ordered_sizes}
                row_sizes[size] = max_pcs
                cartons.append(make_carton_row(box_no, line, color, row_sizes, ean_map, tare_kg, unit_kg))
                box_no += 1
            if rem:
                remainder[size] = rem
        if remainder:
            current: Dict[str, int] = {s: 0 for s in ordered_sizes}
            current_total = 0
            for size in ordered_sizes:
                qty = remainder.get(size, 0)
                while qty > 0:
                    take = min(qty, max_pcs - current_total)
                    current[size] = current.get(size, 0) + take
                    qty -= take
                    current_total += take
                    if current_total >= max_pcs:
                        cartons.append(make_carton_row(box_no, line, color, current, ean_map, tare_kg, unit_kg))
                        box_no += 1
                        current = {s: 0 for s in ordered_sizes}
                        current_total = 0
            if current_total:
                cartons.append(make_carton_row(box_no, line, color, current, ean_map, tare_kg, unit_kg))
        return cartons

    current = {s: 0 for s in ordered_sizes}
    current_total = 0
    for size in ordered_sizes:
        qty = sizes[size]
        while qty > 0:
            take = min(qty, max_pcs - current_total)
            current[size] = current.get(size, 0) + take
            qty -= take
            current_total += take
            if current_total >= max_pcs:
                cartons.append(make_carton_row(box_no, line, color, current, ean_map, tare_kg, unit_kg))
                box_no += 1
                current = {s: 0 for s in ordered_sizes}
                current_total = 0
    if current_total:
        cartons.append(make_carton_row(box_no, line, color, current, ean_map, tare_kg, unit_kg))
    return cartons


def make_carton_row(box_no: int, line: Dict[str, Any], color: str, sizes: Dict[str, int], ean_map: Dict[str, str], tare_kg: float, unit_kg: float) -> Dict[str, Any]:
    total = sum(int(v or 0) for v in sizes.values())
    eans = []
    missing = []
    for size in sort_sizes(sizes.keys()):
        if int(sizes.get(size, 0)) > 0:
            ean = lookup_ean(ean_map, color, size)
            if ean:
                eans.append(ean)
            else:
                missing.append(size)
    return {
        "line_index": line.get("line_index", 0),
        "batch": line.get("batch", ""),
        "reference": line.get("reference") or line.get("product_name") or "",
        "product": line.get("product_name") or line.get("reference") or "",
        "color": color,
        "supplier_color": line.get("supplier_color", ""),
        "box_no": box_no,
        "sizes": sizes,
        "total_pcs": total,
        "ean": ",".join(eans),
        "missing_ean_sizes": missing,
        "gross_weight": round(tare_kg + total * unit_kg, 2),
        "pl_weight": round(total * float(line.get("pl_unit_kg", 0.75) or 0.75), 2),
    }


def generate_cartons(parsed_order: Dict[str, Any], ean_map: Dict[str, str], max_pcs: int = 15, mode: str = "sequential", tare_kg: float = 1.2, unit_kg: float = 1.3) -> List[Dict[str, Any]]:
    all_cartons: List[Dict[str, Any]] = []
    for line_index, line in enumerate(parsed_order.get("order_lines", [])):
        line = dict(line)
        line["line_index"] = line.get("line_index", line_index)
        line["product_name"] = parsed_order.get("metadata", {}).get("product_name", "")
        # Per-style overrides are used immediately when rebuilding from the editor.
        local_max_pcs = int(float(line.get("pieces_per_carton_override") or max_pcs or 15))
        local_tare_kg = float(line.get("tare_kg_override") or tare_kg or 1.2)
        local_unit_kg = float(line.get("unit_kg_override") or unit_kg or 1.3)
        line["pl_unit_kg"] = float(line.get("pl_unit_kg_override") or parsed_order.get("metadata", {}).get("pl_unit_kg", 0.75) or 0.75)
        cartons = generate_cartons_for_line(line, ean_map, max_pcs=local_max_pcs, mode=mode, tare_kg=local_tare_kg, unit_kg=local_unit_kg, start_box=1)
        all_cartons.extend(cartons)
    return all_cartons


def extract_identification_from_workbook(xlsx_path: str) -> str:
    try:
        wb = load_workbook(xlsx_path, data_only=True, read_only=True)
        for ws in wb.worksheets:
            for row_idx in range(1, min(5, ws.max_row) + 1):
                for col_idx in range(1, min(12, ws.max_column) + 1):
                    val = clean_cell(ws.cell(row_idx, col_idx).value)
                    m = re.search(r"identification\s+code\s+is\s*:?\s*(\d{8,20})", val, re.I)
                    if m:
                        return m.group(1)
            # Do not read the lower "reception identification" row as the real ID.
    except Exception:
        pass
    return ""


def extract_template_metadata(xlsx_path: Optional[str]) -> Dict[str, Any]:
    if not xlsx_path or not os.path.exists(xlsx_path):
        return {}
    meta: Dict[str, Any] = {}
    try:
        wb = load_workbook(xlsx_path, data_only=True)
        ws = next((s for s in wb.worksheets if s.title.upper() != "PACKING LIST"), wb.worksheets[0])
        ident = extract_identification_from_workbook(xlsx_path)
        if ident:
            meta["identification_code"] = ident
        # Match the original merged structure: shipper values in B5:B8, destination values in J5:J8.
        shipper_lines = [clean_cell(ws[f"B{r}"].value) for r in range(5, 9)]
        receiver_lines = [clean_cell(ws[f"J{r}"].value) for r in range(5, 9)]
        if any(shipper_lines):
            meta["shipper"] = "\n".join([x for x in shipper_lines if x])
        if any(receiver_lines):
            meta["receiver"] = "\n".join([x for x in receiver_lines if x])
        ship_date = ws["L11"].value
        if isinstance(ship_date, datetime):
            meta["shipment_date"] = ship_date.strftime("%d/%m/%Y")
        else:
            parsed = parse_date(clean_cell(ship_date))
            if parsed:
                meta["shipment_date"] = parsed.strftime("%d/%m/%Y")
    except Exception:
        return meta
    return meta


def style_range_border(ws, cell_range: str, border: Border):
    for row in ws[cell_range]:
        for cell in row:
            cell.border = border


def set_row_values(ws, row: int, values: List[Any], start_col: int = 1):
    for i, v in enumerate(values, start_col):
        ws.cell(row, i).value = v


def setup_sheet_style(ws, sizes: List[str]):
    """Asphalte packing-list styling tuned to the uploaded Excel samples.

    The generated workbook intentionally keeps the same visual structure:
    A1 identification row, shipper/destination blocks, forwarder/date blocks,
    row 23 size headers, row 24 table header, grey total row and EAN summary.
    """
    thin = Side(style="thin", color="000000")
    hair = Side(style="thin", color="D9D9D9")
    medium = Side(style="medium", color="000000")
    border_all = Border(left=thin, right=thin, top=thin, bottom=thin)
    medium_box = Border(left=medium, right=medium, top=medium, bottom=medium)
    grey_fill = PatternFill("solid", fgColor="F2F2F2")
    dark_grey_fill = PatternFill("solid", fgColor="7F7F7F")
    white_fill = PatternFill("solid", fgColor="FFFFFF")

    # Fixed A:L layout like the uploaded packing lists. Extra sizes continue after L if needed.
    base_widths = {
        "A": 32.44, "B": 13.78, "C": 9.89, "D": 41.55, "E": 10.66,
        "F": 10.00, "G": 13.00, "H": 13.00, "I": 13.00, "J": 13.00,
        "K": 13.00, "L": 11.66, "M": 10.66,
    }
    max_col = max(12, 5 + len(sizes))
    for c in range(1, max_col + 1):
        col = get_column_letter(c)
        ws.column_dimensions[col].width = base_widths.get(col, 12.0)

    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins = PageMargins(left=0.7086614173, right=0.7086614173, top=0.7480314961, bottom=0.7480314961, header=0.3, footer=0.3)
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = None

    for r in range(1, 260):
        ws.row_dimensions[r].height = 15.75
    row_heights = {1: 25.5, 3: 15, 4: 14.4, 5: 14.4, 6: 14.4, 7: 14.4, 8: 14.4, 10: 14.4, 11: 14.4, 12: 14.4, 13: 14.4, 14: 14.4, 15: 14.4, 16: 14.4, 17: 14.4, 18: 14.4, 19: 14.4, 20: 14.4, 23: 14.4, 24: 28.8}
    for r, h in row_heights.items():
        ws.row_dimensions[r].height = h

    for row in ws.iter_rows(min_row=1, max_row=260, min_col=1, max_col=max_col):
        for cell in row:
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            cell.font = Font(name="Aptos Narrow", size=11)
            cell.border = Border()
            cell.fill = white_fill
    return {
        "thin": thin, "hair": hair, "medium": medium, "border_all": border_all, "medium_box": medium_box,
        "grey_fill": grey_fill, "dark_grey_fill": dark_grey_fill, "white_fill": white_fill,
    }


def safe_sheet_title(title: str) -> str:
    title = re.sub(r"[\\/*?:\[\]]", " ", title).strip()
    title = re.sub(r"\s+", " ", title)
    return title[:31] or "Sheet"


def apply_box_border(ws, min_row: int, min_col: int, max_row: int, max_col: int, side: Side):
    """Apply a strong outside border and thin inside borders to a rectangular block."""
    thin = Side(style="thin", color="000000")
    for r in range(min_row, max_row + 1):
        for c in range(min_col, max_col + 1):
            cell = ws.cell(r, c)
            cell.border = Border(
                left=side if c == min_col else thin,
                right=side if c == max_col else thin,
                top=side if r == min_row else thin,
                bottom=side if r == max_row else thin,
            )


def write_header_block(ws, metadata: Dict[str, Any], identification_code: str, shipment_date: Optional[datetime], total_cartons: int, total_weight: float, of_text: str, styles: Dict[str, Any], sizes: List[str]):
    max_col = max(12, 5 + len(sizes))
    merge_ranges = ["A1:L1", "A3:E3", "I3:L3", "A10:E10", "G10:L10", "G12:L12", "G15:L15", "G17:L17", "G19:L19", "A23:E23"]
    for rng in merge_ranges:
        try:
            ws.merge_cells(rng)
        except Exception:
            pass
    try:
        ws.merge_cells(start_row=24, start_column=6, end_row=24, end_column=max_col)
    except Exception:
        pass

    ws["A1"] = f"ASPHALTE - identification code is: {identification_code or ''}"
    ws["A1"].font = Font(name="Aptos Narrow", size=24, bold=True)
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws["A1"].fill = styles["white_fill"]

    ws["A3"] = "EXPEDITEUR/ SHIPPER"
    ws["I3"] = "DESTINATAIRE"
    for cell in (ws["A3"], ws["I3"]):
        cell.fill = styles["grey_fill"]
        cell.font = Font(name="Aptos Narrow", size=11, bold=True)
        cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)

    shipper = [x for x in str(metadata.get("shipper") or DEFAULT_SHIPPER).splitlines() if x.strip()]
    receiver = [x for x in str(metadata.get("receiver") or DEFAULT_RECEIVER).splitlines() if x.strip()]
    labels_ship = ["Nom/ Name", "Adresse / Address", "", ""]
    labels_recv = ["Nom", "Adresse", "", ""]
    for idx in range(4):
        r = 5 + idx
        ws.cell(r, 1).value = labels_ship[idx]
        ws.cell(r, 2).value = shipper[idx] if idx < len(shipper) else ""
        ws.cell(r, 9).value = labels_recv[idx]
        ws.cell(r, 10).value = receiver[idx] if idx < len(receiver) else ""
        try:
            ws.merge_cells(start_row=r, start_column=10, end_row=r, end_column=12)
        except Exception:
            pass
        ws.cell(r, 1).font = Font(name="Aptos Narrow", size=11, bold=True)
        ws.cell(r, 9).font = Font(name="Aptos Narrow", size=11, bold=True)
        ws.cell(r, 2).font = Font(name="Aptos Narrow", size=11 if idx else 14, bold=(idx == 0))
        ws.cell(r, 10).font = Font(name="Aptos Narrow", size=11 if idx else 14, bold=(idx == 0))

    # Forwarder/date metadata in the same cell coordinates as the uploaded template.
    ws["A10"] = "TRANSPORTEUR / FORWARDER"
    ws["G10"] = "DATE D'EXPEDITION/ SHIPMENT DATE"
    ws["L11"] = shipment_date
    if shipment_date:
        ws["L11"].number_format = "DD/MM/YYYY"
    ws["A12"] = "Nom/ Name"
    ws["B12"] = metadata.get("forwarder_name", "BULSTAR OUTDOOR Ltd")
    ws["A13"] = "Plate N "
    ws["B13"] = metadata.get("plate", "")
    ws["A14"] = "Contact"
    ws["B14"] = metadata.get("contact", "")
    ws["G12"] = "DATE DE LIVRAISON ESTIMÉE/ DELIVERY ESTIMATED DATE"
    ws["G15"] = "NB COLIS TOTAL/ TOTAL NUMBER OF PARCEL"
    ws["L16"] = total_cartons
    ws["G17"] = "NB PALETTE TOTAL/ TOTAL NUMBER OF PALLET"
    ws["A16"] = "N° IDENTIFICATION RECEPTION/ RECEPTION IDENTIFICATION NUMBER"
    ws["E16"] = "XXXX"
    ws["A18"] = "N°OF"
    ws["B18"] = of_text
    ws["G19"] = "TOTAL POIDS KG"
    # Formula-like text if unit is unknown; otherwise number. Keep as number from builder.
    ws["L20"] = total_weight

    for cell in ["A10", "G10", "A12", "G12", "G15", "G17", "A18", "G19"]:
        ws[cell].font = Font(name="Aptos Narrow", size=11, bold=True)
        ws[cell].fill = styles["grey_fill"]
    for cell in ["B12", "B13", "B14", "B18", "L11", "L16", "L20", "E16"]:
        ws[cell].font = Font(name="Aptos Narrow", size=11)
        ws[cell].alignment = Alignment(horizontal="left" if cell.startswith("B") else "center", vertical="center", wrap_text=True)

    apply_box_border(ws, 3, 1, 8, 5, styles["medium"])
    apply_box_border(ws, 3, 9, 8, 12, styles["medium"])
    apply_box_border(ws, 10, 1, 14, 5, styles["medium"])
    apply_box_border(ws, 10, 7, 20, 12, styles["medium"])
    apply_box_border(ws, 16, 1, 20, 12, styles["medium"])

    # Table header block exactly A23:L24 for default XS-3XL sizes.
    for c in range(1, max_col + 1):
        ws.cell(23, c).border = styles["border_all"]
        ws.cell(24, c).border = styles["border_all"]
        ws.cell(23, c).font = Font(name="Aptos Narrow", size=11, bold=True)
        ws.cell(24, c).font = Font(name="Aptos Narrow", size=11, bold=True)
        ws.cell(23, c).alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.cell(24, c).alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for c in range(1, 6):
        ws.cell(23, c).fill = styles["dark_grey_fill"]
    for idx, sz in enumerate(sizes, 6):
        ws.cell(23, idx).value = sz
        ws.cell(23, idx).fill = styles["white_fill"]
    headers = ["Référence", "Coloris", "N° de colis", "EAN ", "Nbre de Pcs/Carton"]
    set_row_values(ws, 24, headers, 1)
    ws.cell(24, 6).value = "Quantité / Taille"


def parse_existing_packing_list_xlsx(xlsx_path: str) -> Dict[str, Any]:
    wb = load_workbook(xlsx_path, data_only=False)
    cartons: List[Dict[str, Any]] = []
    warnings: List[str] = []
    identification_code = extract_identification_from_workbook(xlsx_path)
    meta = extract_template_metadata(xlsx_path)
    meta.setdefault("identification_code", identification_code)
    for ws in wb.worksheets:
        if ws.title.upper().strip() == "PACKING LIST":
            continue
        matrix = [[clean_cell(ws.cell(r, c).value) for c in range(1, ws.max_column + 1)] for r in range(1, ws.max_row + 1)]
        size_row = -1
        size_cols: Dict[str, int] = {}
        for ridx, row in enumerate(matrix, 1):
            temp = {}
            for cidx, val in enumerate(row, 1):
                sz = normalize_size(val)
                if sz in SIZE_ORDER:
                    temp[sz] = cidx
            if len(temp) >= 2:
                size_row = ridx
                size_cols = temp
                break
        if size_row == -1:
            warnings.append(f"Sheet {ws.title}: size header not found.")
            continue
        header_row = size_row + 1
        cols = {"ref": 1, "color": 2, "box": 3, "ean": 4, "pcs": 5}
        for cidx in range(1, ws.max_column + 1):
            val = clean_cell(ws.cell(header_row, cidx).value).lower()
            if "référence" in val or "reference" in val:
                cols["ref"] = cidx
            elif "color" in val or "coloris" in val:
                cols["color"] = cidx
            elif "colis" in val or "carton" in val:
                cols["box"] = cidx
            elif "ean" in val:
                cols["ean"] = cidx
            elif "pcs" in val or "nbre" in val:
                cols["pcs"] = cidx
        r = header_row + 1
        while r <= ws.max_row:
            row_text = " ".join(clean_cell(ws.cell(r, c).value) for c in range(1, min(ws.max_column, 12) + 1)).lower()
            if "total" in row_text and clean_cell(ws.cell(r, 1).value).lower().startswith("total"):
                break
            box_val = clean_cell(ws.cell(r, cols["box"]).value)
            if not box_val:
                r += 1
                continue
            try:
                box_no = int(float(re.sub(r"[^0-9.]", "", box_val)))
            except Exception:
                r += 1
                continue
            sizes = {}
            for sz, cidx in size_cols.items():
                val = ws.cell(r, cidx).value
                try:
                    q = int(float(val)) if val not in (None, "") else 0
                except Exception:
                    q = 0
                if q > 0:
                    sizes[sz] = q
            if sizes:
                cartons.append({
                    "sheet": ws.title,
                    "batch": ws.title,
                    "reference": clean_cell(ws.cell(r, cols["ref"]).value),
                    "product": clean_cell(ws.cell(r, cols["ref"]).value),
                    "color": normalize_color(clean_cell(ws.cell(r, cols["color"]).value)),
                    "box_no": box_no,
                    "sizes": sizes,
                    "total_pcs": sum(sizes.values()),
                    "ean": clean_cell(ws.cell(r, cols["ean"]).value),
                    "gross_weight": 0,
                    "missing_ean_sizes": [],
                })
            r += 1
    return {"metadata": meta, "cartons": cartons, "warnings": warnings}


def draw_wrapped(c: canvas.Canvas, text: str, x: float, y: float, max_width: float, font: str = "Helvetica", size: int = 9, leading: float = 11):
    c.setFont(font, size)
    lines_out = []
    for line in str(text or "").splitlines():
        words = line.split()
        current = ""
        for word in words:
            test = (current + " " + word).strip()
            if c.stringWidth(test, font, size) <= max_width or not current:
                current = test
            else:
                lines_out.append(current)
                current = word
        if current:
            lines_out.append(current)
    for idx, line in enumerate(lines_out):
        c.drawString(x, y - idx * leading, line)
    return y - len(lines_out) * leading


def draw_ean13(c: canvas.Canvas, ean: str, x: float, y: float, width: float = 50 * mm, height: float = 18 * mm):
    ean = re.sub(r"\D", "", str(ean or ""))
    if len(ean) not in (12, 13):
        return False
    try:
        widget = eanbc.Ean13BarcodeWidget(ean[:12])
        bounds = widget.getBounds()
        bw = bounds[2] - bounds[0]
        bh = bounds[3] - bounds[1]
        d = Drawing(width, height, transform=[width / bw, 0, 0, height / bh, 0, 0])
        d.add(widget)
        renderPDF.draw(d, c, x, y)
        return True
    except Exception:
        return False


def _ean13_check_digit(first12: str) -> str:
    total = 0
    for i, ch in enumerate(first12):
        n = int(ch)
        total += n if i % 2 == 0 else n * 3
    return str((10 - (total % 10)) % 10)


def _barcode_png_path(ean: str, tmp_dir: str) -> Optional[str]:
    """Create a PNG EAN-13 barcode for Excel insertion.

    Uses python-barcode when available, otherwise a small built-in EAN-13 renderer
    with Pillow, so labels still contain real barcodes even on machines without the
    optional barcode package already installed.
    """
    ean_digits = re.sub(r"\D", "", str(ean or ""))
    if len(ean_digits) == 12:
        ean_full = ean_digits + _ean13_check_digit(ean_digits)
    elif len(ean_digits) == 13:
        ean_full = ean_digits
        ean_digits = ean_full[:12]
    else:
        return None
    os.makedirs(tmp_dir, exist_ok=True)

    # Preferred: external writer, if installed.
    if BarcodeEAN13 is not None and BarcodeImageWriter is not None:
        try:
            base = os.path.join(tmp_dir, f"ean_{ean_full}_{uuid.uuid4().hex[:8]}")
            full = BarcodeEAN13(ean_digits, writer=BarcodeImageWriter()).save(base, options={
                "module_width": 0.26,
                "module_height": 13.0,
                "quiet_zone": 2.0,
                "font_size": 8,
                "text_distance": 2.0,
                "write_text": True,
                "dpi": 160,
            })
            return full
        except Exception:
            pass

    # Built-in fallback using Pillow.
    if Image is None:
        return None
    try:
        L = {
            "0":"0001101","1":"0011001","2":"0010011","3":"0111101","4":"0100011",
            "5":"0110001","6":"0101111","7":"0111011","8":"0110111","9":"0001011",
        }
        G = {
            "0":"0100111","1":"0110011","2":"0011011","3":"0100001","4":"0011101",
            "5":"0111001","6":"0000101","7":"0010001","8":"0001001","9":"0010111",
        }
        R = {
            "0":"1110010","1":"1100110","2":"1101100","3":"1000010","4":"1011100",
            "5":"1001110","6":"1010000","7":"1000100","8":"1001000","9":"1110100",
        }
        PAR = {
            "0":"LLLLLL","1":"LLGLGG","2":"LLGGLG","3":"LLGGGL","4":"LGLLGG",
            "5":"LGGLLG","6":"LGGGLL","7":"LGLGLG","8":"LGLGGL","9":"LGGLGL",
        }
        first = ean_full[0]
        left = ean_full[1:7]
        right = ean_full[7:13]
        bits = "101"
        for digit, parity in zip(left, PAR[first]):
            bits += (L if parity == "L" else G)[digit]
        bits += "01010"
        for digit in right:
            bits += R[digit]
        bits += "101"
        module = 3
        quiet = 12
        bar_h = 74
        text_h = 22
        w = len(bits) * module + 2 * quiet
        h = bar_h + text_h
        img = Image.new("RGB", (w, h), "white")
        from PIL import ImageDraw, ImageFont
        draw = ImageDraw.Draw(img)
        for i, bit in enumerate(bits):
            if bit == "1":
                x = quiet + i * module
                draw.rectangle([x, 4, x + module - 1, bar_h], fill="black")
        try:
            font = ImageFont.truetype("arial.ttf", 13)
        except Exception:
            font = ImageFont.load_default()
        text = ean_full
        try:
            bbox = draw.textbbox((0, 0), text, font=font)
            tw = bbox[2] - bbox[0]
        except Exception:
            tw = len(text) * 7
        draw.text(((w - tw) // 2, bar_h + 3), text, fill="black", font=font)
        out = os.path.join(tmp_dir, f"ean_{ean_full}_{uuid.uuid4().hex[:8]}.png")
        img.save(out)
        return out
    except Exception:
        return None


def build_labels_pdf(output_path: str, cartons: List[Dict[str, Any]], metadata: Dict[str, Any], identification_code: str = "") -> Dict[str, Any]:
    c = canvas.Canvas(output_path, pagesize=landscape(A4))
    w, h = landscape(A4)
    for idx, carton in enumerate(cartons, 1):
        c.setFillColor(colors.white)
        c.rect(0, 0, w, h, fill=1, stroke=0)
        margin = 12 * mm
        c.setStrokeColor(colors.black)
        c.setLineWidth(1.2)
        c.rect(margin, margin, w - 2 * margin, h - 2 * margin)
        c.setFont("Helvetica-Bold", 18)
        c.drawCentredString(w / 2, h - 22 * mm, "ASPHALTE CARTON LABEL")
        c.setFont("Helvetica", 9)
        c.drawRightString(w - margin, h - 16 * mm, f"Page {idx}/{len(cartons)}")
        c.setFont("Helvetica-Bold", 10)
        c.drawString(margin + 5 * mm, h - 34 * mm, "Identification code")
        c.setFont("Helvetica", 10)
        c.drawString(margin + 50 * mm, h - 34 * mm, identification_code or metadata.get("identification_code", ""))
        c.setFont("Helvetica-Bold", 10)
        c.drawString(margin + 5 * mm, h - 43 * mm, "N° OF / Batch")
        c.setFont("Helvetica", 10)
        c.drawString(margin + 50 * mm, h - 43 * mm, str(carton.get("batch", ""))[:80])
        c.setFont("Helvetica-Bold", 28)
        c.drawRightString(w - margin - 5 * mm, h - 36 * mm, f"BOX {carton.get('box_no', idx)}")

        # Shipper / receiver blocks
        top_y = h - 60 * mm
        c.setFont("Helvetica-Bold", 10)
        c.drawString(margin + 5 * mm, top_y, "EXPEDITEUR / SHIPPER")
        c.drawString(w / 2 + 5 * mm, top_y, "DESTINATAIRE")
        c.setStrokeColor(colors.lightgrey)
        c.rect(margin + 4 * mm, top_y - 40 * mm, w / 2 - margin - 10 * mm, 36 * mm)
        c.rect(w / 2 + 4 * mm, top_y - 40 * mm, w / 2 - margin - 10 * mm, 36 * mm)
        draw_wrapped(c, metadata.get("shipper", DEFAULT_SHIPPER), margin + 8 * mm, top_y - 10 * mm, w / 2 - margin - 18 * mm, size=10)
        draw_wrapped(c, metadata.get("receiver", DEFAULT_RECEIVER), w / 2 + 8 * mm, top_y - 10 * mm, w / 2 - margin - 18 * mm, size=10)

        # Product summary
        prod_y = top_y - 52 * mm
        c.setFont("Helvetica-Bold", 11)
        c.drawString(margin + 5 * mm, prod_y, "Reference")
        c.drawString(margin + 90 * mm, prod_y, "Color")
        c.drawString(margin + 140 * mm, prod_y, "Total pcs")
        c.drawString(margin + 175 * mm, prod_y, "Gross weight")
        c.setFont("Helvetica", 11)
        c.drawString(margin + 5 * mm, prod_y - 8 * mm, str(carton.get("reference", ""))[:45])
        c.drawString(margin + 90 * mm, prod_y - 8 * mm, str(carton.get("color", ""))[:25])
        c.drawString(margin + 140 * mm, prod_y - 8 * mm, str(carton.get("total_pcs", "")))
        weight = carton.get("gross_weight", "")
        c.drawString(margin + 175 * mm, prod_y - 8 * mm, f"{weight} kg" if weight else "")

        # Size table
        table_y = prod_y - 24 * mm
        active_sizes = [(s, int(q or 0)) for s, q in carton.get("sizes", {}).items() if int(q or 0) > 0]
        active_sizes = [(s, q) for s, q in sorted(active_sizes, key=lambda x: SIZE_ORDER.index(x[0]) if x[0] in SIZE_ORDER else 99)]
        cell_w = 22 * mm
        c.setFont("Helvetica-Bold", 10)
        for i, (sz, _) in enumerate(active_sizes):
            x = margin + 5 * mm + i * cell_w
            c.rect(x, table_y, cell_w, 8 * mm)
            c.drawCentredString(x + cell_w / 2, table_y + 2.4 * mm, sz)
        c.setFont("Helvetica", 11)
        for i, (_, qty) in enumerate(active_sizes):
            x = margin + 5 * mm + i * cell_w
            c.rect(x, table_y - 8 * mm, cell_w, 8 * mm)
            c.drawCentredString(x + cell_w / 2, table_y - 5.5 * mm, str(qty))

        # EAN barcodes for every active size, not only first three.
        barcode_y = margin + 18 * mm
        c.setFont("Helvetica-Bold", 9)
        c.drawString(margin + 5 * mm, barcode_y + 28 * mm, "EAN barcodes")
        eans = [e.strip() for e in str(carton.get("ean", "")).split(",") if e.strip()]
        x0 = margin + 5 * mm
        for i, ean in enumerate(eans[:8]):
            bx = x0 + (i % 4) * 65 * mm
            by = barcode_y + (1 - (i // 4)) * 22 * mm if i < 8 else barcode_y
            ok = draw_ean13(c, ean, bx, by, width=48 * mm, height=14 * mm)
            c.setFont("Helvetica", 7)
            c.drawCentredString(bx + 24 * mm, by - 3 * mm, ean if ok else f"Invalid EAN: {ean}")
        if not eans:
            c.setFont("Helvetica", 10)
            c.setFillColor(colors.red)
            c.drawString(x0, barcode_y + 14 * mm, "No EAN barcode for this carton. Add EAN manually or upload EAN/SKU file.")
            c.setFillColor(colors.black)
        c.showPage()
    c.save()
    return {"path": output_path, "labels": len(cartons)}


def build_pallet_plan(cartons: List[Dict[str, Any]], pallet_capacity: int = 16) -> List[Dict[str, Any]]:
    """Split carton runs into pallet labels like the uploaded Asphalte pallet template.

    Pallets are filled sequentially with a default 16 cartons per pallet. If a pallet
    finishes one style/color and still has free places, the next style/color starts on
    the same pallet as an additional block. Each block shows its own Box N range.
    """
    pallet_capacity = max(1, int(pallet_capacity or 16))
    groups: List[Tuple[Any, List[Dict[str, Any]]]] = []
    lookup: Dict[Any, int] = {}
    for c in cartons or []:
        key = c.get("line_index")
        if key is None:
            key = f"{c.get('batch','')}|{c.get('reference','')}|{c.get('color','')}"
        if key not in lookup:
            lookup[key] = len(groups)
            groups.append((key, []))
        groups[lookup[key]][1].append(c)
    pallets: List[Dict[str, Any]] = []
    cur = {"number": 1, "boxes": 0, "segments": []}
    for key, rows in groups:
        rows = sorted(rows, key=lambda x: int(x.get("box_no", 0) or 0))
        i = 0
        while i < len(rows):
            if cur["boxes"] >= pallet_capacity:
                pallets.append(cur)
                cur = {"number": len(pallets) + 1, "boxes": 0, "segments": []}
            free = pallet_capacity - cur["boxes"]
            take = min(free, len(rows) - i)
            if take <= 0:
                continue
            seg_rows = rows[i:i + take]
            first, last = seg_rows[0], seg_rows[-1]
            cur["segments"].append({
                "key": key,
                "line_index": first.get("line_index"),
                "batch": first.get("batch", ""),
                "reference": first.get("reference") or first.get("product") or "",
                "color": first.get("color", ""),
                "from_box": int(first.get("box_no", 1) or 1),
                "to_box": int(last.get("box_no", first.get("box_no", 1)) or 1),
                "boxes": take,
                "pieces": sum(int(r.get("total_pcs", 0) or 0) for r in seg_rows),
            })
            cur["boxes"] += take
            i += take
            if cur["boxes"] >= pallet_capacity:
                pallets.append(cur)
                cur = {"number": len(pallets) + 1, "boxes": 0, "segments": []}
    if cur.get("segments"):
        pallets.append(cur)
    total = len(pallets)
    for idx, pal in enumerate(pallets, 1):
        pal["number"] = idx
        pal["total_pallets"] = total
    return pallets


def build_pallet_labels_xlsx(output_path: str, cartons: List[Dict[str, Any]], metadata: Dict[str, Any], pallet_capacity: int = 16) -> Dict[str, Any]:
    """Create Asphalte-style pallet labels in Excel.

    Layout follows the uploaded Pallet LABEL.xlsx: one sheet per pallet, stacked
    blocks when one pallet contains several batch/color ranges.
    """
    pallets = build_pallet_plan(cartons, pallet_capacity=pallet_capacity)
    wb = Workbook()
    if wb.active:
        wb.remove(wb.active)
    medium = Side(style="medium", color="000000")
    thin = Side(style="thin", color="000000")
    white = PatternFill("solid", fgColor="FFFFFF")
    grey = PatternFill("solid", fgColor="F2F2F2")

    def set_cell(cell, value="", size=18, bold=False, fill=None, align="center"):
        cell.value = value
        cell.font = Font(name="Arial", size=size, bold=bold)
        cell.alignment = Alignment(horizontal=align, vertical="center", wrap_text=True)
        if fill:
            cell.fill = fill

    def merge(ws, r1, c1, r2, c2):
        try:
            ws.merge_cells(start_row=r1, start_column=c1, end_row=r2, end_column=c2)
        except Exception:
            pass
        return ws.cell(r1, c1)

    def outside(ws, r1, c1, r2, c2):
        for rr in range(r1, r2 + 1):
            for cc in range(c1, c2 + 1):
                ws.cell(rr, cc).border = Border(
                    left=medium if cc == c1 else thin,
                    right=medium if cc == c2 else thin,
                    top=medium if rr == r1 else thin,
                    bottom=medium if rr == r2 else thin,
                )

    if not pallets:
        ws = wb.create_sheet("pallet labels")
        ws["B1"] = "NO CARTONS / NO PALLET LABELS"
        ws["B1"].font = Font(name="Arial", size=22, bold=True)
        wb.save(output_path)
        return {"path": output_path, "pallets": 0}

    for pal in pallets:
        title = safe_sheet_title(f"pallet {pal['number']}-{pal['total_pallets']}")
        ws = wb.create_sheet(title)
        ws.sheet_view.showGridLines = False
        ws.page_setup.orientation = "landscape"
        ws.page_setup.paperSize = ws.PAPERSIZE_A4
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 1
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_margins = PageMargins(left=0.25, right=0.25, top=0.25, bottom=0.25, header=0.1, footer=0.1)
        widths = {"A": 4, "B": 22, "C": 22, "D": 22, "E": 24, "F": 16, "G": 38}
        for col, w in widths.items():
            ws.column_dimensions[col].width = w
        for rr in range(1, 80):
            ws.row_dimensions[rr].height = 44 if rr % 7 in (1,2,4,5,6) else 28
            for cc in range(1, 8):
                ws.cell(rr, cc).fill = white
                ws.cell(rr, cc).alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        for bi, seg in enumerate(pal.get("segments", [])):
            r = 1 + bi * 7
            merge(ws, r, 2, r, 4); set_cell(ws.cell(r, 2), "PALLET No", 34, True)
            set_cell(ws.cell(r, 5), pal["number"], 34, True)
            set_cell(ws.cell(r, 6), "of", 28, True)
            set_cell(ws.cell(r, 7), pal["total_pallets"], 34, True)
            merge(ws, r + 1, 2, r + 1, 4); set_cell(ws.cell(r + 1, 2), "TOTAL BOXES", 26, False)
            merge(ws, r + 1, 5, r + 1, 7); set_cell(ws.cell(r + 1, 5), pal["boxes"], 28, True)
            merge(ws, r + 2, 2, r + 2, 7); set_cell(ws.cell(r + 2, 2), "", 12)
            merge(ws, r + 3, 2, r + 3, 7)
            batch = compact_spaces(str(seg.get("batch") or "")) or "BATCH"
            set_cell(ws.cell(r + 3, 2), f"PACKING LIST ASPHALT / {batch.upper()}", 26, True)
            merge(ws, r + 4, 2, r + 4, 4); set_cell(ws.cell(r + 4, 2), "Refference", 24, False)
            merge(ws, r + 4, 5, r + 4, 7)
            ref = compact_spaces(str(seg.get("reference") or metadata.get("product_name") or ""))
            color = compact_spaces(str(seg.get("color") or ""))
            ref_text = (ref + (f"  col.{color}" if color else "")).strip()
            set_cell(ws.cell(r + 4, 5), ref_text, 22, True)
            merge(ws, r + 5, 2, r + 5, 4); set_cell(ws.cell(r + 5, 2), "Box N", 24, False)
            merge(ws, r + 5, 5, r + 5, 7)
            set_cell(ws.cell(r + 5, 5), f"FROM {seg.get('from_box')} to {seg.get('to_box')}", 24, True)
            outside(ws, r, 2, r + 5, 7)
            for c in range(2, 8):
                ws.cell(r + 2, c).border = Border(bottom=medium, left=thin, right=thin, top=thin)
        last_row = 6 + (len(pal.get("segments", [])) - 1) * 7
        ws.print_area = f"B1:G{last_row}"
    wb.save(output_path)
    return {"path": output_path, "pallets": len(pallets), "capacity": max(1, int(pallet_capacity or 16))}

def make_job_zip(zip_path: str, files: Dict[str, str]):
    with zipfile.ZipFile(zip_path, "w", zipfile.ZIP_DEFLATED) as zf:
        for arcname, fpath in files.items():
            if fpath and os.path.exists(fpath):
                zf.write(fpath, arcname)
    return zip_path


def safe_filename_part(value: str, fallback: str = "file") -> str:
    value = compact_spaces(str(value or ""))
    value = re.sub(r"[^A-Za-z0-9._ -]+", "", value).strip().replace(" ", "_")
    value = re.sub(r"_+", "_", value)
    return value[:80] or fallback


def ean_rows_to_map(ean_rows: List[Dict[str, str]]) -> Dict[str, str]:
    result: Dict[str, str] = {}
    for row in ean_rows or []:
        ean = str(row.get("ean", "")).strip()
        size = normalize_size(row.get("size", ""))
        color = row.get("color", "")
        if ean and size:
            result[f"{normalize_key(color)}|{size}"] = ean
            result.setdefault(f"*|{size}", ean)
    return result


def build_individual_outputs(work_dir: str, parsed_order: Dict[str, Any], cartons: List[Dict[str, Any]], ean_rows: List[Dict[str, str]], identification_code: str, shipment_date_text: str = "", prefix: str = "order_line", pl_unit_kg: float = 0.75, pallet_capacity: int = 16) -> List[Dict[str, Any]]:
    out_dir = os.path.join(work_dir, "individual")
    os.makedirs(out_dir, exist_ok=True)
    outputs: List[Dict[str, Any]] = []
    groups: Dict[Any, List[Dict[str, Any]]] = {}
    for c in cartons:
        key = c.get("line_index")
        if key is None:
            key = f"{c.get('batch','')}|{c.get('color','')}|{c.get('sheet','')}"
        groups.setdefault(key, []).append(c)

    for n, (key, rows) in enumerate(groups.items(), 1):
        if not rows:
            continue
        try:
            line_index = int(key)
        except Exception:
            line_index = n - 1
        order_lines = parsed_order.get("order_lines", [])
        line = dict(order_lines[line_index]) if 0 <= line_index < len(order_lines) else {
            "batch": rows[0].get("batch", ""),
            "asphalte_color": rows[0].get("color", ""),
            "supplier_color": rows[0].get("supplier_color", ""),
            "sizes": {},
            "total": sum(int(r.get("total_pcs", 0) or 0) for r in rows),
        }
        line["line_index"] = line_index
        parsed_one = dict(parsed_order)
        parsed_one["order_lines"] = [line]
        parsed_one["sizes"] = sort_sizes({s for r in rows for s, q in r.get("sizes", {}).items() if int(q or 0) > 0})
        name = safe_filename_part(f"{line.get('batch','Line_'+str(n))}_{line.get('asphalte_color') or rows[0].get('color','')}", f"line_{n}")
        xlsx_path = os.path.join(out_dir, f"{name}_packing_list.xlsx")
        labels_xlsx_path = os.path.join(out_dir, f"{name}_labels.xlsx")
        pallet_labels_xlsx_path = os.path.join(out_dir, f"{name}_pallet_labels.xlsx")
        build_packing_list_xlsx(xlsx_path, parsed_one, rows, ean_rows, metadata_overrides=parsed_order.get("metadata", {}), identification_code=identification_code, shipment_date_text=shipment_date_text, pl_unit_kg=pl_unit_kg)
        build_labels_xlsx(labels_xlsx_path, rows, parsed_order.get("metadata", {}), identification_code=identification_code)
        build_pallet_labels_xlsx(pallet_labels_xlsx_path, rows, parsed_order.get("metadata", {}), pallet_capacity=pallet_capacity)
        outputs.append({
            "id": f"line_{line_index}",
            "line_index": line_index,
            "name": name.replace("_", " "),
            "batch": line.get("batch", ""),
            "color": line.get("asphalte_color") or rows[0].get("color", ""),
            "cartons": len(rows),
            "pieces": sum(int(r.get("total_pcs", 0) or 0) for r in rows),
            "xlsx": xlsx_path,
            "labels_xlsx": labels_xlsx_path,
            "pallet_labels_xlsx": pallet_labels_xlsx_path,
        })
    return outputs


def analyze_orders_zip(zip_path: str, work_dir: str, ocr_mode: str = "never") -> Dict[str, Any]:
    """Analyze many Asphalte order PDFs and create JSON/CSV parser audit reports."""
    extract_dir = os.path.join(work_dir, "training_orders")
    os.makedirs(extract_dir, exist_ok=True)
    with zipfile.ZipFile(zip_path, "r") as zf:
        zf.extractall(extract_dir)
    rows: List[Dict[str, Any]] = []
    for pdf in sorted(Path(extract_dir).rglob("*.pdf")):
        try:
            parsed, ocr_warnings = parse_order_pdf(str(pdf), ocr_mode=ocr_mode)
            if not parsed.get("metadata", {}).get("production_order"):
                m_file = re.search(r"\b(OF[-_ ]?\d+)\b", pdf.name, re.I)
                if m_file:
                    parsed["metadata"]["production_order"] = m_file.group(1).replace("_", "-").replace(" ", "-").upper()
            positive_lines = [l for l in parsed.get("order_lines", []) if int(l.get("total", 0) or 0) > 0]
            zero_lines = [l for l in parsed.get("order_lines", []) if int(l.get("total", 0) or 0) == 0]
            missing_breakdown_lines = [l for l in parsed.get("order_lines", []) if l.get("size_breakdown_missing")]
            rows.append({
                "file": str(pdf.relative_to(extract_dir)),
                "production_order": parsed.get("metadata", {}).get("production_order", ""),
                "product": parsed.get("metadata", {}).get("product_name", ""),
                "order_lines": len(parsed.get("order_lines", [])),
                "positive_lines": len(positive_lines),
                "zero_lines": len(zero_lines),
                "missing_size_breakdown_lines": len(missing_breakdown_lines),
                "total_pieces": parsed.get("total_pieces", 0),
                "shipper_first_line": (parsed.get("metadata", {}).get("shipper", "").split("\n") or [""])[0],
                "receiver_first_line": (parsed.get("metadata", {}).get("receiver", "").split("\n") or [""])[0],
                "warnings": "; ".join((ocr_warnings or []) + parsed.get("warnings", [])),
                "status": "OK" if parsed.get("order_lines") else "CHECK",
            })
        except Exception as exc:
            rows.append({"file": str(pdf.relative_to(extract_dir)), "status": "ERROR", "warnings": str(exc)})
    json_path = os.path.join(work_dir, "training_audit.json")
    csv_path = os.path.join(work_dir, "training_audit.csv")
    with open(json_path, "w", encoding="utf-8") as f:
        json.dump(rows, f, ensure_ascii=False, indent=2)
    import csv
    fieldnames = ["file", "production_order", "product", "order_lines", "positive_lines", "zero_lines", "missing_size_breakdown_lines", "total_pieces", "shipper_first_line", "receiver_first_line", "status", "warnings"]
    with open(csv_path, "w", encoding="utf-8-sig", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=fieldnames)
        writer.writeheader()
        for row in rows:
            writer.writerow({k: row.get(k, "") for k in fieldnames})
    ok = sum(1 for r in rows if r.get("status") == "OK")
    errors = sum(1 for r in rows if r.get("status") == "ERROR")
    positive_files = sum(1 for r in rows if int(r.get("total_pieces", 0) or 0) > 0)
    return {
        "json": json_path,
        "csv": csv_path,
        "rows": rows,
        "summary": {
            "files": len(rows),
            "ok": ok,
            "errors": errors,
            "positive_files": positive_files,
            "zero_or_planning_files": len(rows) - positive_files - errors,
            "total_detected_pieces": sum(int(r.get("total_pieces", 0) or 0) for r in rows),
        }
    }

def build_from_edited_payload(base_payload: Dict[str, Any], work_dir: str, order_lines: List[Dict[str, Any]], max_pcs: int = 15, carton_mode: str = "sequential", tare_kg: float = 1.2, unit_kg: float = 1.3, pl_unit_kg: float = 0.75, shipment_date: str = "", identification_code: str = "", pallet_capacity: int = 16) -> Dict[str, Any]:
    parsed = dict(base_payload.get("parsed_order", {}))
    metadata = dict(parsed.get("metadata", {}))
    if shipment_date:
        metadata["shipment_date"] = shipment_date
    if identification_code:
        metadata["identification_code"] = identification_code
    metadata["pl_unit_kg"] = pl_unit_kg
    parsed["metadata"] = metadata
    clean_lines: List[Dict[str, Any]] = []
    for idx, line in enumerate(order_lines):
        line = dict(line)
        sizes = {normalize_size(k): int(v or 0) for k, v in (line.get("sizes") or {}).items() if is_supported_size(k) and int(v or 0) > 0}
        line["sizes"] = sizes
        line["total"] = sum(sizes.values())
        line["line_index"] = idx
        if line.get("pl_unit_kg_override") in (None, ""):
            line["pl_unit_kg_override"] = pl_unit_kg
        line["asphalte_color"] = normalize_color(line.get("asphalte_color") or line.get("color") or "")
        clean_lines.append(line)
    parsed["order_lines"] = clean_lines
    parsed["sizes"] = sort_sizes({s for line in clean_lines for s in line.get("sizes", {}).keys()})
    parsed["total_pieces"] = sum(int(line.get("total", 0) or 0) for line in clean_lines)
    ean_rows = base_payload.get("ean_rows", [])
    ean_map = ean_rows_to_map(ean_rows)
    cartons = generate_cartons(parsed, ean_map, max_pcs=max_pcs, mode=carton_mode, tare_kg=tare_kg, unit_kg=unit_kg)
    for c in cartons:
        c["reference"] = parsed.get("metadata", {}).get("product_name", c.get("reference", ""))
        c["product"] = parsed.get("metadata", {}).get("product_name", c.get("product", ""))
    shipment_date_text = shipment_date or parsed.get("metadata", {}).get("shipment_date", "")
    if not identification_code:
        identification_code = parsed.get("metadata", {}).get("identification_code", "") or identification_from_shipment_date(parse_date(shipment_date_text or parsed.get("metadata", {}).get("export_date", "")))
        parsed["metadata"]["identification_code"] = identification_code
    xlsx_path = os.path.join(work_dir, "generated_packing_list.xlsx")
    labels_xlsx_path = os.path.join(work_dir, "generated_labels.xlsx")
    pallet_labels_xlsx_path = os.path.join(work_dir, "generated_pallet_labels.xlsx")
    json_path = os.path.join(work_dir, "parsed_data.json")
    zip_path = os.path.join(work_dir, "asphalte_output.zip")
    build_packing_list_xlsx(xlsx_path, parsed, cartons, ean_rows, metadata_overrides=parsed.get("metadata", {}), identification_code=identification_code, shipment_date_text=shipment_date_text, pl_unit_kg=pl_unit_kg)
    build_labels_xlsx(labels_xlsx_path, cartons, parsed.get("metadata", {}), identification_code=identification_code)
    pallet_result = build_pallet_labels_xlsx(pallet_labels_xlsx_path, cartons, parsed.get("metadata", {}), pallet_capacity=pallet_capacity)
    individual = build_individual_outputs(work_dir, parsed, cartons, ean_rows, identification_code, shipment_date_text=shipment_date_text, pl_unit_kg=pl_unit_kg, pallet_capacity=pallet_capacity)
    payload = {"parsed_order": parsed, "cartons": cartons, "ean_rows": ean_rows, "warnings": base_payload.get("warnings", []), "identification_code": identification_code, "pallets": pallet_result.get("pallets", 0), "pallet_capacity": pallet_capacity}
    with open(json_path, "w", encoding="utf-8") as f:
        json.dump(payload, f, ensure_ascii=False, indent=2)
    make_job_zip(zip_path, {"generated_packing_list.xlsx": xlsx_path, "generated_labels.xlsx": labels_xlsx_path, "generated_pallet_labels.xlsx": pallet_labels_xlsx_path, "parsed_data.json": json_path})
    return {"xlsx": xlsx_path, "labels_xlsx": labels_xlsx_path, "pallet_labels_xlsx": pallet_labels_xlsx_path, "json": json_path, "zip": zip_path, "individual": individual, "payload": payload}


def build_from_edited_cartons(base_payload: Dict[str, Any], work_dir: str, cartons: List[Dict[str, Any]], tare_kg: float = 1.2, unit_kg: float = 1.3, pl_unit_kg: float = 0.75, shipment_date: str = "", identification_code: str = "", pallet_capacity: int = 16) -> Dict[str, Any]:
    """Rebuild every output straight from a manually edited packing (carton) list.

    Unlike build_from_edited_payload, this never re-derives cartons from the order
    lines. The cartons are taken exactly as edited in the packing editor - box
    numbers, quantities, EAN, batch, color and reference can be hand-edited, and
    boxes can be manually added, duplicated/split or deleted - so a direct edit to
    the packing list is never silently overwritten by the automatic line-based
    cartonization.
    """
    # Two shapes of base_payload reach this function: the full order-build payload
    # (has "parsed_order") and the flatter "existing packing list" payload from
    # build_labels_from_existing_pl (metadata/cartons at the top level, no PDF-derived
    # order lines). Normalize both to the same working shape so the packing editor
    # can rebuild from either starting point.
    if "parsed_order" in base_payload:
        parsed = dict(base_payload.get("parsed_order", {}))
    else:
        parsed = {"metadata": dict(base_payload.get("metadata", {})), "order_lines": base_payload.get("order_lines", [])}
    metadata = dict(parsed.get("metadata", {}))
    if shipment_date:
        metadata["shipment_date"] = shipment_date
    if identification_code:
        metadata["identification_code"] = identification_code
    metadata["pl_unit_kg"] = pl_unit_kg
    parsed["metadata"] = metadata

    clean_cartons: List[Dict[str, Any]] = []
    for idx, raw in enumerate(cartons or []):
        c = dict(raw)
        sizes = {normalize_size(k): int(v or 0) for k, v in (c.get("sizes") or {}).items() if is_supported_size(k) and int(v or 0) > 0}
        total = sum(sizes.values())
        if total <= 0:
            continue
        try:
            line_index = int(c.get("line_index"))
        except Exception:
            line_index = None
        try:
            box_no = int(float(c.get("box_no") or (idx + 1)))
        except Exception:
            box_no = idx + 1
        clean_cartons.append({
            "line_index": line_index,
            "batch": str(c.get("batch", "") or ""),
            "reference": str(c.get("reference") or c.get("product") or ""),
            "product": str(c.get("product") or c.get("reference") or ""),
            "color": str(c.get("color", "") or ""),
            "supplier_color": str(c.get("supplier_color", "") or ""),
            "box_no": box_no,
            "sizes": sizes,
            "total_pcs": total,
            "ean": str(c.get("ean", "") or ""),
            "missing_ean_sizes": [],
            "gross_weight": round(float(tare_kg or 0) + total * float(unit_kg or 0), 2),
            "pl_weight": round(total * float(pl_unit_kg or 0), 2),
        })

    ean_rows = base_payload.get("ean_rows", [])
    if ean_rows and any(not c["ean"] for c in clean_cartons):
        apply_ean_map_to_cartons([c for c in clean_cartons if not c["ean"]], ean_rows_to_map(ean_rows), overwrite=False)
    parsed["total_pieces"] = sum(c["total_pcs"] for c in clean_cartons)
    shipment_date_text = shipment_date or metadata.get("shipment_date", "")
    if not identification_code:
        identification_code = metadata.get("identification_code", "") or identification_from_shipment_date(parse_date(shipment_date_text or metadata.get("export_date", "")))
        parsed["metadata"]["identification_code"] = identification_code

    xlsx_path = os.path.join(work_dir, "generated_packing_list.xlsx")
    labels_xlsx_path = os.path.join(work_dir, "generated_labels.xlsx")
    pallet_labels_xlsx_path = os.path.join(work_dir, "generated_pallet_labels.xlsx")
    json_path = os.path.join(work_dir, "parsed_data.json")
    zip_path = os.path.join(work_dir, "asphalte_output.zip")

    build_packing_list_xlsx(xlsx_path, parsed, clean_cartons, ean_rows, metadata_overrides=parsed.get("metadata", {}), identification_code=identification_code, shipment_date_text=shipment_date_text, pl_unit_kg=pl_unit_kg)
    build_labels_xlsx(labels_xlsx_path, clean_cartons, parsed.get("metadata", {}), identification_code=identification_code)
    pallet_result = build_pallet_labels_xlsx(pallet_labels_xlsx_path, clean_cartons, parsed.get("metadata", {}), pallet_capacity=pallet_capacity)
    individual = build_individual_outputs(work_dir, parsed, clean_cartons, ean_rows, identification_code, shipment_date_text=shipment_date_text, pl_unit_kg=pl_unit_kg, pallet_capacity=pallet_capacity)

    payload = {
        "parsed_order": parsed,
        "cartons": clean_cartons,
        "ean_rows": ean_rows,
        "warnings": base_payload.get("warnings", []),
        "identification_code": identification_code,
        "pallets": pallet_result.get("pallets", 0),
        "pallet_capacity": pallet_capacity,
    }
    with open(json_path, "w", encoding="utf-8") as f:
        json.dump(payload, f, ensure_ascii=False, indent=2)
    make_job_zip(zip_path, {"generated_packing_list.xlsx": xlsx_path, "generated_labels.xlsx": labels_xlsx_path, "generated_pallet_labels.xlsx": pallet_labels_xlsx_path, "parsed_data.json": json_path})
    return {"xlsx": xlsx_path, "labels_xlsx": labels_xlsx_path, "pallet_labels_xlsx": pallet_labels_xlsx_path, "json": json_path, "zip": zip_path, "individual": individual, "payload": payload}


def build_from_order(order_pdf: str, work_dir: str, ean_xlsx: Optional[str] = None, template_xlsx: Optional[str] = None, manual_eans: str = "", ocr_mode: str = "auto", max_pcs: int = 15, carton_mode: str = "sequential", shipment_date: str = "", identification_code: str = "", tare_kg: float = 1.2, unit_kg: float = 1.3, pl_unit_kg: float = 0.75, pallet_capacity: int = 16) -> Dict[str, Any]:
    parsed, ocr_warnings = parse_order_pdf(order_pdf, ocr_mode=ocr_mode)
    if not parsed.get("metadata", {}).get("production_order"):
        m_file = re.search(r"\b(OF[-_ ]?\d+)\b", os.path.basename(order_pdf), re.I)
        if m_file:
            parsed["metadata"]["production_order"] = m_file.group(1).replace("_", "-").replace(" ", "-").upper()
    template_meta = extract_template_metadata(template_xlsx)
    if template_meta:
        parsed["metadata"].update({k: v for k, v in template_meta.items() if v})
    if shipment_date:
        parsed["metadata"]["shipment_date"] = shipment_date
    if identification_code:
        parsed["metadata"]["identification_code"] = identification_code
    elif template_meta.get("identification_code"):
        identification_code = template_meta["identification_code"]
    else:
        identification_code = parsed["metadata"].get("identification_code", "") or identification_from_shipment_date(parse_date(shipment_date or parsed["metadata"].get("shipment_date", "") or parsed["metadata"].get("export_date", "")))
        parsed["metadata"]["identification_code"] = identification_code

    parsed["metadata"]["pl_unit_kg"] = pl_unit_kg

    ean_parts = []
    if ean_xlsx:
        ean_parts.append(parse_eans_xlsx(ean_xlsx))
    if manual_eans:
        ean_parts.append(parse_manual_eans(manual_eans))
    ean = merge_ean_maps(*ean_parts)
    cartons = generate_cartons(parsed, ean.get("map", {}), max_pcs=max_pcs, mode=carton_mode, tare_kg=tare_kg, unit_kg=unit_kg)
    # Fill reference/product fields
    for c in cartons:
        c["reference"] = parsed["metadata"].get("product_name", c.get("reference", ""))
        c["product"] = parsed["metadata"].get("product_name", c.get("product", ""))

    all_warnings = ocr_warnings + parsed.get("warnings", []) + ean.get("warnings", [])
    for c in cartons:
        if c.get("missing_ean_sizes"):
            all_warnings.append(f"{c.get('batch')} {c.get('color')} box {c.get('box_no')}: missing EAN for {', '.join(c['missing_ean_sizes'])}")

    xlsx_path = os.path.join(work_dir, "generated_packing_list.xlsx")
    labels_xlsx_path = os.path.join(work_dir, "generated_labels.xlsx")
    pallet_labels_xlsx_path = os.path.join(work_dir, "generated_pallet_labels.xlsx")
    json_path = os.path.join(work_dir, "parsed_data.json")
    zip_path = os.path.join(work_dir, "asphalte_output.zip")
    shipment_date_text = shipment_date or parsed["metadata"].get("shipment_date", "")
    build_packing_list_xlsx(xlsx_path, parsed, cartons, ean.get("rows", []), metadata_overrides=parsed.get("metadata", {}), identification_code=identification_code, shipment_date_text=shipment_date_text, pl_unit_kg=pl_unit_kg)
    build_labels_xlsx(labels_xlsx_path, cartons, parsed.get("metadata", {}), identification_code=identification_code)
    pallet_result = build_pallet_labels_xlsx(pallet_labels_xlsx_path, cartons, parsed.get("metadata", {}), pallet_capacity=pallet_capacity)
    individual = build_individual_outputs(work_dir, parsed, cartons, ean.get("rows", []), identification_code, shipment_date_text=shipment_date_text, pl_unit_kg=pl_unit_kg, pallet_capacity=pallet_capacity)
    payload = {"parsed_order": parsed, "cartons": cartons, "ean_rows": ean.get("rows", []), "warnings": all_warnings, "identification_code": identification_code, "pallets": pallet_result.get("pallets", 0), "pallet_capacity": pallet_capacity}
    with open(json_path, "w", encoding="utf-8") as f:
        json.dump(payload, f, ensure_ascii=False, indent=2)
    make_job_zip(zip_path, {"generated_packing_list.xlsx": xlsx_path, "generated_labels.xlsx": labels_xlsx_path, "generated_pallet_labels.xlsx": pallet_labels_xlsx_path, "parsed_data.json": json_path})
    return {"xlsx": xlsx_path, "labels_xlsx": labels_xlsx_path, "pallet_labels_xlsx": pallet_labels_xlsx_path, "json": json_path, "zip": zip_path, "individual": individual if 'individual' in locals() else [], "payload": payload}


def _active_sizes_sorted(carton: Dict[str, Any]) -> List[str]:
    act = [normalize_size(s) for s, q in (carton.get("sizes") or {}).items() if int(q or 0) > 0]
    return sorted(act, key=lambda x: SIZE_ORDER.index(x) if x in SIZE_ORDER else 999)


def apply_ean_map_to_cartons(cartons: List[Dict[str, Any]], ean_map: Dict[str, str], overwrite: bool = True) -> int:
    """Fill each carton's comma-joined `ean` (aligned to its active sizes) from an EAN map.
    Existing per-size values are kept where the map has no match. Returns number of matched sizes."""
    matched = 0
    for c in cartons or []:
        sizes = _active_sizes_sorted(c)
        existing = [e.strip() for e in str(c.get("ean", "")).replace("\n", ",").split(",")]
        out: List[str] = []
        missing: List[str] = []
        for i, sz in enumerate(sizes):
            found = lookup_ean(ean_map, c.get("color", "") or c.get("supplier_color", ""), sz)
            old = existing[i] if i < len(existing) else ""
            if found and (overwrite or not old):
                out.append(found)
                matched += 1
            else:
                out.append(old)
                if not old:
                    missing.append(sz)
        c["ean"] = ",".join(out) if any(out) else ""
        c["missing_ean_sizes"] = missing
    return matched


def build_labels_from_existing_pl(pl_xlsx: str, work_dir: str, ocr_mode: str = "auto", ean_xlsx: Optional[str] = None, manual_eans: str = "", tare_kg: float = 1.2, unit_kg: float = 1.3, pl_unit_kg: float = 0.75, pallet_capacity: int = 16, shipment_date: str = "", identification_code: str = "") -> Dict[str, Any]:
    parsed = parse_existing_packing_list_xlsx(pl_xlsx)
    meta = parsed.setdefault("metadata", {})
    if shipment_date:
        meta["shipment_date"] = shipment_date
    if identification_code:
        meta["identification_code"] = identification_code
    meta["pl_unit_kg"] = pl_unit_kg
    ean_parts = []
    if ean_xlsx:
        ean_parts.append(parse_eans_xlsx(ean_xlsx))
    if manual_eans:
        ean_parts.append(parse_manual_eans(manual_eans))
    ean = merge_ean_maps(*ean_parts)
    cartons = parsed.get("cartons", [])
    warnings = list(parsed.get("warnings", [])) + list(ean.get("warnings", []))
    if ean.get("map"):
        apply_ean_map_to_cartons(cartons, ean["map"], overwrite=True)
        for c in cartons:
            if c.get("missing_ean_sizes"):
                warnings.append(f"{c.get('batch')} {c.get('color')} box {c.get('box_no')}: missing EAN for {', '.join(c['missing_ean_sizes'])}")
    parsed["warnings"] = warnings
    parsed["ean_rows"] = ean.get("rows", [])
    for c in cartons:
        tot = int(c.get("total_pcs", 0) or 0)
        if not float(c.get("gross_weight") or 0):
            c["gross_weight"] = round(float(tare_kg or 0) + tot * float(unit_kg or 0), 2)
        c["pl_weight"] = round(tot * float(pl_unit_kg or 0), 2)
    ident = meta.get("identification_code", "")
    labels_xlsx_path = os.path.join(work_dir, "labels_from_existing_packing_list.xlsx")
    pallet_labels_xlsx_path = os.path.join(work_dir, "pallet_labels_from_existing_packing_list.xlsx")
    json_path = os.path.join(work_dir, "parsed_existing_packing_list.json")
    zip_path = os.path.join(work_dir, "labels_from_existing_packing_list.zip")
    build_labels_xlsx(labels_xlsx_path, cartons, meta, identification_code=ident)
    pallet_result = build_pallet_labels_xlsx(pallet_labels_xlsx_path, cartons, meta, pallet_capacity=pallet_capacity)
    parsed["pallets"] = pallet_result.get("pallets", 0)
    parsed["pallet_capacity"] = pallet_capacity
    with open(json_path, "w", encoding="utf-8") as f:
        json.dump(parsed, f, ensure_ascii=False, indent=2)
    make_job_zip(zip_path, {"labels_from_existing_packing_list.xlsx": labels_xlsx_path, "pallet_labels_from_existing_packing_list.xlsx": pallet_labels_xlsx_path, "parsed_existing_packing_list.json": json_path})
    return {"labels_xlsx": labels_xlsx_path, "pallet_labels_xlsx": pallet_labels_xlsx_path, "json": json_path, "zip": zip_path, "payload": parsed}

# ------------------------------
# v13 formatting overrides
# ------------------------------
# These functions intentionally override the earlier v12 builders.  The goal is to
# match the uploaded Asphalte Excel/PDF samples more closely: combined PKL workbook
# with SUMMARY + PACKING LIST + all order-line sheets, and carton labels as one
# clean A4-landscape Excel sheet per carton.

def _v13_unique_sheet_title(wb, base: str) -> str:
    base = safe_sheet_title(base)[:31] or "Sheet"
    if base not in wb.sheetnames:
        return base
    stem = base[:26]
    i = 2
    while True:
        name = f"{stem} {i}"
        if name not in wb.sheetnames:
            return name
        i += 1


def _v13_side(style="thin", color="000000"):
    return Side(style=style, color=color)


def _v13_border(style="thin", color="000000"):
    s = _v13_side(style, color)
    return Border(left=s, right=s, top=s, bottom=s)


def _v13_set(cell, value="", size=10, bold=False, italic=False, fill=None, align="left", color="000000", wrap=True):
    cell.value = value
    cell.font = Font(name="Aptos Narrow", size=size, bold=bold, italic=italic, color=color)
    cell.alignment = Alignment(horizontal=align, vertical="center", wrap_text=wrap)
    if fill is not None:
        cell.fill = fill
    return cell


def _v13_merge(ws, r1, c1, r2, c2):
    try:
        ws.merge_cells(start_row=r1, start_column=c1, end_row=r2, end_column=c2)
    except Exception:
        pass
    return ws.cell(r1, c1)


def _v13_apply_border(ws, r1, c1, r2, c2, outside_style="medium", inside_style="thin", color="000000"):
    out = _v13_side(outside_style, color)
    inside = _v13_side(inside_style, color)
    for rr in range(r1, r2 + 1):
        for cc in range(c1, c2 + 1):
            ws.cell(rr, cc).border = Border(
                left=out if cc == c1 else inside,
                right=out if cc == c2 else inside,
                top=out if rr == r1 else inside,
                bottom=out if rr == r2 else inside,
            )


def _v13_clean_lines(value: str, default: str = "") -> List[str]:
    return [compact_spaces(x) for x in str(value or default or "").splitlines() if compact_spaces(x)]


def _v13_group_cartons(cartons: List[Dict[str, Any]]) -> Dict[Any, List[Dict[str, Any]]]:
    groups: Dict[Any, List[Dict[str, Any]]] = {}
    for c in cartons or []:
        if c.get("line_index") is not None:
            key = ("line", c.get("line_index"), c.get("batch", ""), c.get("color", ""))
        else:
            key = ("sheet", c.get("sheet") or c.get("batch", ""), c.get("color", ""))
        groups.setdefault(key, []).append(c)
    return groups


def _v13_carton_ean_by_size(carton: Dict[str, Any], size: str) -> str:
    """Map the comma-separated carton EAN list back to the active sizes order."""
    size = normalize_size(size)
    active = [(normalize_size(s), int(q or 0)) for s, q in (carton.get("sizes") or {}).items() if int(q or 0) > 0]
    active = sorted(active, key=lambda x: SIZE_ORDER.index(x[0]) if x[0] in SIZE_ORDER else 999)
    eans = [e.strip() for e in str(carton.get("ean", "")).replace("\n", ",").split(",") if e.strip()]
    for idx, (sz, _q) in enumerate(active):
        if sz == size and idx < len(eans):
            return eans[idx]
    return ""


def _v13_ean_for_size(ean_rows: List[Dict[str, str]], color: str, size: str, carton: Optional[Dict[str, Any]] = None) -> str:
    ean_map = ean_rows_to_map(ean_rows or [])
    ean = lookup_ean(ean_map, color, normalize_size(size))
    if ean:
        return ean
    if carton:
        return _v13_carton_ean_by_size(carton, size)
    return ""


def _v13_prepare_sheet(ws, orientation="landscape", fit_height=1, margins=(0.18, 0.18, 0.22, 0.18)):
    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = orientation
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = fit_height
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins = PageMargins(left=margins[0], right=margins[1], top=margins[2], bottom=margins[3], header=0.1, footer=0.1)


def _write_packing_sheet_base(ws, metadata: Dict[str, Any], rows: List[Dict[str, Any]], ean_rows: List[Dict[str, str]], identification_code: str, shipment_date: Optional[datetime], pl_unit_kg: float, pallet_capacity: int = 16):
    """Write one original-like Asphalte packing-list sheet for one batch/color.

    This is the base layer that the layout refinements below (
    _write_packing_sheet_extended, then the final _v13_write_packing_sheet)
    build on top of.
    """
    grey = PatternFill("solid", fgColor="F2F2F2")
    dark = PatternFill("solid", fgColor="7F7F7F")
    pale_blue = PatternFill("solid", fgColor="DDEBF7")
    red_fill = PatternFill("solid", fgColor="F4CCCC")
    black_fill = PatternFill("solid", fgColor="000000")
    white = PatternFill("solid", fgColor="FFFFFF")
    thin_border = _v13_border("thin")

    _v13_prepare_sheet(ws, "landscape", 1, margins=(0.18, 0.18, 0.16, 0.16))
    # Match the original workbook proportions: a wide A:G left side and J:M destination/right side.
    widths = {
        "A": 27, "B": 12, "C": 9, "D": 36, "E": 10, "F": 8, "G": 8,
        "H": 10, "I": 10, "J": 10, "K": 10, "L": 10, "M": 10,
    }
    for col, width in widths.items():
        ws.column_dimensions[col].width = width
    for rr in range(1, 52):
        ws.row_dimensions[rr].height = 15
        for cc in range(1, 14):
            ws.cell(rr, cc).fill = white
            ws.cell(rr, cc).font = Font(name="Aptos Narrow", size=10)
            ws.cell(rr, cc).alignment = Alignment(vertical="center", wrap_text=True)

    # Header ID row.
    _v13_merge(ws, 1, 1, 1, 13)
    _v13_set(ws["A1"], f"ASPHALTE - identification code is: {identification_code or ''}", size=22, bold=True, align="center")
    ws.row_dimensions[1].height = 25

    shipper = _v13_clean_lines(metadata.get("shipper"), DEFAULT_SHIPPER)
    receiver = _v13_clean_lines(metadata.get("receiver"), DEFAULT_RECEIVER)
    batch = rows[0].get("batch", "") if rows else ""
    color = rows[0].get("color", "") if rows else ""
    ref = rows[0].get("reference") or metadata.get("product_name", "") if rows else metadata.get("product_name", "")
    production_order = metadata.get("production_order", "") or metadata.get("order", "")
    order_label = production_order.replace("OF-", "-", 1) if production_order.upper().startswith("OF-") else production_order
    if order_label and not order_label.startswith("-") and re.search(r"\d", order_label):
        order_label = "-" + re.sub(r"[^0-9]", "", order_label) if not order_label.startswith("-") else order_label
    of_detail = compact_spaces(f"{batch} {ref}  {color}")
    total_cartons = len(rows)
    total_pcs = sum(int(c.get("total_pcs", 0) or 0) for c in rows)
    total_weight = round(sum(float(c.get("pl_weight", int(c.get("total_pcs", 0) or 0) * pl_unit_kg) or 0) for c in rows), 2)
    pallets = math.ceil(total_cartons / max(1, int(pallet_capacity or 16))) if total_cartons else 0

    # Shipper block A3:G8.
    _v13_merge(ws, 3, 1, 3, 7); _v13_set(ws["A3"], "EXPEDITEUR/ SHIPPER", 10, True, fill=grey)
    for rr in range(4, 9):
        for cc in range(1, 8): ws.cell(rr, cc).border = thin_border
    _v13_set(ws["A5"], "Nom/ Name", 10, True); _v13_set(ws["A6"], "Adresse / Address", 10, True)
    for idx, line in enumerate(shipper[:4], 5):
        _v13_set(ws.cell(idx, 2), line, 10 if idx else 12, bold=(idx == 5))
    _v13_apply_border(ws, 3, 1, 8, 7, "medium")

    # Destination J3:M8, with values starting K5 like the original.
    _v13_merge(ws, 3, 10, 3, 13); _v13_set(ws["J3"], "DESTINATAIRE", 10, True, fill=grey)
    for rr in range(4, 9):
        for cc in range(10, 14): ws.cell(rr, cc).border = thin_border
    _v13_set(ws["J5"], "Nom", 10, True); _v13_set(ws["J6"], "Adresse", 10, True)
    for idx, line in enumerate(receiver[:4], 5):
        _v13_merge(ws, idx, 11, idx, 13)
        _v13_set(ws.cell(idx, 11), line, 10, bold=(idx == 5))
    _v13_apply_border(ws, 3, 10, 8, 13, "medium")

    # Forwarder A10:G15.
    _v13_merge(ws, 10, 1, 10, 7); _v13_set(ws["A10"], "TRANSPORTEUR / FORWARDER", 10, True, fill=grey)
    _v13_set(ws["A12"], "Nom/ Name", 10, True); _v13_set(ws["B12"], metadata.get("forwarder_name", "BULSTAR OUTDOOR Ltd"), 10, True)
    _v13_set(ws["A13"], "Plate N", 10); _v13_set(ws["B13"], metadata.get("plate", "A7269PM"), 10)
    _v13_set(ws["A14"], "Contact", 10); _v13_set(ws["B14"], metadata.get("contact", "Miroslav Nikolov +359 888782258"), 10)
    _v13_apply_border(ws, 10, 1, 15, 7, "medium")

    # Right date/totals block H10:M20.
    _v13_merge(ws, 10, 8, 10, 13); _v13_set(ws["H10"], "DATE D'EXPEDITION/ SHIPMENT DATE", 10, True, fill=grey)
    ship_text = shipment_date.strftime("%d.%m.%Y") if shipment_date else ""
    _v13_merge(ws, 11, 8, 11, 13); _v13_set(ws["H11"], ship_text, 10, False, align="right")
    _v13_merge(ws, 12, 8, 12, 13); _v13_set(ws["H12"], "DATE DE LIVRAISON ESTIMÉE/ DELIVERY ESTIMATED DATE", 10, True, fill=grey)
    _v13_merge(ws, 15, 8, 15, 13); _v13_set(ws["H15"], "NB COLIS TOTAL/ TOTAL NUMBER OF PARCEL", 10, True, fill=grey)
    _v13_merge(ws, 16, 8, 16, 13); _v13_set(ws["H16"], total_cartons, 10, False, align="right")
    _v13_merge(ws, 17, 8, 17, 13); _v13_set(ws["H17"], "NB PALETTE TOTAL/ TOTAL NUMBER OF PALLET", 10, True, fill=grey)
    _v13_merge(ws, 18, 8, 18, 13); _v13_set(ws["H18"], pallets or "", 10, False, align="right")
    _v13_merge(ws, 19, 8, 19, 13); _v13_set(ws["H19"], "TOTAL POIDS KG", 10, True, fill=grey)
    _v13_merge(ws, 20, 8, 20, 13); _v13_set(ws["H20"], total_weight, 10, False, align="right")
    _v13_apply_border(ws, 10, 8, 20, 13, "medium")

    # Identification reception + order line block.
    _v13_merge(ws, 16, 1, 16, 6); _v13_set(ws["A16"], "N° IDENTIFICATION RECEPTION/ RECEPTION IDENTIFICATION NUMBER", 10, False, fill=grey)
    _v13_set(ws["G16"], "XXXX", 10, True, fill=red_fill, align="center", color="FF0000")
    _v13_apply_border(ws, 16, 1, 16, 7, "medium")
    _v13_set(ws["A18"], f"N°OF{order_label}", 10, True, fill=grey)
    _v13_merge(ws, 18, 2, 18, 7); _v13_set(ws["B18"], of_detail, 10, True, italic=True, color="FF0000")
    _v13_apply_border(ws, 18, 1, 20, 7, "medium")

    # Main carton table.
    table_top = 23
    sizes = ["XS", "S", "M", "L", "XL", "XXL", "3XL"]
    _v13_merge(ws, table_top, 1, table_top, 5)
    for cc in range(1, 14):
        ws.cell(table_top, cc).fill = dark if cc <= 5 else white
        ws.cell(table_top, cc).border = thin_border
        ws.cell(table_top, cc).alignment = Alignment(horizontal="center", vertical="center")
    for i, sz in enumerate(sizes, 7):
        _v13_set(ws.cell(table_top, i), sz, 10, True, align="center")
    # column F remains gap/pcs divider like original; sizes start G:M.
    headers = {1:"Référence", 2:"Coloris", 3:"N° de\ncolis", 4:"EAN", 5:"Nbre de\nPcs/Carton"}
    for cc, txt in headers.items():
        _v13_set(ws.cell(table_top + 1, cc), txt, 10, True, align="center")
    _v13_merge(ws, table_top + 1, 7, table_top + 1, 13); _v13_set(ws.cell(table_top + 1, 7), "Quantité / Taille", 10, True, align="center")
    for cc in range(1, 14):
        ws.cell(table_top + 1, cc).border = thin_border
    start = table_top + 2
    for ridx, carton in enumerate(rows, start):
        _v13_set(ws.cell(ridx, 1), carton.get("reference") or ref, 10, False)
        _v13_set(ws.cell(ridx, 2), color, 10, False, align="center")
        _v13_set(ws.cell(ridx, 3), carton.get("box_no", ridx - start + 1), 10, False, align="center")
        _v13_set(ws.cell(ridx, 4), carton.get("ean", ""), 9, False, fill=pale_blue, align="center")
        _v13_set(ws.cell(ridx, 5), int(carton.get("total_pcs", 0) or 0), 10, False, align="right")
        for sidx, size in enumerate(sizes, 7):
            q = int((carton.get("sizes") or {}).get(size, 0) or 0)
            _v13_set(ws.cell(ridx, sidx), q if q else "", 10, False, align="center")
        for cc in range(1, 14):
            ws.cell(ridx, cc).border = thin_border
        ws.row_dimensions[ridx].height = 20 if len(str(carton.get("ean", ""))) < 55 else 34
    total_row = start + len(rows)
    _v13_merge(ws, total_row, 1, total_row, 4); _v13_set(ws.cell(total_row, 1), "TOTAL", 10, True, fill=grey, align="center")
    _v13_set(ws.cell(total_row, 5), total_pcs, 10, True, fill=dark, align="right")
    ws.cell(total_row, 5).font = Font(name="Aptos Narrow", size=10, bold=True, color="FFFFFF")
    for sidx, size in enumerate(sizes, 7):
        tq = sum(int((c.get("sizes") or {}).get(size, 0) or 0) for c in rows)
        _v13_set(ws.cell(total_row, sidx), tq if tq else "", 10, True, fill=dark, align="center")
        ws.cell(total_row, sidx).font = Font(name="Aptos Narrow", size=10, bold=True, color="FFFFFF")
    for cc in range(1, 14):
        ws.cell(total_row, cc).border = thin_border
    _v13_apply_border(ws, table_top, 1, total_row, 13, "medium")

    # EAN summary below the main table.
    erow = total_row + 3
    _v13_set(ws.cell(erow, 1), "EAN", 10, True); _v13_set(ws.cell(erow, 2), "QUANTITIES", 10, True)
    ws.cell(erow, 1).border = thin_border; ws.cell(erow, 2).border = thin_border
    erow += 1
    for size in sizes:
        tq = sum(int((c.get("sizes") or {}).get(size, 0) or 0) for c in rows)
        if tq <= 0:
            continue
        ean = _v13_ean_for_size(ean_rows, color, size, rows[0] if rows else None)
        _v13_set(ws.cell(erow, 1), ean, 10, False, align="center")
        _v13_set(ws.cell(erow, 2), tq, 10, False, align="right")
        ws.cell(erow, 1).border = thin_border; ws.cell(erow, 2).border = thin_border
        erow += 1
    ws.print_area = f"A1:M{max(total_row + 10, 37)}"


def build_packing_list_xlsx(output_path: str, parsed_order: Dict[str, Any], cartons: List[Dict[str, Any]], ean_rows: List[Dict[str, str]], metadata_overrides: Optional[Dict[str, Any]] = None, identification_code: str = "", shipment_date_text: str = "", pl_unit_kg: float = 0.75) -> Dict[str, Any]:
    metadata = dict(parsed_order.get("metadata", {}))
    if metadata_overrides:
        metadata.update({k: v for k, v in metadata_overrides.items() if v not in (None, "")})
    shipment_date = parse_date(shipment_date_text or metadata.get("shipment_date", "") or metadata.get("export_date", ""))
    if not identification_code:
        identification_code = metadata.get("identification_code", "") or identification_from_shipment_date(shipment_date)

    wb = Workbook()
    ws_summary = wb.active
    ws_summary.title = "SUMMARY"
    _v13_prepare_sheet(ws_summary, "landscape", 1)
    ws_summary.sheet_view.showGridLines = False
    for col in range(1, 10):
        ws_summary.column_dimensions[get_column_letter(col)].width = [20,18,18,18,18,18,18,18,18][col-1]
    ws_summary.merge_cells("A1:I1")
    _v13_set(ws_summary["A1"], "ASPHALTE PACKING SUMMARY", 22, True, align="center")
    summary_data = [
        ("Identification code", identification_code or ""),
        ("Production order", metadata.get("production_order", "")),
        ("Product", metadata.get("product_name", "")),
        ("Shipment date", shipment_date.strftime("%d.%m.%Y") if shipment_date else ""),
        ("Total cartons", len(cartons)),
        ("Total pieces", sum(int(c.get("total_pcs", 0) or 0) for c in cartons)),
    ]
    for ridx, (k, v) in enumerate(summary_data, 3):
        _v13_set(ws_summary.cell(ridx,1), k, 11, True, fill=PatternFill("solid", fgColor="F2F2F2"))
        _v13_set(ws_summary.cell(ridx,2), v, 11, False)
        ws_summary.cell(ridx,1).border = _v13_border("thin"); ws_summary.cell(ridx,2).border = _v13_border("thin")
    headers = ["Sheet", "Batch", "Color", "Cartons", "Pieces", "Pallets", "Weight kg"]
    for cidx, h in enumerate(headers, 1):
        _v13_set(ws_summary.cell(11, cidx), h, 11, True, fill=PatternFill("solid", fgColor="D9EAD3"), align="center")
        ws_summary.cell(11, cidx).border = _v13_border("thin")

    groups = _v13_group_cartons(cartons)
    summary_row = 12

    # PACKING LIST sheet = EAN/SKU/PL quantity summary exactly like the uploaded summary sheet.
    ws_ean = wb.create_sheet("PACKING LIST")
    _v13_prepare_sheet(ws_ean, "portrait", 1, margins=(0.2, 0.2, 0.2, 0.2))
    ws_ean.sheet_view.showGridLines = False
    for col, width in zip(range(1, 6), [18, 36, 14, 26, 15]):
        ws_ean.column_dimensions[get_column_letter(col)].width = width
    ws_ean.merge_cells("A1:E1")
    _v13_set(ws_ean["A1"], f"ASPHALTE - identification code is: {identification_code or ''}", 20, True, align="center")
    headers = ["EAN", "SKU", "PL\nQUANTITY", "PRODUCT", "BATCH\nNUMBER"]
    for cidx, h in enumerate(headers, 1):
        _v13_set(ws_ean.cell(3, cidx), h, 12, True, fill=PatternFill("solid", fgColor="F2F2F2"), align="center")
        ws_ean.cell(3, cidx).border = _v13_border("thin")
    ws_ean.row_dimensions[3].height = 31

    # Build a summary per color/size/batch from cartons.
    summary_rows: Dict[Tuple[str, str, str, str], int] = {}
    for c in cartons:
        batch = c.get("batch", "")
        color = c.get("color", "")
        for size, qty in (c.get("sizes") or {}).items():
            qty = int(qty or 0)
            if qty <= 0:
                continue
            ean = _v13_ean_for_size(ean_rows, color, size, c)
            key = (batch, color, normalize_size(size), ean)
            summary_rows[key] = summary_rows.get(key, 0) + qty
    row = 4
    for (batch, color, size, ean), qty in sorted(summary_rows.items(), key=lambda x: (str(x[0][0]), normalize_key(x[0][1]), size_sort_key(x[0][2]))):
        sku = ""
        for er in ean_rows or []:
            if str(er.get("ean", "")).strip() == str(ean).strip():
                sku = er.get("sku", "")
                break
        vals = [ean, sku, qty, metadata.get("product_name", ""), batch]
        for cidx, val in enumerate(vals, 1):
            _v13_set(ws_ean.cell(row, cidx), val, 11, False)
            ws_ean.cell(row, cidx).border = _v13_border("thin")
        row += 1
    ws_ean.print_area = f"A1:E{max(row, 30)}"

    for key, rows in groups.items():
        batch = rows[0].get("batch", "") if rows else ""
        color = rows[0].get("color", "") if rows else ""
        title = _v13_unique_sheet_title(wb, f"{batch} {color}".strip() or "Packing")
        ws = wb.create_sheet(title)
        _v13_write_packing_sheet(ws, metadata, rows, ean_rows, identification_code, shipment_date, pl_unit_kg, pallet_capacity=16)
        pieces = sum(int(c.get("total_pcs", 0) or 0) for c in rows)
        pallets = math.ceil(len(rows)/16) if rows else 0
        vals = [title, batch, color, len(rows), pieces, pallets, round(sum(float(c.get("pl_weight", int(c.get("total_pcs",0))*pl_unit_kg) or 0) for c in rows),2)]
        for cidx, val in enumerate(vals, 1):
            _v13_set(ws_summary.cell(summary_row, cidx), val, 11, False, align="center" if cidx >=4 else "left")
            ws_summary.cell(summary_row, cidx).border = _v13_border("thin")
        summary_row += 1
    ws_summary.print_area = f"A1:I{max(summary_row+1, 20)}"
    wb.save(output_path)
    return {"path": output_path, "sheets": wb.sheetnames, "identification_code": identification_code}


# ---------------------------------------------------------------------------
# v14 visual fidelity fixes
# - packing text is given proper merged space / row height
# - carton labels use only the borders visible in the approved PDF design
# - EAN-13 and identification Code128 barcodes are rendered at high DPI,
#   aspect-ratio preserved, and cell-anchored instead of stretched
# ---------------------------------------------------------------------------


def _v14_outline_border(ws, r1: int, c1: int, r2: int, c2: int, style: str = "thick", color: str = "000000"):
    """Apply a border only to the perimeter. Internal cells stay border-free."""
    side = _v13_side(style, color)
    for rr in range(r1, r2 + 1):
        for cc in range(c1, c2 + 1):
            cell = ws.cell(rr, cc)
            old = cell.border
            cell.border = Border(
                left=side if cc == c1 else old.left,
                right=side if cc == c2 else old.right,
                top=side if rr == r1 else old.top,
                bottom=side if rr == r2 else old.bottom,
            )


def _v14_line(ws, row: int, c1: int, c2: int, style: str = "thick", color: str = "000000"):
    side = _v13_side(style, color)
    for cc in range(c1, c2 + 1):
        cell = ws.cell(row, cc)
        old = cell.border
        cell.border = Border(left=old.left, right=old.right, top=old.top, bottom=side)


def _v14_right_line(ws, c: int, r1: int, r2: int, style: str = "thin", color: str = "D4DAE2"):
    side = _v13_side(style, color)
    for rr in range(r1, r2 + 1):
        cell = ws.cell(rr, c)
        old = cell.border
        cell.border = Border(left=old.left, right=side, top=old.top, bottom=old.bottom)


def _v14_safe_merge(ws, cell_range: str):
    """Merge a range even when the old version has a smaller overlapping merge."""
    from openpyxl.utils.cell import range_boundaries
    min_col, min_row, max_col, max_row = range_boundaries(cell_range)
    to_remove = []
    for merged in list(ws.merged_cells.ranges):
        if not (merged.max_col < min_col or merged.min_col > max_col or merged.max_row < min_row or merged.min_row > max_row):
            to_remove.append(str(merged))
    for merged in to_remove:
        try:
            ws.unmerge_cells(merged)
        except Exception:
            pass
    ws.merge_cells(cell_range)


def _v14_add_image_locked(ws, image_path: str, row: int, col: int, max_width_px: int, max_height_px: int,
                          x_offset_px: int = 2, y_offset_px: int = 1) -> bool:
    """Anchor an image to a cell and fit it proportionally inside a fixed box."""
    if not image_path or not os.path.exists(image_path):
        return False
    try:
        with Image.open(image_path) as im:
            iw, ih = im.size
        if iw <= 0 or ih <= 0:
            return False
        scale = min(max_width_px / iw, max_height_px / ih)
        width = max(1, int(iw * scale))
        height = max(1, int(ih * scale))
        # Center in the target pixel box.
        xoff = x_offset_px + max(0, (max_width_px - width) // 2)
        yoff = y_offset_px + max(0, (max_height_px - height) // 2)
        img = XLImage(image_path)
        marker = AnchorMarker(
            col=col - 1,
            row=row - 1,
            colOff=pixels_to_EMU(xoff),
            rowOff=pixels_to_EMU(yoff),
        )
        img.anchor = OneCellAnchor(
            _from=marker,
            ext=XDRPositiveSize2D(pixels_to_EMU(width), pixels_to_EMU(height)),
        )
        ws.add_image(img)
        return True
    except Exception:
        return False


# ReportLab raster fallback is bundled with the application and provides correct
# guard bars / human-readable digits even when python-barcode is unavailable.
def _v14_barcode_png_path(value: str, tmp_dir: str, barcode_type: str = "ean13") -> Optional[str]:
    raw = str(value or "").strip()
    if not raw:
        return None
    os.makedirs(tmp_dir, exist_ok=True)
    kind = str(barcode_type or "ean13").lower()
    try:
        from reportlab.graphics.barcode import createBarcodeDrawing
        from reportlab.graphics import renderPM
        if kind == "code128":
            safe = re.sub(r"[^A-Za-z0-9._-]", "_", raw)[:48]
            path = os.path.join(tmp_dir, f"code128_{safe}_{uuid.uuid4().hex[:8]}.png")
            drawing = createBarcodeDrawing(
                "Code128", value=raw, barHeight=34, barWidth=0.78,
                humanReadable=True, quiet=True,
            )
            renderPM.drawToFile(drawing, path, fmt="PNG", dpi=300)
            return path

        digits = re.sub(r"\D", "", raw)
        if len(digits) == 12:
            digits += _ean13_check_digit(digits)
        if len(digits) != 13:
            return None
        path = os.path.join(tmp_dir, f"ean13_{digits}_{uuid.uuid4().hex[:8]}.png")
        drawing = createBarcodeDrawing(
            "EAN13", value=digits, barHeight=35, barWidth=0.82,
            humanReadable=True, quiet=True,
        )
        renderPM.drawToFile(drawing, path, fmt="PNG", dpi=300)
        return path
    except Exception:
        # Last-resort EAN fallback; identification code remains visible as text
        # if a machine has a broken ReportLab raster backend.
        if kind != "code128":
            return _barcode_png_path(raw, tmp_dir)
        return None

# ---------------------------------------------------------------------------
# v15 landscape label rebuild
# - true wide A4 landscape sheet proportions
# - one label per worksheet
# - clean PDF-like hierarchy and restrained borders
# - shared, robust Code128 / EAN-13 rendering anchored to cell ranges
# ---------------------------------------------------------------------------

def _v15_barcode_png_path(value: str, tmp_dir: str, barcode_type: str = "ean13") -> Optional[str]:
    raw = str(value or "").strip()
    if not raw:
        return None
    os.makedirs(tmp_dir, exist_ok=True)
    kind = str(barcode_type or "ean13").lower()

    # Preferred renderer: python-barcode + Pillow.
    try:
        if kind == "code128" and BarcodeCode128 is not None and BarcodeImageWriter is not None:
            safe = re.sub(r"[^A-Za-z0-9._-]", "_", raw)[:50]
            base = os.path.join(tmp_dir, f"id_{safe}_{uuid.uuid4().hex[:8]}")
            return BarcodeCode128(raw, writer=BarcodeImageWriter()).save(base, options={
                "module_width": 0.26,
                "module_height": 10.0,
                "quiet_zone": 1.2,
                "font_size": 7,
                "text_distance": 1.0,
                "write_text": True,
                "dpi": 300,
                "background": "white",
                "foreground": "black",
            })
        if kind != "code128":
            digits = re.sub(r"\D", "", raw)
            if len(digits) == 13:
                digits = digits[:12]
            if len(digits) == 12 and BarcodeEAN13 is not None and BarcodeImageWriter is not None:
                base = os.path.join(tmp_dir, f"ean_{digits}_{uuid.uuid4().hex[:8]}")
                return BarcodeEAN13(digits, writer=BarcodeImageWriter()).save(base, options={
                    "module_width": 0.25,
                    "module_height": 11.0,
                    "quiet_zone": 1.0,
                    "font_size": 7,
                    "text_distance": 0.9,
                    "write_text": True,
                    "dpi": 300,
                    "background": "white",
                    "foreground": "black",
                })
    except Exception:
        pass

    # Bundled ReportLab fallback, including Code128 identification barcode.
    try:
        from reportlab.graphics.barcode import createBarcodeDrawing
        from reportlab.graphics import renderPM
        if kind == "code128":
            safe = re.sub(r"[^A-Za-z0-9._-]", "_", raw)[:50]
            path = os.path.join(tmp_dir, f"id_rl_{safe}_{uuid.uuid4().hex[:8]}.png")
            drawing = createBarcodeDrawing(
                "Code128", value=raw, barHeight=34, barWidth=0.74,
                humanReadable=True, quiet=True,
            )
            renderPM.drawToFile(drawing, path, fmt="PNG", dpi=300)
            return path
        digits = re.sub(r"\D", "", raw)
        if len(digits) == 12:
            digits += _ean13_check_digit(digits)
        if len(digits) == 13:
            path = os.path.join(tmp_dir, f"ean_rl_{digits}_{uuid.uuid4().hex[:8]}.png")
            drawing = createBarcodeDrawing(
                "EAN13", value=digits, barHeight=35, barWidth=0.72,
                humanReadable=True, quiet=True,
            )
            renderPM.drawToFile(drawing, path, fmt="PNG", dpi=300)
            return path
    except Exception:
        pass

    if kind == "ean13":
        return _barcode_png_path(raw, tmp_dir)
    return None


def _v15_pad_image(image_path: str, target_width: int, target_height: int, out_path: str, padding: int = 4) -> Optional[str]:
    """Fit without distortion, then pad to the exact target aspect ratio."""
    if Image is None or not image_path or not os.path.exists(image_path):
        return image_path if image_path and os.path.exists(image_path) else None
    try:
        with Image.open(image_path) as src:
            src = src.convert("RGB")
            # Crop excess white canvas while retaining a small quiet zone.
            inv = Image.eval(src.convert("L"), lambda p: 255 - p)
            bbox = inv.getbbox()
            if bbox:
                l, t, r, b = bbox
                l = max(0, l - padding)
                t = max(0, t - padding)
                r = min(src.width, r + padding)
                b = min(src.height, b + padding)
                src = src.crop((l, t, r, b))
            scale = min((target_width - 2 * padding) / max(1, src.width), (target_height - 2 * padding) / max(1, src.height))
            nw = max(1, int(src.width * scale))
            nh = max(1, int(src.height * scale))
            src = src.resize((nw, nh), Image.Resampling.LANCZOS)
            canvas = Image.new("RGB", (target_width, target_height), "white")
            canvas.paste(src, ((target_width - nw) // 2, (target_height - nh) // 2))
            canvas.save(out_path, "PNG", dpi=(300, 300))
        return out_path
    except Exception:
        return image_path


def _v15_add_image_in_cells(ws, image_path: str, start_row: int, start_col: int, end_row: int, end_col: int,
                            target_width: int, target_height: int, tmp_dir: str, key: str) -> bool:
    """Use a two-cell anchor so the barcode moves and resizes with its cell block."""
    if not image_path or not os.path.exists(image_path):
        return False
    try:
        from openpyxl.drawing.spreadsheet_drawing import TwoCellAnchor
        padded = os.path.join(tmp_dir, f"placed_{key}_{uuid.uuid4().hex[:8]}.png")
        image_path = _v15_pad_image(image_path, target_width, target_height, padded, padding=3) or image_path
        img = XLImage(image_path)
        img.anchor = TwoCellAnchor(
            editAs="twoCell",
            _from=AnchorMarker(col=start_col - 1, row=start_row - 1, colOff=pixels_to_EMU(2), rowOff=pixels_to_EMU(1)),
            to=AnchorMarker(col=end_col, row=end_row, colOff=0, rowOff=0),
        )
        ws.add_image(img)
        return True
    except Exception:
        return _v14_add_image_locked(ws, image_path, start_row, start_col, target_width, target_height, 2, 1)


def _v15_set(cell, value="", size=9, bold=False, color="000000", align="left", fill=None, wrap=True, vertical="center"):
    cell.value = value
    cell.font = Font(name="Arial", size=size, bold=bold, color=color)
    cell.alignment = Alignment(horizontal=align, vertical=vertical, wrap_text=wrap, shrink_to_fit=False)
    if fill is not None:
        cell.fill = fill
    return cell


def _v15_merge_set(ws, cell_range: str, value="", **kwargs):
    _v14_safe_merge(ws, cell_range)
    start = cell_range.split(":")[0]
    return _v15_set(ws[start], value, **kwargs)


# ---------------------------------------------------------------------------
# v16 polishing
# - cleaner packing workbook merged value areas
# - barcode rendering switched to ReportLab-first output for crisp human-readable text
# - label header gets the heavy bottom rule from the approved PDF
# - bottom barcode section tightened so the lower border sits closer to the barcodes
# ---------------------------------------------------------------------------

def _v16_barcode_png_path(value: str, tmp_dir: str, barcode_type: str = "ean13") -> Optional[str]:
    raw = str(value or "").strip()
    if not raw:
        return None
    os.makedirs(tmp_dir, exist_ok=True)
    kind = str(barcode_type or 'ean13').lower()
    # Prefer ReportLab because it produces stable human-readable digits for both
    # EAN-13 and Code128 in the generated Excel workbooks.
    try:
        from reportlab.graphics.barcode import createBarcodeDrawing
        from reportlab.graphics import renderPM
        safe = re.sub(r"[^A-Za-z0-9._-]", "_", raw)[:50]
        if kind == 'code128':
            path = os.path.join(tmp_dir, f"id_v16_{safe}_{uuid.uuid4().hex[:8]}.png")
            drawing = createBarcodeDrawing(
                'Code128', value=raw, barHeight=42, barWidth=1.05,
                humanReadable=True, quiet=True,
            )
            renderPM.drawToFile(drawing, path, fmt='PNG', dpi=600)
            return path
        digits = re.sub(r"\D", "", raw)
        if len(digits) == 12:
            digits += _ean13_check_digit(digits)
        elif len(digits) == 13:
            digits = digits
        else:
            digits = ''
        if digits:
            path = os.path.join(tmp_dir, f"ean_v16_{digits}_{uuid.uuid4().hex[:8]}.png")
            drawing = createBarcodeDrawing(
                'EAN13', value=digits, barHeight=46, barWidth=0.95,
                humanReadable=True, quiet=True,
            )
            renderPM.drawToFile(drawing, path, fmt='PNG', dpi=600)
            return path
    except Exception:
        pass
    return _v15_barcode_png_path(raw, tmp_dir, barcode_type)


def _v16_pad_image(image_path: str, target_width: int, target_height: int, out_path: str, padding: int = 6) -> Optional[str]:
    """Fit source image into a clean white canvas without aggressive cropping.

    The previous crop-first approach could push the human-readable barcode digits
    too close to the bars when Excel scaled the image. v16 preserves more of the
    source whitespace so the printed result matches the approved PDF more closely.
    """
    if Image is None or not image_path or not os.path.exists(image_path):
        return image_path if image_path and os.path.exists(image_path) else None
    try:
        with Image.open(image_path) as src:
            src = src.convert('RGB')
            inner_w = max(1, target_width - 2 * padding)
            inner_h = max(1, target_height - 2 * padding)
            scale = min(inner_w / max(1, src.width), inner_h / max(1, src.height))
            nw = max(1, int(src.width * scale))
            nh = max(1, int(src.height * scale))
            src = src.resize((nw, nh), Image.Resampling.LANCZOS)
            canvas = Image.new('RGB', (target_width, target_height), 'white')
            canvas.paste(src, ((target_width - nw) // 2, (target_height - nh) // 2))
            canvas.save(out_path, 'PNG', dpi=(600, 600))
        return out_path
    except Exception:
        return image_path


def _v16_add_image_in_cells(ws, image_path: str, start_row: int, start_col: int, end_row: int, end_col: int,
                            target_width: int, target_height: int, tmp_dir: str, key: str) -> bool:
    if not image_path or not os.path.exists(image_path):
        return False
    try:
        from openpyxl.drawing.spreadsheet_drawing import TwoCellAnchor
        padded = os.path.join(tmp_dir, f"placed_v16_{key}_{uuid.uuid4().hex[:8]}.png")
        image_path = _v16_pad_image(image_path, target_width, target_height, padded, padding=8) or image_path
        img = XLImage(image_path)
        img.anchor = TwoCellAnchor(
            editAs='twoCell',
            _from=AnchorMarker(col=start_col - 1, row=start_row - 1, colOff=pixels_to_EMU(1), rowOff=pixels_to_EMU(1)),
            to=AnchorMarker(col=end_col, row=end_row, colOff=-pixels_to_EMU(1), rowOff=-pixels_to_EMU(1)),
        )
        ws.add_image(img)
        return True
    except Exception:
        return _v14_add_image_locked(ws, image_path, start_row, start_col, target_width, target_height, 1, 1)


# Cleaner packing workbook: preserve the overall structure but merge the value
# areas so the sheet is visually closer to the approved Asphalte sample.
def _write_packing_sheet_extended(ws, metadata: Dict[str, Any], rows: List[Dict[str, Any]], ean_rows: List[Dict[str, str]], identification_code: str, shipment_date: Optional[datetime], pl_unit_kg: float, pallet_capacity: int = 16):
    """Layout refinements on top of _write_packing_sheet_base (merged value areas, taller rows for long text, wider EAN/description columns)."""
    _write_packing_sheet_base(ws, metadata, rows, ean_rows, identification_code, shipment_date, pl_unit_kg, pallet_capacity)

    # Merge shipper value area across the available width, matching the cleaner sample.
    for rng in ('B5:G5', 'B6:G6', 'B7:G7', 'B8:G8'):
        _v14_safe_merge(ws, rng)
    for rr in range(5, 9):
        ws.cell(rr, 2).alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)

    # Forwarder value area.
    for rng in ('B12:G12', 'B13:G13', 'B14:G15'):
        _v14_safe_merge(ws, rng)
    ws['B12'].alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)
    ws['B13'].alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)
    ws['B14'].alignment = Alignment(horizontal='left', vertical='top', wrap_text=True)

    # Keep the right-side destination/address block tidy.
    for rr in range(5, 9):
        ws.cell(rr, 11).alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)

    # Long bilingual headers and address rows need explicit height.
    for rr, height in {3:18, 5:20, 6:20, 7:20, 8:20, 10:20, 12:28, 15:24, 17:24, 19:22, 23:20, 24:32}.items():
        ws.row_dimensions[rr].height = max(ws.row_dimensions[rr].height or 0, height)

    # Better balance of the carton table.
    widths = {
        'A': 29, 'B': 12, 'C': 10, 'D': 42, 'E': 12, 'F': 3,
        'G': 8, 'H': 8, 'I': 8, 'J': 8, 'K': 8, 'L': 8, 'M': 8,
    }
    for col, width in widths.items():
        ws.column_dimensions[col].width = width

    total_row = None
    for rr in range(25, ws.max_row + 1):
        if str(ws.cell(rr, 1).value or '').strip().upper() == 'TOTAL':
            total_row = rr
            break
        ean_text = str(ws.cell(rr, 4).value or '')
        line_count = max(1, math.ceil(len(ean_text) / 42))
        ws.row_dimensions[rr].height = max(22, 15 * line_count + 5)
        ws.cell(rr, 1).alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)
        ws.cell(rr, 4).alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)
        ws.cell(rr, 5).alignment = Alignment(horizontal='right', vertical='center', wrap_text=True)
    if total_row:
        ws.row_dimensions[total_row].height = 22

    ws.sheet_view.zoomScale = 85
    ws.sheet_view.zoomScaleNormal = 85
    return ws


# ---------------------------------------------------------------------------
# v17 barcode-card renderer
# Uses only local, free libraries. EAN-13 is drawn module-by-module so guard bars
# and human-readable digits always match the standard. Code128 bars are rendered
# locally with ReportLab, while the readable identification number is drawn by
# Pillow to avoid Excel scaling artefacts.
# ---------------------------------------------------------------------------

def _v17_load_font(size: int, bold: bool = False):
    if Image is None:
        return None
    try:
        from PIL import ImageFont
        candidates = []
        windir = os.environ.get('WINDIR', r'C:\Windows')
        if bold:
            candidates.extend([
                os.path.join(windir, 'Fonts', 'arialbd.ttf'),
                os.path.join(windir, 'Fonts', 'calibrib.ttf'),
                '/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf',
                '/usr/share/fonts/truetype/liberation2/LiberationSans-Bold.ttf',
            ])
        else:
            candidates.extend([
                os.path.join(windir, 'Fonts', 'arial.ttf'),
                os.path.join(windir, 'Fonts', 'calibri.ttf'),
                '/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf',
                '/usr/share/fonts/truetype/liberation2/LiberationSans-Regular.ttf',
            ])
        for font_path in candidates:
            if os.path.exists(font_path):
                return ImageFont.truetype(font_path, size=size)
        return ImageFont.load_default()
    except Exception:
        return None


def _v17_text_bbox(draw, xy, text, font):
    try:
        return draw.textbbox(xy, text, font=font)
    except Exception:
        w, h = draw.textsize(text, font=font)
        return (xy[0], xy[1], xy[0] + w, xy[1] + h)


def _v17_ean13_digits(value: str) -> Optional[str]:
    digits = re.sub(r'\D', '', str(value or ''))
    if len(digits) == 12:
        digits += _ean13_check_digit(digits)
    if len(digits) != 13:
        return None
    # Keep the supplied code when the checksum is valid; repair a bad final digit.
    expected = _ean13_check_digit(digits[:12])
    if digits[-1] != expected:
        digits = digits[:12] + expected
    return digits


def _v17_ean13_modules(digits: str) -> str:
    l_codes = {
        '0':'0001101','1':'0011001','2':'0010011','3':'0111101','4':'0100011',
        '5':'0110001','6':'0101111','7':'0111011','8':'0110111','9':'0001011',
    }
    g_codes = {
        '0':'0100111','1':'0110011','2':'0011011','3':'0100001','4':'0011101',
        '5':'0111001','6':'0000101','7':'0010001','8':'0001001','9':'0010111',
    }
    r_codes = {
        '0':'1110010','1':'1100110','2':'1101100','3':'1000010','4':'1011100',
        '5':'1001110','6':'1010000','7':'1000100','8':'1001000','9':'1110100',
    }
    parity = {
        '0':'LLLLLL','1':'LLGLGG','2':'LLGGLG','3':'LLGGGL','4':'LGLLGG',
        '5':'LGGLLG','6':'LGGGLL','7':'LGLGLG','8':'LGLGGL','9':'LGGLGL',
    }
    left = ''.join((l_codes if p == 'L' else g_codes)[d] for d, p in zip(digits[1:7], parity[digits[0]]))
    right = ''.join(r_codes[d] for d in digits[7:13])
    return '101' + left + '01010' + right + '101'


def _v17_ean13_card_image(ean: str, color: str, size: str, tmp_dir: str) -> Optional[str]:
    if Image is None:
        return None
    digits = _v17_ean13_digits(ean)
    if not digits:
        return None
    try:
        from PIL import ImageDraw
        os.makedirs(tmp_dir, exist_ok=True)
        width, height = 600, 250
        img = Image.new('RGB', (width, height), 'white')
        draw = ImageDraw.Draw(img)
        # Rounded light-grey card matching the source PDF.
        draw.rounded_rectangle((3, 3, width-4, height-4), radius=18, outline='#D9DEE5', width=4, fill='white')
        header_font = _v17_load_font(25, bold=True)
        digit_font = _v17_load_font(25, bold=False)
        color_text = str(color or '').upper()
        size_text = f'T.{normalize_size(size)}'
        draw.text((22, 14), color_text, font=header_font, fill='#4B5563')
        sb = _v17_text_bbox(draw, (0,0), size_text, header_font)
        draw.text((width - 22 - (sb[2]-sb[0]), 14), size_text, font=header_font, fill='#1D4ED8')

        bits = _v17_ean13_modules(digits)
        module = 5
        quiet_modules = 10
        total_modules = len(bits) + quiet_modules * 2
        barcode_width = total_modules * module
        left = (width - barcode_width) // 2 + quiet_modules * module
        top = 58
        normal_bottom = 184
        guard_bottom = 202
        guard_positions = set(range(0,3)) | set(range(45,50)) | set(range(92,95))
        for idx, bit in enumerate(bits):
            if bit == '1':
                bottom = guard_bottom if idx in guard_positions else normal_bottom
                x1 = left + idx * module
                draw.rectangle((x1, top, x1 + module - 1, bottom), fill='black')

        # Standard EAN-13 human-readable grouping: 1 digit, 6 digits, 6 digits.
        baseline = 207
        first = digits[0]
        fb = _v17_text_bbox(draw, (0,0), first, digit_font)
        draw.text((left - 34 - (fb[2]-fb[0])//2, baseline), first, font=digit_font, fill='black')
        # Center each digit under its seven-module symbol.
        for i, d in enumerate(digits[1:7]):
            center = left + 3*module + (i*7 + 3.5)*module
            bb = _v17_text_bbox(draw, (0,0), d, digit_font)
            draw.text((center - (bb[2]-bb[0])/2, baseline), d, font=digit_font, fill='black')
        right_start = left + 3*module + 42*module + 5*module
        for i, d in enumerate(digits[7:13]):
            center = right_start + (i*7 + 3.5)*module
            bb = _v17_text_bbox(draw, (0,0), d, digit_font)
            draw.text((center - (bb[2]-bb[0])/2, baseline), d, font=digit_font, fill='black')

        safe = re.sub(r'[^A-Za-z0-9._-]', '_', f'{digits}_{color}_{size}')[:80]
        path = os.path.join(tmp_dir, f'ean_card_{safe}_{uuid.uuid4().hex[:8]}.png')
        img.save(path, 'PNG', dpi=(600,600))
        return path
    except Exception:
        return None


def _v17_ident_barcode_image(value: str, tmp_dir: str) -> Optional[str]:
    if Image is None:
        return None
    raw = str(value or '').strip()
    if not raw:
        return None
    try:
        from PIL import ImageDraw, ImageOps
        from reportlab.graphics.barcode import createBarcodeDrawing
        from reportlab.graphics import renderPM
        os.makedirs(tmp_dir, exist_ok=True)
        temp = os.path.join(tmp_dir, f'id_bars_{uuid.uuid4().hex[:8]}.png')
        drawing = createBarcodeDrawing('Code128', value=raw, barHeight=55, barWidth=1.15, humanReadable=False, quiet=True)
        renderPM.drawToFile(drawing, temp, fmt='PNG', dpi=600)
        with Image.open(temp) as bars_src:
            bars_src = bars_src.convert('RGB')
            grey = ImageOps.grayscale(bars_src)
            inv = ImageOps.invert(grey)
            bbox = inv.getbbox()
            if bbox:
                bars_src = bars_src.crop(bbox)
            width, height = 620, 190
            img = Image.new('RGB', (width, height), 'white')
            max_w, max_h = 560, 112
            scale = min(max_w/max(1,bars_src.width), max_h/max(1,bars_src.height))
            bars_src = bars_src.resize((max(1,int(bars_src.width*scale)), max(1,int(bars_src.height*scale))), Image.Resampling.NEAREST)
            x = 10
            img.paste(bars_src, (x, 10))
            draw = ImageDraw.Draw(img)
            font = _v17_load_font(29, bold=False)
            bb = _v17_text_bbox(draw, (0,0), raw, font)
            draw.text(((width-(bb[2]-bb[0]))/2, 145), raw, font=font, fill='black')
            safe = re.sub(r'[^A-Za-z0-9._-]', '_', raw)[:50]
            path = os.path.join(tmp_dir, f'ident_card_{safe}_{uuid.uuid4().hex[:8]}.png')
            img.save(path, 'PNG', dpi=(600,600))
            return path
    except Exception:
        return None


# ---------------------------------------------------------------------------
# v18 label polish
# - 2x2 EAN card layout when only a few active sizes exist
# - color/size captions align with the actual barcode width
# - identification barcode prefers GS1-128 / EAN-128 rendering when available
# - audit controls closer to the approved PDF
# ---------------------------------------------------------------------------

def _v18_ean13_card_image(ean: str, color: str, size: str, tmp_dir: str) -> Optional[str]:
    if Image is None:
        return None
    digits = _v17_ean13_digits(ean)
    if not digits:
        return None
    try:
        from PIL import ImageDraw
        os.makedirs(tmp_dir, exist_ok=True)
        width, height = 560, 205
        img = Image.new('RGB', (width, height), 'white')
        draw = ImageDraw.Draw(img)
        draw.rectangle((0, 0, width-1, height-1), fill='white')

        header_font = _v17_load_font(22, bold=True)
        digit_font = _v17_load_font(20, bold=False)
        color_text = str(color or '').upper()
        size_text = f'T.{normalize_size(size)}'

        bits = _v17_ean13_modules(digits)
        module = 4
        quiet_modules = 10
        total_modules = len(bits) + quiet_modules * 2
        barcode_width = total_modules * module
        left = (width - barcode_width) // 2 + quiet_modules * module
        top = 42
        normal_bottom = 124
        guard_bottom = 138
        guard_positions = set(range(0,3)) | set(range(45,50)) | set(range(92,95))
        for idx, bit in enumerate(bits):
            if bit == '1':
                bottom = guard_bottom if idx in guard_positions else normal_bottom
                x1 = left + idx * module
                draw.rectangle((x1, top, x1 + module - 1, bottom), fill='black')

        # captions aligned to barcode start and end
        draw.text((left, 12), color_text, font=header_font, fill='#4B5563')
        sb = _v17_text_bbox(draw, (0,0), size_text, header_font)
        draw.text((left + len(bits)*module - (sb[2]-sb[0]), 12), size_text, font=header_font, fill='#1D4ED8')

        baseline = 145
        first = digits[0]
        fb = _v17_text_bbox(draw, (0,0), first, digit_font)
        draw.text((left - 24 - (fb[2]-fb[0])//2, baseline), first, font=digit_font, fill='black')
        for i, d in enumerate(digits[1:7]):
            center = left + 3*module + (i*7 + 3.5)*module
            bb = _v17_text_bbox(draw, (0,0), d, digit_font)
            draw.text((center - (bb[2]-bb[0])/2, baseline), d, font=digit_font, fill='black')
        right_start = left + 3*module + 42*module + 5*module
        for i, d in enumerate(digits[7:13]):
            center = right_start + (i*7 + 3.5)*module
            bb = _v17_text_bbox(draw, (0,0), d, digit_font)
            draw.text((center - (bb[2]-bb[0])/2, baseline), d, font=digit_font, fill='black')

        safe = re.sub(r'[^A-Za-z0-9._-]', '_', f'{digits}_{color}_{size}')[:80]
        path = os.path.join(tmp_dir, f'ean_card_v18_{safe}_{uuid.uuid4().hex[:8]}.png')
        img.save(path, 'PNG', dpi=(600,600))
        return path
    except Exception:
        return _v17_ean13_card_image(ean, color, size, tmp_dir)


_CODE128_PATTERNS = [
    "212222","222122","222221","121223","121322","131222","122213","122312","132212","221213",
    "221312","231212","112232","122132","122231","113222","123122","123221","223211","221132",
    "221231","213212","223112","312131","311222","321122","321221","312212","322112","322211",
    "212123","212321","232121","111323","131123","131321","112313","132113","132311","211313",
    "231113","231311","112133","112331","132131","113123","113321","133121","313121","211331",
    "231131","213113","213311","213131","311123","311321","331121","312113","312311","332111",
    "314111","221411","431111","111224","111422","121124","121421","141122","141221","112214",
    "112412","122114","122411","142112","142211","241211","221114","413111","241112","134111",
    "111242","121142","121241","114212","124112","124211","411212","421112","421211","212141",
    "214121","412121","111143","111341","131141","114113","114311","411113","411311","113141",
    "114131","311141","411131","211412","211214","211232","2331112",
]


def _code128_modules(value: str) -> str:
    value = str(value)
    codes: List[int] = []
    if value.isdigit() and len(value) >= 4 and len(value) % 2 == 0:
        codes.append(105)
        for i in range(0, len(value), 2):
            codes.append(int(value[i:i + 2]))
    else:
        codes.append(104)
        for ch in value:
            o = ord(ch)
            codes.append(o - 32 if 32 <= o < 127 else 0)
    chk = codes[0]
    for i, c in enumerate(codes[1:], 1):
        chk += i * c
    codes.append(chk % 103)
    codes.append(106)
    out = ""
    for c in codes:
        bar = True
        for w in _CODE128_PATTERNS[c]:
            out += ("1" if bar else "0") * int(w)
            bar = not bar
    return out


def _pil_code128_png(value: str, tmp_dir: str, module: int = 4, height: int = 150) -> Optional[str]:
    """Pure-PIL Code128 PNG - last-resort fallback when no barcode library works."""
    if Image is None or not str(value or "").strip():
        return None
    from PIL import ImageDraw
    os.makedirs(tmp_dir, exist_ok=True)
    bits = _code128_modules(str(value).strip())
    quiet = 10
    img = Image.new("RGB", ((len(bits) + quiet * 2) * module, height), "white")
    d = ImageDraw.Draw(img)
    for i, b in enumerate(bits):
        if b == "1":
            x = (quiet + i) * module
            d.rectangle((x, 0, x + module - 1, height), fill="black")
    path = os.path.join(tmp_dir, f"c128_{uuid.uuid4().hex[:8]}.png")
    img.save(path, "PNG")
    return path


def _v18_ident_barcode_image(value: str, tmp_dir: str) -> Optional[str]:
    if Image is None:
        return None
    raw = str(value or '').strip()
    if not raw:
        return None
    try:
        from PIL import ImageDraw, ImageOps
        os.makedirs(tmp_dir, exist_ok=True)
        temp = None
        # Prefer GS1-128 / EAN-128 from python-barcode when available.
        try:
            import barcode as pybarcode
            from barcode.writer import ImageWriter
            writer = ImageWriter()
            options = {
                'module_width': 0.42,
                'module_height': 18.0,
                'quiet_zone': 2.0,
                'font_size': 0,
                'text_distance': 3,
                'write_text': False,
                'dpi': 600,
                'background': 'white',
                'foreground': 'black',
            }
            try:
                cls = pybarcode.get_barcode_class('gs1_128')
                code = cls(raw, writer=writer)
            except Exception:
                cls = pybarcode.get_barcode_class('code128')
                code = cls(raw, writer=writer)
            tmpbase = os.path.join(tmp_dir, f'gs1tmp_{uuid.uuid4().hex[:8]}')
            temp = code.save(tmpbase, options)
        except Exception:
            try:
                from reportlab.graphics.barcode import createBarcodeDrawing
                from reportlab.graphics import renderPM
                temp = os.path.join(tmp_dir, f'id_bars_v18_{uuid.uuid4().hex[:8]}.png')
                drawing = createBarcodeDrawing('Code128', value=raw, barHeight=56, barWidth=1.05, humanReadable=False, quiet=True)
                renderPM.drawToFile(drawing, temp, fmt='PNG', dpi=600)
            except Exception:
                temp = _pil_code128_png(raw, tmp_dir)

        with Image.open(temp) as bars_src:
            bars_src = bars_src.convert('RGB')
            grey = ImageOps.grayscale(bars_src)
            inv = ImageOps.invert(grey)
            bbox = inv.getbbox()
            if bbox:
                bars_src = bars_src.crop(bbox)
            width, height = 620, 176
            img = Image.new('RGB', (width, height), 'white')
            max_w, max_h = 560, 92
            scale = min(max_w/max(1,bars_src.width), max_h/max(1,bars_src.height))
            bars_src = bars_src.resize((max(1,int(bars_src.width*scale)), max(1,int(bars_src.height*scale))), Image.Resampling.NEAREST)
            x = 10
            img.paste(bars_src, (x, 10))
            draw = ImageDraw.Draw(img)
            font = _v17_load_font(28, bold=False)
            bb = _v17_text_bbox(draw, (0,0), raw, font)
            draw.text((10, 118), raw, font=font, fill='black')
            safe = re.sub(r'[^A-Za-z0-9._-]', '_', raw)[:50]
            path = os.path.join(tmp_dir, f'ident_card_v18_{safe}_{uuid.uuid4().hex[:8]}.png')
            img.save(path, 'PNG', dpi=(600,600))
            return path
    except Exception:
        return _v17_ident_barcode_image(value, tmp_dir)


def build_labels_xlsx(output_path: str, cartons: List[Dict[str, Any]], metadata: Dict[str, Any], identification_code: str = '') -> Dict[str, Any]:
    wb = Workbook()
    wb.remove(wb.active)

    white = PatternFill('solid', fgColor='FFFFFF')
    header_fill = PatternFill('solid', fgColor='F5F6F7')
    active_fill = PatternFill('solid', fgColor='FFFBEA')
    green_fill = PatternFill('solid', fgColor='D9F5E5')
    black_fill = PatternFill('solid', fgColor='000000')
    grey_blue = '58687D'
    blue = '24457A'
    red = 'D8272F'
    light = 'CBD3DE'
    tmp_dir = os.path.join(os.path.dirname(output_path), 'barcode_images_v18')
    os.makedirs(tmp_dir, exist_ok=True)

    shipper_lines = _v13_clean_lines(metadata.get('shipper'), DEFAULT_SHIPPER)
    receiver_lines = _v13_clean_lines(metadata.get('receiver'), DEFAULT_RECEIVER)
    shipper_text = '\n'.join(shipper_lines[:4])
    receiver_text = '\n'.join(receiver_lines[:4])
    ident = str(identification_code or metadata.get('identification_code', '') or '').strip()

    group_counts: Dict[Any, int] = {}
    group_positions: Dict[Any, int] = {}
    for carton in cartons or []:
        key = carton.get('line_index') if carton.get('line_index') is not None else (carton.get('batch', ''), carton.get('color', ''))
        group_counts[key] = group_counts.get(key, 0) + 1

    if not cartons:
        ws = wb.create_sheet('NO LABELS')
        ws['A1'] = 'NO CARTONS / NO LABELS'
        wb.save(output_path)
        return {'path': output_path, 'labels': 0, 'sheets': wb.sheetnames}

    for global_idx, carton in enumerate(cartons, 1):
        key = carton.get('line_index') if carton.get('line_index') is not None else (carton.get('batch', ''), carton.get('color', ''))
        group_positions[key] = group_positions.get(key, 0) + 1
        carton_position = group_positions[key]
        total_in_group = group_counts.get(key, 1)

        ws = wb.create_sheet(_v13_unique_sheet_title(wb, f"LBL {global_idx:03d} C{carton.get('box_no', carton_position)}"))
        _v13_prepare_sheet(ws, 'landscape', 1, margins=(0.18, 0.18, 0.14, 0.16))
        ws.print_options.horizontalCentered = True
        ws.print_options.verticalCentered = True
        ws.sheet_view.showGridLines = False
        ws.sheet_view.zoomScale = 75
        ws.sheet_view.zoomScaleNormal = 75

        ws.column_dimensions['A'].width = 1.2
        for col_idx in range(2, 24):
            ws.column_dimensions[get_column_letter(col_idx)].width = 6.35
        ws.column_dimensions['X'].width = 1.2
        for rr in range(1, 41):
            ws.row_dimensions[rr].height = 10.5
            for cc in range(1, 25):
                cell = ws.cell(rr, cc)
                cell.fill = white
                cell.border = Border()
                cell.font = Font(name='Arial', size=9, color='000000')
                cell.alignment = Alignment(vertical='center', wrap_text=True)

        _v15_merge_set(ws, 'B2:K3', 'E T I Q U E T T E   C O L I S', size=23, bold=True)
        _v15_merge_set(ws, 'R2:W3', 'ASPHALTE LOGISTICS', size=11, bold=True, color='6B7280', align='right')
        _v14_line(ws, 4, 2, 23, 'thick')

        _v14_outline_border(ws, 5, 2, 39, 23, 'thick')
        _v14_line(ws, 11, 2, 23, 'thick')
        _v14_line(ws, 19, 2, 23, 'thick')
        _v14_line(ws, 28, 2, 23, 'medium')

        product = str(carton.get('reference') or metadata.get('product_name', '') or '').strip()
        batch = compact_spaces(f"{carton.get('batch', '')} {product}")
        gross = carton.get('gross_weight', '')
        try:
            gross_text = f"{float(gross):.1f} KG"
        except Exception:
            gross_text = f"{gross} KG".strip()

        _v15_merge_set(ws, 'B6:H6', 'DESTINATAIRE / CONSIGNEE', size=7.5, bold=True, color=grey_blue)
        _v15_merge_set(ws, 'B7:H10', receiver_text, size=9.2, bold=True, vertical='top')
        _v15_merge_set(ws, 'I6:P6', 'N° OF (BATCH)', size=7.5, bold=True, color=grey_blue)
        _v15_merge_set(ws, 'I7:P7', batch, size=8.5, bold=True, color=blue)
        _v15_merge_set(ws, 'I8:P8', 'N° COLIS / CARTON', size=7.5, bold=True, color=grey_blue)
        _v15_merge_set(ws, 'I9:P10', f"{carton_position} of {total_in_group}", size=18, bold=True)
        _v15_merge_set(ws, 'Q6:W6', 'RÉFÉRENCE MODEL', size=7.5, bold=True, color=grey_blue)
        _v15_merge_set(ws, 'Q7:W7', product, size=9.2, bold=True)
        _v15_merge_set(ws, 'Q8:W8', 'POIDS BRUT / GROSS WEIGHT', size=7.5, bold=True, color=grey_blue)
        _v15_merge_set(ws, 'Q9:W10', gross_text, size=20, bold=True, color=red)
        for c1, c2, row in ((2,8,6),(9,16,6),(9,16,8),(17,23,6),(17,23,8)):
            for cc in range(c1, c2+1):
                old = ws.cell(row, cc).border
                ws.cell(row, cc).border = Border(left=old.left,right=old.right,top=old.top,bottom=_v13_side('thin', light))
        _v14_right_line(ws, 8, 6, 10, 'thin', light)
        _v14_right_line(ws, 16, 6, 10, 'thin', light)

        _v15_merge_set(ws, 'B14:H14', 'EXPÉDITEUR / SHIPPER', size=7.5, bold=True, color=grey_blue)
        _v15_merge_set(ws, 'B15:H18', shipper_text, size=8.6, color='5A6574', vertical='top')
        _v15_merge_set(ws, 'I14:P14', 'AUDIT / CONTRÔLE QUALITÉ', size=7.5, bold=True, color=grey_blue)
        # Choice controls styled closer to PDF.
        _v15_merge_set(ws, 'I15:I16', '☑', size=10, bold=True, align='center', fill=header_fill)
        _v15_merge_set(ws, 'J15:K16', 'CHOIX 1', size=8.5, bold=True, color='FFFFFF', align='center', fill=black_fill)
        _v15_merge_set(ws, 'L15:L16', '☐', size=10, bold=True, align='center', fill=header_fill)
        _v15_merge_set(ws, 'M15:N16', 'CHOIX 2', size=8.5, bold=True, color='7A7F87', align='center', fill=header_fill)
        _v15_merge_set(ws, 'T14:W14', 'AUDIT STATUS', size=7.5, bold=True, color=grey_blue, align='center')
        _v15_merge_set(ws, 'U15:W17', 'PASS', size=10, bold=True, color='008A4E', align='center', fill=green_fill)
        for c1, c2 in ((2,8),(9,16),(20,23)):
            for cc in range(c1, c2+1):
                old = ws.cell(14,cc).border
                ws.cell(14,cc).border = Border(left=old.left,right=old.right,top=old.top,bottom=_v13_side('thin',light))
        _v14_right_line(ws, 8, 14, 18, 'thin', light)

        _v15_merge_set(ws, 'B22:W22', 'DISTRIBUTION QUANTITÉS / SIZES BREAKDOWN', size=7.5, bold=True, color=grey_blue)
        label_sizes = sort_sizes((carton.get('sizes') or {}).keys())
        if not label_sizes:
            label_sizes = ['XS', 'S', 'M', 'L', 'XL', 'XXL', '3XL']
        # Keep the original spacious layout for alphabetic size runs. For wider
        # numeric runs (27..40 etc.) make the reference/color blocks a little
        # narrower so up to 11 size columns still fit on the same landscape label.
        if len(label_sizes) <= 7:
            ref_start, ref_end = 2, 9
            color_start, color_end = 10, 12
            size_start = 13
        else:
            ref_start, ref_end = 2, 7
            color_start, color_end = 8, 10
            size_start = 11
        size_end = size_start + len(label_sizes) - 1
        total_start = size_end + 1
        total_end = 23
        _v15_merge_set(ws, f"{get_column_letter(ref_start)}23:{get_column_letter(ref_end)}23", 'RÉFÉRENCE', size=7.5, bold=True, align='center', fill=header_fill)
        _v15_merge_set(ws, f"{get_column_letter(color_start)}23:{get_column_letter(color_end)}23", 'COLORIS', size=7.5, bold=True, align='center', fill=header_fill)
        for idx_sz, sz in enumerate(label_sizes):
            cc = size_start + idx_sz
            _v15_merge_set(ws, f"{get_column_letter(cc)}23:{get_column_letter(cc)}23", sz, size=7.0 if len(label_sizes) > 8 else 7.5, bold=True, align='center', fill=header_fill)
        _v15_merge_set(ws, f"{get_column_letter(total_start)}23:{get_column_letter(total_end)}23", 'TOTAL\nPCS', size=7.5, bold=True, align='center', fill=header_fill)
        _v15_merge_set(ws, f"{get_column_letter(ref_start)}24:{get_column_letter(ref_end)}25", product, size=9.0 if len(label_sizes) > 8 else 9.5, bold=True)
        _v15_merge_set(ws, f"{get_column_letter(color_start)}24:{get_column_letter(color_end)}25", carton.get('color', ''), size=9.0 if len(label_sizes) > 8 else 9.5, bold=True)
        for idx_sz, sz in enumerate(label_sizes):
            cc = size_start + idx_sz
            qty = int((carton.get('sizes') or {}).get(sz, 0) or 0)
            _v15_merge_set(ws, f"{get_column_letter(cc)}24:{get_column_letter(cc)}25", qty if qty else '', size=10 if len(label_sizes) > 8 else 11, bold=bool(qty), align='center', fill=active_fill if qty else None)
        _v15_merge_set(ws, f"{get_column_letter(total_start)}24:{get_column_letter(total_end)}25", int(carton.get('total_pcs',0) or 0), size=14, bold=True, align='center', fill=black_fill, color='FFFFFF')
        _v13_apply_border(ws, 23, 2, 25, 23, 'medium', 'thin')

        _v15_merge_set(ws, 'B30:H30', 'IDENTIFICATION CODE', size=7.5, bold=True, color=grey_blue)
        for cc in range(2,9):
            old = ws.cell(30,cc).border
            ws.cell(30,cc).border = Border(left=old.left,right=old.right,top=old.top,bottom=_v13_side('thin',light))
        _v15_merge_set(ws, 'B31:H31', ident, size=10.5, bold=True)
        if ident:
            raw_path = _v18_ident_barcode_image(ident, tmp_dir)
            _v14_add_image_locked(ws, raw_path, 32, 2, 215, 80, 2, 1)

        active_sizes = [(normalize_size(s), int(q or 0)) for s,q in (carton.get('sizes') or {}).items() if int(q or 0)>0]
        active_sizes.sort(key=lambda x: size_sort_key(x[0]))

        # 2x2 layout when up to 4 active sizes, matching the provided example.
        def card_slots(n):
            if n <= 2:
                return [(31, 11, 15), (31, 17, 21)][:n]
            if n <= 4:
                return [(31, 10, 14), (31, 16, 20), (35, 10, 14), (35, 16, 20)][:n]
            return [(31, 10, 12), (31, 14, 16), (31, 18, 20), (35, 11, 13), (35, 17, 19)][:n]

        for (size,_qty), (top_row, c1, c2) in zip(active_sizes[:5], card_slots(len(active_sizes[:5]))):
            ean = _v13_carton_ean_by_size(carton,size)
            if ean:
                raw_path = _v18_ean13_card_image(ean, str(carton.get('color','')), size, tmp_dir)
                _v14_add_image_locked(ws, raw_path, top_row, c1, 190 if len(active_sizes[:5])<=4 else 150, 70, 1, 1)
            else:
                _v15_merge_set(ws, f"{get_column_letter(c1)}{top_row+1}:{get_column_letter(c2)}{top_row+3}", 'EAN липсва', size=7, align='center', color='9B1C1C')

        heights = {
            2:18,3:18,4:7,5:5,
            6:13,7:15,8:13,9:17,10:15,11:6,
            12:10,13:10,14:13,15:15,16:15,17:15,18:11,19:6,
            20:11,21:11,22:13,23:18,24:18,25:18,
            26:9,27:9,28:6,29:4,30:13,31:14,
            32:18,33:18,34:18,35:18,36:18,37:18,38:10,39:6,
        }
        for rr, h in heights.items():
            ws.row_dimensions[rr].height = h

        ws.print_area = 'B2:W39'

    wb.save(output_path)
    return {'path': output_path, 'labels': len(cartons), 'sheets': wb.sheetnames}

# ---------------------------------------------------------------------------
# v19 size-system upgrade
# Supports alphabetic sizes (XS..3XL), numeric garment sizes (27, 28, 29...)
# and mixed size systems in order PDFs / EAN workbooks / manual edits.
# ---------------------------------------------------------------------------

def normalize_size(value: Any) -> str:
    if value is None:
        return ""
    # Excel may expose numeric sizes as int/float. Preserve integer numeric labels.
    if isinstance(value, bool):
        return ""
    if isinstance(value, (int, float)):
        try:
            f = float(value)
            if f.is_integer():
                return str(int(f))
            return ("%g" % f).strip()
        except Exception:
            pass
    raw = compact_spaces(str(value or ""))
    if not raw:
        return ""
    # 27.0 -> 27; otherwise preserve numeric text as a size token.
    if re.fullmatch(r"\d+(?:\.0+)?", raw):
        try:
            return str(int(float(raw)))
        except Exception:
            return raw
    s = raw.upper().replace("XXXL", "3XL")
    s = re.sub(r"[^A-Z0-9]", "", s)
    return s


def is_numeric_size(value: Any) -> bool:
    s = normalize_size(value)
    return bool(re.fullmatch(r"\d{1,3}", s))


def is_supported_size(value: Any) -> bool:
    s = normalize_size(value)
    if not s:
        return False
    if s in SIZE_ORDER:
        return True
    # Numeric garment sizes are accepted as long as they are explicitly found in
    # a size header / SIZE column / manual edit. Keep the range broad enough for
    # waist, French/Italian and other apparel systems.
    return bool(re.fullmatch(r"\d{1,3}", s))


def size_sort_key(value: Any):
    s = normalize_size(value)
    if s in SIZE_ORDER:
        return (0, SIZE_ORDER.index(s), 0, s)
    if re.fullmatch(r"\d{1,3}", s):
        return (1, int(s), 0, s)
    m = re.fullmatch(r"(\d+)([A-Z]+)", s)
    if m:
        return (2, int(m.group(1)), 0, m.group(2))
    return (3, 9999, 0, s)


def sort_sizes(sizes: Iterable[str]) -> List[str]:
    normalized = []
    seen = set()
    for value in sizes:
        s = normalize_size(value)
        if s and s not in seen:
            seen.add(s)
            normalized.append(s)
    return sorted(normalized, key=size_sort_key)


def _v19_header_size_candidate(token: str) -> bool:
    s = normalize_size(token)
    if s in SIZE_ORDER:
        return True
    # In a header context numeric values are sizes, even when the scale is small.
    return bool(re.fullmatch(r"\d{1,3}", s))


def extract_sizes_from_section(section_lines: List[str]) -> List[str]:
    """Extract size headers from the actual header row.

    Modern Asphalte PDFs flatten the table into tokens. Numeric sizes such as
    27,28,29... must only be interpreted as sizes while reading the header, not
    later quantity values. We therefore anchor after 'Supplier color' (or the
    closest known table header) and stop at Total / Export date / quantities.
    """
    tokens = [compact_spaces(x) for x in section_lines]
    low = [x.lower() for x in tokens]

    anchors = [i for i, x in enumerate(low) if x in {"supplier color", "supplier colour", "coloris fournisseur"}]
    if not anchors:
        anchors = [i for i, x in enumerate(low) if x in {"asphalte color", "asphalte colour", "coloris", "color"}]

    for anchor in anchors[:3]:
        sizes: List[str] = []
        for token in tokens[anchor + 1: anchor + 30]:
            tl = token.lower().strip()
            if tl in {"total", "quantities", "quantity", "price (€)", "amount (€)"} or tl.startswith("export date"):
                break
            if parse_date(token):
                break
            if _v19_header_size_candidate(token):
                s = normalize_size(token)
                if s not in sizes:
                    sizes.append(s)
            elif sizes:
                # Once a genuine size run started, first non-size token marks its end.
                break
        if sizes:
            return sort_sizes(sizes)

    # Legacy fallback: alphabetic sizes can still be scanned safely because
    # quantities do not resemble XS/S/M/L etc.
    sizes: List[str] = []
    for token in tokens:
        s = normalize_size(token)
        if s in SIZE_ORDER and s not in sizes:
            sizes.append(s)
        if sizes and token.lower().strip() in {"total", "quantities"}:
            break
    return sort_sizes(sizes) if sizes else ["XS", "S", "M", "L", "XL", "XXL", "3XL"]


def detect_size_from_sku(sku: str) -> str:
    sku_up = str(sku or "").upper().strip()
    for size in ["5XL", "4XL", "3XL", "XXXL", "XXL", "XXS", "XS", "XL", "L", "M", "S"]:
        if sku_up.endswith(size):
            return "3XL" if size == "XXXL" else size
    matches = SIZE_RE.findall(sku_up)
    if matches:
        sz = matches[-1].upper()
        return "3XL" if sz == "XXXL" else sz
    # Numeric sizes in the supplied SKU file are appended to the SKU, e.g.
    # PE26H010602PANTV2BLMA27. Prefer 2-3 trailing digits to avoid catching years.
    m = re.search(r"(?<!\d)(\d{2,3})$", sku_up)
    if not m:
        m = re.search(r"([A-Z])(\d{2,3})$", sku_up)
        if m:
            return normalize_size(m.group(2))
    elif m:
        return normalize_size(m.group(1))
    return ""


def parse_manual_eans(text: str) -> Dict[str, Any]:
    mappings: Dict[str, str] = {}
    rows: List[Dict[str, str]] = []
    warnings: List[str] = []
    for raw in str(text or "").splitlines():
        line = raw.strip()
        if not line or line.startswith("#"):
            continue
        ean_match = EAN_RE.search(line)
        if not ean_match:
            warnings.append(f"Manual EAN row ignored - no EAN found: {line}")
            continue
        ean = ean_match.group(1)
        left = line.split("=", 1)[0] if "=" in line else line
        # Accept color|27=EAN, color;27;EAN, color,XS,EAN etc.
        parts = [p.strip() for p in re.split(r"[|;,\t]", left) if p.strip()]
        size = ""
        color_parts: List[str] = []
        for p in parts:
            ps = normalize_size(p)
            if is_supported_size(ps):
                size = ps
            elif not EAN_RE.fullmatch(p):
                color_parts.append(p)
        if not size:
            size = detect_size_from_sku(left)
        if not size:
            warnings.append(f"Manual EAN row ignored - no size found: {line}")
            continue
        color = normalize_color(" ".join(color_parts)) if color_parts else ""
        key = f"{normalize_key(color)}|{size}" if color else f"*|{size}"
        mappings[key] = ean
        mappings.setdefault(f"*|{size}", ean)
        rows.append({"color": color, "size": size, "ean": ean, "source": "manual"})
    return {"map": mappings, "rows": rows, "warnings": warnings}


def parse_eans_xlsx(xlsx_path: str) -> Dict[str, Any]:
    wb = load_workbook(xlsx_path, data_only=True)
    mappings: Dict[str, str] = {}
    rows: List[Dict[str, str]] = []
    warnings: List[str] = []
    for ws in wb.worksheets:
        matrix = [[clean_cell(ws.cell(r, c).value) for c in range(1, ws.max_column + 1)] for r in range(1, min(ws.max_row, 3000) + 1)]
        header_row = -1
        cols: Dict[str, int] = {}
        for r, row in enumerate(matrix[:30]):
            local_cols: Dict[str, int] = {}
            for c, val in enumerate(row):
                low = val.lower().strip()
                if low in ("ean", "ean13") or low.startswith("ean"):
                    local_cols["ean"] = c
                if "sku" in low or "code article" in low or low == "article":
                    local_cols["sku"] = c
                if "color" in low or "couleur" in low or "coloris" in low:
                    local_cols["color"] = c
                if "size" in low or "taille" in low:
                    local_cols["size"] = c
                if "batch" in low:
                    local_cols["batch"] = c
            if "ean" in local_cols and ("sku" in local_cols or "size" in local_cols):
                cols = local_cols
                header_row = r
                break
        if header_row == -1:
            for row in matrix:
                row_text = " ".join(row)
                eans = EAN_RE.findall(row_text)
                if not eans:
                    continue
                size = detect_size_from_sku(row_text)
                color = detect_color_from_sku_or_text(row_text)
                if size:
                    key = f"{normalize_key(color)}|{size}" if color else f"*|{size}"
                    mappings[key] = eans[0]
                    mappings.setdefault(f"*|{size}", eans[0])
                    rows.append({"color": color, "size": size, "ean": eans[0], "source": ws.title})
            continue

        for row in matrix[header_row + 1:]:
            if not any(row):
                continue
            ean_val = row[cols.get("ean", -1)] if 0 <= cols.get("ean", -1) < len(row) else ""
            ean_m = EAN_RE.search(str(ean_val)) or EAN_RE.search(" ".join(row))
            if not ean_m:
                continue
            ean = ean_m.group(1)
            sku = row[cols.get("sku", -1)] if 0 <= cols.get("sku", -1) < len(row) else ""
            raw_size = row[cols.get("size", -1)] if 0 <= cols.get("size", -1) < len(row) else ""
            size = normalize_size(raw_size) if raw_size and is_supported_size(raw_size) else detect_size_from_sku(sku or " ".join(row))
            raw_color = row[cols.get("color", -1)] if 0 <= cols.get("color", -1) < len(row) else ""
            color = normalize_color(raw_color) if raw_color else detect_color_from_sku_or_text(sku or " ".join(row))
            if not size:
                warnings.append(f"EAN {ean} in sheet {ws.title} has no detectable size.")
                continue
            key = f"{normalize_key(color)}|{size}" if color else f"*|{size}"
            mappings[key] = ean
            mappings.setdefault(f"*|{size}", ean)
            rows.append({"color": color, "size": size, "ean": ean, "source": ws.title, "sku": sku})
    return {"map": mappings, "rows": rows, "warnings": warnings}


# Final packing-sheet writer: applies _write_packing_sheet_extended, then
# rebuilds the carton-size table dynamically when the order uses numeric or
# otherwise non-standard size headers instead of the default XS..3XL set.
def _v13_write_packing_sheet(ws, metadata: Dict[str, Any], rows: List[Dict[str, Any]], ean_rows: List[Dict[str, str]], identification_code: str, shipment_date: Optional[datetime], pl_unit_kg: float, pallet_capacity: int = 16):
    _write_packing_sheet_extended(ws, metadata, rows, ean_rows, identification_code, shipment_date, pl_unit_kg, pallet_capacity)
    dynamic_sizes = sort_sizes({s for c in rows for s in (c.get("sizes") or {}).keys()})
    if not dynamic_sizes:
        return ws
    standard = ["XS", "S", "M", "L", "XL", "XXL", "3XL"]
    if dynamic_sizes == standard:
        return ws

    thin_border = _v13_border("thin")
    grey = PatternFill("solid", fgColor="F2F2F2")
    dark = PatternFill("solid", fgColor="7F7F7F")
    pale_blue = PatternFill("solid", fgColor="DDEBF7")
    white = PatternFill("solid", fgColor="FFFFFF")
    table_top = 23
    start = table_top + 2
    total_row = start + len(rows)
    size_start_col = 7  # G
    size_end_col = size_start_col + len(dynamic_sizes) - 1
    end_col = max(13, size_end_col)

    # Clear old fixed-size table/summary area while leaving upper metadata blocks intact.
    for rr in range(table_top, min(ws.max_row, total_row + 18) + 1):
        for cc in range(1, max(18, end_col) + 1):
            cell = ws.cell(rr, cc)
            if not isinstance(cell, MergedCell):
                cell.value = None
                cell.border = Border()
                cell.fill = white

    # Remove merges that may intersect the old table area.
    for merged in list(ws.merged_cells.ranges):
        if merged.min_row >= table_top:
            try:
                ws.unmerge_cells(str(merged))
            except Exception:
                pass

    # Wider sheet for numeric runs, but keep size columns compact.
    for i in range(size_start_col, size_end_col + 1):
        ws.column_dimensions[get_column_letter(i)].width = 6.0
    ws.column_dimensions["A"].width = 29
    ws.column_dimensions["B"].width = 12
    ws.column_dimensions["C"].width = 10
    ws.column_dimensions["D"].width = 42
    ws.column_dimensions["E"].width = 12
    ws.column_dimensions["F"].width = 3

    # Row 23: grey band over descriptive columns + size headers.
    _v13_merge(ws, table_top, 1, table_top, 5)
    for cc in range(1, end_col + 1):
        ws.cell(table_top, cc).fill = dark if cc <= 5 else white
        ws.cell(table_top, cc).border = thin_border
        ws.cell(table_top, cc).alignment = Alignment(horizontal="center", vertical="center")
    for idx, sz in enumerate(dynamic_sizes, size_start_col):
        _v13_set(ws.cell(table_top, idx), sz, 9 if len(dynamic_sizes) > 9 else 10, True, align="center")

    headers = {1:"Référence", 2:"Coloris", 3:"N° de\ncolis", 4:"EAN", 5:"Nbre de\nPcs/Carton"}
    for cc, txt in headers.items():
        _v13_set(ws.cell(table_top + 1, cc), txt, 10, True, align="center")
    _v13_merge(ws, table_top + 1, size_start_col, table_top + 1, size_end_col)
    _v13_set(ws.cell(table_top + 1, size_start_col), "Quantité / Taille", 10, True, align="center")
    for cc in range(1, end_col + 1):
        ws.cell(table_top + 1, cc).border = thin_border

    color = rows[0].get("color", "") if rows else ""
    ref = rows[0].get("reference") or metadata.get("product_name", "") if rows else metadata.get("product_name", "")
    total_pcs = sum(int(c.get("total_pcs", 0) or 0) for c in rows)
    for ridx, carton in enumerate(rows, start):
        _v13_set(ws.cell(ridx, 1), carton.get("reference") or ref, 10, False)
        _v13_set(ws.cell(ridx, 2), color, 10, False, align="center")
        _v13_set(ws.cell(ridx, 3), carton.get("box_no", ridx - start + 1), 10, False, align="center")
        _v13_set(ws.cell(ridx, 4), carton.get("ean", ""), 9, False, fill=pale_blue, align="left")
        _v13_set(ws.cell(ridx, 5), int(carton.get("total_pcs", 0) or 0), 10, False, align="right")
        for cidx, sz in enumerate(dynamic_sizes, size_start_col):
            q = int((carton.get("sizes") or {}).get(sz, 0) or 0)
            _v13_set(ws.cell(ridx, cidx), q if q else "", 9 if len(dynamic_sizes)>9 else 10, False, align="center")
        for cc in range(1, end_col + 1):
            ws.cell(ridx, cc).border = thin_border
        ean_text = str(carton.get("ean", ""))
        ws.row_dimensions[ridx].height = max(22, 15 * max(1, math.ceil(len(ean_text)/42)) + 5)

    _v13_merge(ws, total_row, 1, total_row, 4)
    _v13_set(ws.cell(total_row, 1), "TOTAL", 10, True, fill=grey, align="center")
    _v13_set(ws.cell(total_row, 5), total_pcs, 10, True, fill=dark, align="right")
    ws.cell(total_row, 5).font = Font(name="Aptos Narrow", size=10, bold=True, color="FFFFFF")
    for cidx, sz in enumerate(dynamic_sizes, size_start_col):
        tq = sum(int((c.get("sizes") or {}).get(sz, 0) or 0) for c in rows)
        _v13_set(ws.cell(total_row, cidx), tq if tq else "", 9 if len(dynamic_sizes)>9 else 10, True, fill=dark, align="center")
        ws.cell(total_row, cidx).font = Font(name="Aptos Narrow", size=9 if len(dynamic_sizes)>9 else 10, bold=True, color="FFFFFF")
    for cc in range(1, end_col + 1):
        ws.cell(total_row, cc).border = thin_border
    _v13_apply_border(ws, table_top, 1, total_row, end_col, "medium")

    # Rebuild EAN summary using the same numeric sizes.
    erow = total_row + 3
    _v13_set(ws.cell(erow, 1), "EAN", 10, True)
    _v13_set(ws.cell(erow, 2), "QUANTITIES", 10, True)
    ws.cell(erow, 1).border = thin_border; ws.cell(erow, 2).border = thin_border
    erow += 1
    for sz in dynamic_sizes:
        tq = sum(int((c.get("sizes") or {}).get(sz, 0) or 0) for c in rows)
        if tq <= 0:
            continue
        ean = _v13_ean_for_size(ean_rows, color, sz, rows[0] if rows else None)
        _v13_set(ws.cell(erow, 1), ean, 10, False, align="center")
        _v13_set(ws.cell(erow, 2), tq, 10, False, align="right")
        ws.cell(erow, 1).border = thin_border; ws.cell(erow, 2).border = thin_border
        erow += 1

    # Center the identification title across the full dynamic table width.
    try:
        for merged in list(ws.merged_cells.ranges):
            if merged.min_row == 1 and merged.max_row == 1:
                ws.unmerge_cells(str(merged))
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=end_col)
        ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    except Exception:
        pass
    ws.print_area = f"A1:{get_column_letter(end_col)}{max(erow + 1, 37)}"
    return ws
