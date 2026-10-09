import io
import json
import os
import re
import tempfile
import zipfile
from datetime import datetime
from pathlib import Path
from typing import Dict, List, Any, Optional

try:
    from pypdf import PdfReader
except Exception:
    PdfReader = None

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.page import PageMargins
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont

try:
    import pytesseract  # optional, only used if PyMuPDF + Pillow are available locally
    from PIL import Image
    OCR_AVAILABLE = True
except Exception:
    OCR_AVAILABLE = False

APP_DIR = Path(__file__).resolve().parent
UPLOAD_DIR = Path(os.environ.get("TEOKROZE_DATA_DIR", APP_DIR)) / "uploads"  # TEOKROZE: папка за запис
UPLOAD_DIR.mkdir(exist_ok=True)
PACKING_TEMPLATE_XLSX = APP_DIR / "templates" / "acne_packing_template.xlsx"

PACKING_TEMPLATES = {
    "SE": APP_DIR / "templates" / "acne_packing_template_se.xlsx",
    "US": APP_DIR / "templates" / "acne_packing_template_us.xlsx",
    "KR": APP_DIR / "templates" / "acne_packing_template_kr.xlsx",
    "JP": APP_DIR / "templates" / "acne_packing_template_jp.xlsx",
    "CN": APP_DIR / "templates" / "acne_packing_template_cn.xlsx",
}

DESTINATION_PRESETS = {
    "SE": {
        "display": "Sweden / WH SE Spånga",
        "summary": "Spanga",
        "label_address": ["AIU Logistics", "Finspångsgatan 49", "163 53 Spånga", "Sweden"],
        "packing_rows": {11: "AIU Logistics", 12: "Finspångsgatan 49", 13: "163 53 Spånga", 14: "Sweden"},
    },
    "US": {
        "display": "United States / 3PL US",
        "summary": "USA",
        "label_address": ["Acne Corp", "c/o Bergen Logistics", "Attn: Receiving Dept", "5903 West Side Avenue", "North Bergen, NJ 07047", "United States", "(201) 854-1512 ext 418"],
        "packing_rows": {10: "Acne Corp", 11: "c/o Bergen Logistics", 12: "Attn: Receiving Dept", 13: "5903 West Side Avenue", 14: "North Bergen, NJ 07047", 15: "United States", 16: "(201) 854-1512 ext 418"},
    },
    "KR": {
        "display": "Korea / 3PL KR",
        "summary": "Korea",
        "label_address": ["Acne Studios Korea LLC", "c/o Maersk Logistics DC8", "Dock 1, 2F, 1911, Hwangmu-ro, Bubal-eup, Icheon-si", "Gyeonggi-do 17405", "Republic of Korea"],
        "packing_rows": {10: "Acne Studios Korea LLC", 11: "c/o Maersk Logistics DC8", 12: "Dock 1, 2F, 1911, Hwangmu-ro, Bubal-eup, Icheon-si", 13: "Gyeonggi-do 17405", 14: "Republic of Korea"},
    },
    "JP": {
        "display": "Japan / 3PL JP",
        "summary": "Japan",
        "label_address": ["Acne Aoyama c/o Maersk Logistics", "GLP MFLP Ichikawa Shiohama", "5F 1-6-3 Shiohama, Ichikawa City", "Chiba 272-0127", "Japan"],
        "packing_rows": {11: "Acne Aoyama c/o Maersk Logistics", 12: "GLP MFLP Ichikawa Shiohama", 13: "5F 1-6-3 Shiohama, Ichikawa City", 14: "Chiba 272-0127", 15: "Japan"},
    },
    "CN": {
        "display": "China / 3PL CN",
        "summary": "China",
        "label_address": ["Acne Studios China Co., Ltd.", "c/o Cargo Services China (Shanghai)", "1800 Wenchuan Road,", "Shanghai Baoshan City Industrial Park", "201901 Shanghai", "China"],
        "packing_rows": {10: "Acne Studios China Co., Ltd.", 11: "c/o Cargo Services China (Shanghai)", 12: "1800 Wenchuan Road,", 13: "Shanghai Baoshan City Industrial Park", 14: "201901 Shanghai", 15: "China"},
    },
}


THIN = Side(style="thin", color="000000")
MEDIUM = Side(style="medium", color="000000")
BORDER_THIN = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
BORDER_MEDIUM = Border(left=MEDIUM, right=MEDIUM, top=MEDIUM, bottom=MEDIUM)
BLACK_FILL = PatternFill("solid", fgColor="000000")
GRAY_FILL = PatternFill("solid", fgColor="F3F4F6")
LIGHT_BLUE_FILL = PatternFill("solid", fgColor="EAF3FF")
LIGHT_GREEN_FILL = PatternFill("solid", fgColor="EAF7EA")
WHITE_FILL = PatternFill("solid", fgColor="FFFFFF")


def clean_line(s: str) -> str:
    return re.sub(r"\s+", " ", str(s or "").strip())


def euro_number_to_float(value: str) -> float:
    value = clean_line(str(value)).replace(" ", "")
    if not value:
        return 0.0
    if "," in value and "." in value:
        value = value.replace(".", "").replace(",", ".")
    else:
        value = value.replace(",", ".")
    try:
        return float(value)
    except Exception:
        return 0.0


def int_qty(value: Any) -> int:
    try:
        f = float(value)
    except Exception:
        f = euro_number_to_float(str(value))
    if abs(f - int(f)) < 0.0001:
        return int(f)
    return round(f, 3)


def is_article_code(token: str) -> bool:
    token = clean_line(token)
    if len(token) < 6:
        return False
    # Examples: RW-WN-SKIR000070, ABC-12-XYZ
    return bool(re.match(r"^[A-Z0-9]{1,6}-[A-Z0-9]{1,6}-[A-Z0-9]{3,}$", token))


def is_item_code(token: str) -> bool:
    token = clean_line(token)
    return bool(re.match(r"^[A-Z0-9]{2,}-[A-Z0-9]{3,}$", token))


def is_size_token(token: str) -> bool:
    token = clean_line(token)
    if not token:
        return False
    if re.match(r"^\d{1,3}([./,]\d)?$", token):
        return True
    if re.match(r"^(XXS|XS|S|M|L|XL|XXL|XXXL|OS|ONE|ONESIZE|ONE SIZE)$", token, re.I):
        return True
    return False


def infer_color_code(item_number: str, size: str = "") -> str:
    """Derive Acne color code from item number where possible.

    Examples from the uploaded PDFs:
    AF0617-EA1032 + size 32 -> EA10
    AF0617-J83034 + size 34 -> J830
    """
    item_number = clean_line(item_number)
    size = clean_line(str(size)).replace(".0", "")
    if "-" not in item_number:
        return ""
    tail = item_number.split("-")[-1]
    if size and tail.endswith(size) and len(tail) > len(size):
        return tail[:-len(size)]
    m = re.match(r"^([A-Z]+\d+)", tail)
    return m.group(1) if m else tail


def extract_text_with_ocr_fallback(pdf_path: Path) -> Dict[str, Any]:
    """Extract text from PDF without requiring compiled PyMuPDF.

    V14 intentionally uses pypdf first because it is pure Python and does not need
    Visual Studio/C++ build tools.  The old version imported PyMuPDF (`fitz`) at
    startup; on some PCs/Python versions pip tried to build PyMuPDF from source
    and failed with "Unable to find Visual Studio".

    OCR is kept as an optional fallback only when the user already has compatible
    PyMuPDF + Pillow + Tesseract available.  Normal Acne PO PDFs from the tests
    have selectable text, so pypdf is enough and installation is much safer.
    """
    warnings: List[str] = []
    text_parts: List[str] = []
    used_ocr = False

    # 1) Safe default: pure-Python text extraction.
    if PdfReader is not None:
        try:
            reader = PdfReader(str(pdf_path))
            for page_no, page in enumerate(reader.pages, start=1):
                try:
                    txt = page.extract_text() or ""
                except Exception as exc:
                    warnings.append(f"Page {page_no}: pypdf text extraction failed: {exc}")
                    txt = ""
                text_parts.append(txt)
            text = "\n".join(text_parts)
            if len(text.strip()) >= 30:
                return {"text": text, "used_ocr": False, "warnings": warnings}
        except Exception as exc:
            warnings.append(f"pypdf failed: {exc}")
    else:
        warnings.append("pypdf is not installed. Run the BAT file again to install requirements.")

    # 2) Optional legacy fallback: use PyMuPDF if already available.
    #    It is NOT a required dependency anymore, so app startup never fails because
    #    of missing Visual Studio / missing fitz.
    try:
        import fitz  # type: ignore
        doc = fitz.open(str(pdf_path))
        fallback_parts: List[str] = []
        for page_no, page in enumerate(doc, start=1):
            page_text = page.get_text("text") or ""
            if len(page_text.strip()) >= 30:
                fallback_parts.append(page_text)
                continue

            if not OCR_AVAILABLE:
                warnings.append(f"Page {page_no}: no selectable text and optional OCR libraries are not available.")
                fallback_parts.append(page_text)
                continue

            try:
                used_ocr = True
                pix = page.get_pixmap(matrix=fitz.Matrix(2.2, 2.2), alpha=False)
                image = Image.open(io.BytesIO(pix.tobytes("png")))
                ocr_text = pytesseract.image_to_string(image, lang="eng")
                fallback_parts.append(ocr_text)
            except Exception as exc:
                warnings.append(f"Page {page_no}: optional OCR failed: {exc}")
                fallback_parts.append(page_text)
        return {"text": "\n".join(fallback_parts), "used_ocr": used_ocr, "warnings": warnings}
    except Exception:
        # No PyMuPDF; this is fine for normal selectable Acne PDFs.
        pass

    warnings.append("No readable PDF text was found. This may be a scanned PDF; install OCR/Tesseract support or send the PDF for parser tuning.")
    return {"text": "\n".join(text_parts), "used_ocr": used_ocr, "warnings": warnings}


def get_after_label(lines: List[str], label: str, default: str = "") -> str:
    label_low = label.lower()
    for i, line in enumerate(lines):
        if line.lower() == label_low and i + 1 < len(lines):
            return lines[i + 1]
        if line.lower().startswith(label_low + " "):
            return line[len(label):].strip()
    return default


def extract_between(lines: List[str], start_label: str, end_patterns: List[str], max_lines: int = 8) -> str:
    start = -1
    for i, line in enumerate(lines):
        if line.lower() == start_label.lower():
            start = i + 1
            break
    if start == -1:
        return ""
    out = []
    for line in lines[start:start + max_lines]:
        if any(re.search(pat, line, flags=re.I) for pat in end_patterns):
            break
        if line:
            out.append(line)
    return " ".join(out).strip()


def extract_customer_address(lines: List[str]) -> str:
    # In provided PDFs the full customer block appears after the supplier address and before Telephone.
    for i, line in enumerate(lines):
        if line.upper() == "ACNE STUDIOS AB":
            block = []
            for j in range(i, min(i + 6, len(lines))):
                if re.search(r"^(Telephone|Fax|Giro|Tax registration number|Purchase order)$", lines[j], re.I):
                    break
                if lines[j]:
                    block.append(lines[j])
            if len(block) >= 2:
                return "\n".join(block)
    # Fallback around Delivery address.
    for i, line in enumerate(lines):
        if line.lower() == "delivery address" and i + 1 < len(lines):
            return lines[i + 1]
    return "Acne Studios AB"


def parse_order_lines(lines: List[str]) -> List[Dict[str, Any]]:
    start = -1
    for i, line in enumerate(lines):
        low = line.lower()
        # PyMuPDF usually returns the whole header on one line.
        if low.startswith("line nr item number"):
            start = i + 1
            break
        # pypdf returns the same header as multiple lines:
        # Line Nr / Item number / Color/Length / Size / ... / Ex Factory.
        if low == "line nr":
            for k in range(i + 1, min(i + 14, len(lines))):
                if lines[k].lower() == "ex factory":
                    start = k + 1
                    break
            if start != -1:
                break
    if start == -1:
        # Last-resort fallback: start at the first numbered row followed by an item code.
        for i in range(len(lines) - 1):
            if re.match(r"^\d+$", lines[i]) and is_item_code(lines[i + 1]):
                start = i
                break
    if start == -1:
        return []

    items: List[Dict[str, Any]] = []
    i = start
    while i < len(lines):
        line = lines[i]
        if not re.match(r"^\d+$", line):
            i += 1
            continue
        line_nr = int(line)
        if i + 1 >= len(lines) or not is_item_code(lines[i + 1]):
            i += 1
            continue
        item_number = lines[i + 1]
        article_idx = -1
        for k in range(i + 3, min(i + 12, len(lines))):
            if is_article_code(lines[k]):
                article_idx = k
                break
        if article_idx == -1 or article_idx + 1 >= len(lines):
            i += 1
            continue

        size = lines[article_idx - 1]
        color_lines = [x for x in lines[i + 2:article_idx - 1] if x]
        color = " ".join(color_lines).replace(" /", "/").replace("/ ", "/")
        qty = int_qty(euro_number_to_float(lines[article_idx + 1]))
        article_name = lines[article_idx]

        items.append({
            "line_nr": line_nr,
            "item_number": item_number,
            "color": color,
            "color_code": infer_color_code(item_number, size),
            "size": size,
            "article_name": article_name,
            "quantity": qty,
            "unit_price": lines[article_idx + 2] if article_idx + 2 < len(lines) else "",
            "discount": lines[article_idx + 3] if article_idx + 3 < len(lines) else "",
            "amount": lines[article_idx + 4] if article_idx + 4 < len(lines) else "",
        })

        next_i = None
        for p in range(article_idx + 2, len(lines)):
            if lines[p] == str(line_nr + 1) and p + 1 < len(lines) and is_item_code(lines[p + 1]):
                next_i = p
                break
        if next_i is None:
            break
        i = next_i
    return items



def excel_value_to_text(value: Any) -> str:
    """Convert Excel cell values to clean display text used by the web editor."""
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d")
    # date-only objects also have strftime but are not datetime
    if hasattr(value, "strftime") and value.__class__.__name__ == "date":
        try:
            return value.strftime("%Y-%m-%d")
        except Exception:
            pass
    if isinstance(value, float) and abs(value - int(value)) < 0.0001:
        return str(int(value))
    return clean_line(str(value))


def _find_cell(ws, patterns: List[str], min_row: int = 1, max_row: Optional[int] = None) -> Optional[tuple]:
    """Find first cell whose text contains any pattern, case-insensitive."""
    max_row = max_row or ws.max_row
    pats = [p.upper() for p in patterns]
    for row in ws.iter_rows(min_row=min_row, max_row=max_row):
        for cell in row:
            txt = excel_value_to_text(cell.value).upper()
            if txt and any(p in txt for p in pats):
                return cell.row, cell.column
    return None


def _value_right_of_label(ws, patterns: List[str], default: str = "") -> str:
    found = _find_cell(ws, patterns, max_row=min(ws.max_row, 80))
    if not found:
        return default
    r, c = found
    for cc in range(c + 1, min(ws.max_column, c + 8) + 1):
        txt = excel_value_to_text(ws.cell(r, cc).value)
        if txt:
            return txt
    return default


def _detect_destination_from_packing(ws, filename: str = "") -> str:
    name = (filename or "").upper()
    if "3PL CN" in name or " CN" in name:
        return "CN"
    if "3PL US" in name or " US" in name:
        return "US"
    if "3PL JP" in name or " JP" in name:
        return "JP"
    if "3PL KR" in name or " KR" in name:
        return "KR"
    # Address-based fallback from the uploaded destination templates.
    text_parts = []
    for r in range(8, min(ws.max_row, 25) + 1):
        for c in range(1, min(ws.max_column, 24) + 1):
            txt = excel_value_to_text(ws.cell(r, c).value)
            if txt:
                text_parts.append(txt.upper())
    blob = " ".join(text_parts)
    if "CARGO SERVICES CHINA" in blob or "SHANGHAI" in blob:
        return "CN"
    if "BERGEN LOGISTICS" in blob or "NORTH BERGEN" in blob:
        return "US"
    if "MAERSK LOGISTICS DC8" in blob or "REPUBLIC OF KOREA" in blob or "GYEONGGI" in blob:
        return "KR"
    if "ICHIKAWA" in blob or "AOYAMA" in blob or "JAPAN" in blob:
        return "JP"
    if "SPÅNGA" in blob or "SPANGA" in blob or "AIU LOGISTICS" in blob or "SWEDEN" in blob:
        return "SE"
    return "SE"


def _read_delivery_address_from_packing(ws, dest_key: str) -> str:
    header = _find_cell(ws, ["DELIVERY ADDRESS"], max_row=30)
    if not header:
        return default_destination_address(dest_key)
    hr, hc = header
    # Destination lines are usually below the header, around J/K.  Merged templates
    # put the real value in the anchor cell, so scan a small address box.
    lines = []
    for r in range(hr + 1, min(hr + 9, ws.max_row) + 1):
        row_vals = []
        for c in range(max(1, hc - 1), min(ws.max_column, hc + 7) + 1):
            txt = excel_value_to_text(ws.cell(r, c).value)
            if txt and txt not in row_vals:
                row_vals.append(txt)
        line = clean_line(" ".join(row_vals))
        if line and not any(stop in line.upper() for stop in ["NNW", "NET WEIGHT", "GROSS WEIGHT", "CURRENT DATE", "SHIP DATE"]):
            lines.append(line)
    # Remove accidental duplicates while keeping order.
    dedup = []
    for line in lines:
        if line not in dedup:
            dedup.append(line)
    return "\n".join(dedup[:7]) if dedup else default_destination_address(dest_key)


def _is_active_scale_row(ws, row_no: int, start_col: int, end_col: int) -> bool:
    """Detect the chosen red/highlighted size-scale row in existing packing sheets."""
    score = 0
    for c in range(start_col, end_col + 1):
        cell = ws.cell(row_no, c)
        if excel_value_to_text(cell.value):
            fill_type = getattr(cell.fill, "fill_type", None) or getattr(cell.fill, "patternType", None)
            if fill_type:
                score += 2
            fg = getattr(cell.fill.fgColor, "rgb", "") if cell.fill and cell.fill.fgColor else ""
            if isinstance(fg, str) and fg and fg not in ("00000000", "FFFFFFFF"):
                score += 1
            if cell.style_id not in (0, 15, 16, 18, 19):
                score += 1
    return score >= 4


def _detect_active_size_row(ws, header_row: int, start_col: int, end_col: int, used_cols: List[int]) -> int:
    # First preference: the template-selected/highlighted scale row.
    for r in range(max(1, header_row - 20), header_row):
        values = [excel_value_to_text(ws.cell(r, c).value) for c in range(start_col, end_col + 1)]
        if any(values) and _is_active_scale_row(ws, r, start_col, end_col):
            return r
    # Fallback: choose the row with values in all quantity columns.
    candidates = []
    for r in range(max(1, header_row - 20), header_row):
        vals = {c: excel_value_to_text(ws.cell(r, c).value) for c in range(start_col, end_col + 1)}
        if not any(vals.values()):
            continue
        score = sum(1 for c in used_cols if vals.get(c))
        if score:
            candidates.append((score, r))
    if candidates:
        candidates.sort(reverse=True)
        return candidates[0][1]
    return header_row - 6


def _norm_header_text(value: Any) -> str:
    """Normalize worksheet header labels for fuzzy matching."""
    txt = excel_value_to_text(value).upper()
    txt = txt.replace("#", " NO ").replace("/", " ").replace("-", " ")
    txt = re.sub(r"[^A-Z0-9]+", " ", txt)
    return re.sub(r"\s+", " ", txt).strip()


def _find_packing_table_layout(ws) -> Dict[str, Any]:
    """Detect the main carton/box table in Acne packing sheets.

    The customer templates are mostly stable, but some sheets do not literally use
    "CARTON #".  The WH/Spånga Pyxano example has the first header as "Hanging"
    and the dimension column as "Boxes".  This detector therefore anchors on the
    reliable columns (STYLE NAME / COLOR NAME / COLOR CODE / CHOOSE APPLICABLE /
    TTL UNITS) and treats CARTON, CTN, BOX or HANGING as optional carton headers.
    """
    candidates = []
    max_scan_row = min(ws.max_row, 140)
    for r in range(1, max_scan_row + 1):
        cells = [_norm_header_text(ws.cell(r, c).value) for c in range(1, ws.max_column + 1)]
        blob = " ".join([x for x in cells if x])
        if not blob:
            continue
        has_style = "STYLE NAME" in blob
        has_color = "COLOR NAME" in blob
        has_code = "COLOR CODE" in blob
        has_size = "CHOOSE APPLICABLE" in blob or "SIZE SCALE" in blob
        has_ttl = "TTL UNITS" in blob or "TOTAL UNITS" in blob
        score = sum([has_style, has_color, has_code, has_size, has_ttl])
        # Strong table header: style/color/code/size are the reliable anchors.
        if score >= 4:
            candidates.append((score, r, cells))
    if not candidates:
        # Legacy fallback: literal carton/box label row.
        found = _find_cell(ws, ["CARTON #", "CARTON#", "CTN #", "BOX", "BOX #"], max_row=120)
        if not found:
            raise ValueError("Could not find the Acne carton/box table header in this packing list.")
        header_row = found[0]
        cells = [_norm_header_text(ws.cell(header_row, c).value) for c in range(1, ws.max_column + 1)]
    else:
        candidates.sort(reverse=True)
        _score, header_row, cells = candidates[0]

    def first_col(match_fn, default: Optional[int] = None) -> Optional[int]:
        for idx, txt in enumerate(cells, start=1):
            if txt and match_fn(txt):
                return idx
        return default

    style_col = first_col(lambda t: "STYLE NAME" in t)
    color_col = first_col(lambda t: "COLOR NAME" in t)
    color_code_col = first_col(lambda t: "COLOR CODE" in t)
    size_start_col = first_col(lambda t: "CHOOSE APPLICABLE" in t or "SIZE SCALE" in t)
    ttl_col = first_col(lambda t: "TTL UNITS" in t or "TOTAL UNITS" in t, None)

    if not (style_col and color_col and color_code_col and size_start_col):
        raise ValueError("Found a possible table header, but could not identify Style/Color/Size columns.")

    # Search left of STYLE NAME for an explicit carton/box column.  If the template
    # uses "Hanging" or leaves the carton column blank, we still parse rows and
    # assign carton numbers sequentially.
    carton_col = None
    explicit_carton = False
    for c in range(1, style_col):
        txt = cells[c - 1]
        if not txt:
            continue
        if any(k in txt for k in ("CARTON", "CTN", "BOX", "BOXES", "HANGING")):
            carton_col = c
            explicit_carton = any(k in txt for k in ("CARTON", "CTN", "BOX NO", "BOX NUMBER"))
            break
    if carton_col is None:
        carton_col = max(1, style_col - 1)

    size_end_col = (ttl_col - 1) if ttl_col and ttl_col > size_start_col else min(ws.max_column, size_start_col + 15)
    size_end_col = min(size_end_col, size_start_col + 15)

    return {
        "header_row": header_row,
        "carton_col": carton_col,
        "explicit_carton": explicit_carton,
        "style_col": style_col,
        "color_col": color_col,
        "color_code_col": color_code_col,
        "size_start_col": size_start_col,
        "size_end_col": size_end_col,
        "ttl_col": ttl_col,
    }


def parse_packing_workbook(xlsx_path: Path, original_filename: str = "") -> Dict[str, Any]:
    """Import an already-existing Acne packing list so labels can be regenerated/edited.

    V16 parser notes:
    - accepts CARTON #, CTN #, BOX/BOX # and templates where the first column says
      Hanging while rows themselves have no carton number;
    - if the sheet does not contain explicit carton/box numbers, one data row is
      treated as one carton and numbered 1, 2, 3...;
    - anchors the parser on STYLE NAME / COLOR NAME / COLOR CODE / SIZE SCALE / TTL
      columns, so small wording changes still connect the dots.
    """
    wb = load_workbook(xlsx_path, data_only=True)
    ws = wb["Packing List"] if "Packing List" in wb.sheetnames else wb.active

    dest_key = _detect_destination_from_packing(ws, original_filename)
    po_number = _value_right_of_label(ws, ["Acne Studios PO #", "PO #"], "")
    product_name = _value_right_of_label(ws, ["Style Name"], "")
    current_date = _value_right_of_label(ws, ["Current Date"], "")
    ship_date = _value_right_of_label(ws, ["Ship Date"], current_date)

    layout = _find_packing_table_layout(ws)
    header_row = layout["header_row"]
    carton_col = layout["carton_col"]
    style_col = layout["style_col"]
    color_col = layout["color_col"]
    color_code_col = layout["color_code_col"]
    size_start_col = layout["size_start_col"]
    size_end_col = layout["size_end_col"]

    # Find carton/box data rows and the quantity columns that are actually used.
    data_rows = []
    used_cols = set()
    blank_after_data = 0
    for r in range(header_row + 1, ws.max_row + 1):
        row_blob = " ".join(_norm_header_text(ws.cell(r, c).value) for c in range(1, min(ws.max_column, 25) + 1))
        if data_rows and ("SIZE SCALE" in row_blob or "PRODUCT CATEGORIES" in row_blob):
            break

        carton_txt = excel_value_to_text(ws.cell(r, carton_col).value) if carton_col else ""
        style_txt = excel_value_to_text(ws.cell(r, style_col).value)
        color_txt = excel_value_to_text(ws.cell(r, color_col).value)
        code_txt = excel_value_to_text(ws.cell(r, color_code_col).value)

        row_qty_cols = []
        for c in range(size_start_col, size_end_col + 1):
            v = ws.cell(r, c).value
            try:
                q = int_qty(v)
            except Exception:
                q = 0
            if q:
                row_qty_cols.append(c)
                used_cols.add(c)

        has_identity = bool(carton_txt or style_txt or color_txt or code_txt)
        if not row_qty_cols:
            if data_rows:
                blank_after_data += 1
                # Some templates have a few zero/formula rows before the lower block.
                if blank_after_data >= 6 and not has_identity:
                    break
            continue
        blank_after_data = 0

        # Accept both explicit carton-number rows and box-style rows where carton
        # numbers are omitted but style/color/quantity data clearly exist.
        data_rows.append((r, row_qty_cols))

    if not data_rows:
        raise ValueError("No carton/box quantity rows were found in this packing list.")

    active_size_row = _detect_active_size_row(ws, header_row, size_start_col, size_end_col, sorted(used_cols))
    size_by_col = {c: excel_value_to_text(ws.cell(active_size_row, c).value) for c in range(size_start_col, size_end_col + 1)}

    cartons = []
    items_map: Dict[tuple, Dict[str, Any]] = {}
    last_style = product_name
    last_color = ""
    last_code = ""
    for seq_no, (r, qty_cols) in enumerate(data_rows, start=1):
        raw_carton_no = excel_value_to_text(ws.cell(r, carton_col).value) if carton_col else ""
        # If the source uses a non-carton label like "Hanging" and leaves rows blank,
        # number them sequentially so labels still say 1 of N, 2 of N...
        carton_no = raw_carton_no or seq_no
        row_style = excel_value_to_text(ws.cell(r, style_col).value) or last_style or product_name
        row_color = excel_value_to_text(ws.cell(r, color_col).value) or last_color
        row_code = excel_value_to_text(ws.cell(r, color_code_col).value) or last_code
        last_style, last_color, last_code = row_style, row_color, row_code
        skus = []
        for c in qty_cols:
            size = size_by_col.get(c) or f"Size {c - size_start_col + 1}"
            q = int_qty(ws.cell(r, c).value)
            if not q:
                continue
            sku = {
                "size": str(size).replace(".0", ""),
                "units": q,
                "item_number": "",
                "color": row_color,
                "color_code": row_code,
                "article_name": row_style or product_name,
            }
            skus.append(sku)
            key = (row_color, row_code, sku["size"], row_style or product_name)
            if key not in items_map:
                items_map[key] = {
                    "line_nr": len(items_map) + 1,
                    "item_number": "",
                    "color": row_color,
                    "color_code": row_code,
                    "size": sku["size"],
                    "article_name": row_style or product_name,
                    "quantity": 0,
                }
            items_map[key]["quantity"] += q
        if skus:
            cartons.append({
                "carton_no": carton_no,
                "color": row_color,
                "item_numbers": "",
                "skus": skus,
            })

    if not product_name:
        # Fallback to first carton style if the top style cell was blank.
        product_name = clean_line(cartons[0]["skus"][0].get("article_name", "")) if cartons else ""

    import_warning = "Imported from existing packing list. Carton split is preserved for labels. Switch carton draft to rebuild from edited lines."
    if any(not excel_value_to_text(ws.cell(r, carton_col).value) for r, _ in data_rows):
        import_warning += " Some source rows did not contain explicit Carton/Box numbers, so the app numbered those rows sequentially."

    order = {
        "id": safe_name(Path(original_filename or xlsx_path.name).stem) + "_packing",
        "source_filename": original_filename or xlsx_path.name,
        "source_type": "packing_xlsx",
        "destination_key": dest_key,
        "destination": DESTINATION_PRESETS[dest_key]["display"],
        "ship_to": _read_delivery_address_from_packing(ws, dest_key),
        "po_number": po_number,
        "po_line_ref": "",
        "season": "",
        "internal_order_type": "ImportedPackingList",
        "product_name": product_name,
        "style_code": product_name,
        "origin": "BULGARIA",
        "currency": "EUR",
        "ex_factory_date": ship_date,
        "shipment_date": ship_date,
        "total_quantity": sum(int_qty(x.get("quantity", 0)) for x in items_map.values()),
        "items": list(items_map.values()),
        "cartons": cartons,
        "draft_carton_strategy": "from_existing_packing",
        "used_ocr": False,
        "warnings": [import_warning],
        "raw_text_preview": "",
    }
    return ensure_destination_defaults(order, fill_address_if_empty=False)

def parse_pdf_order(pdf_path: Path) -> Dict[str, Any]:
    extract = extract_text_with_ocr_fallback(pdf_path)
    raw_text = extract["text"]
    lines = [clean_line(x) for x in raw_text.splitlines()]
    lines = [x for x in lines if x]

    po_number = get_after_label(lines, "Number")
    season = get_after_label(lines, "Season")
    internal_type = get_after_label(lines, "Internal order type")
    product_name = extract_between(lines, "Internal product name", [r"^\d{4}-\d{2}-\d{2}$", r"^Ex Factory"], max_lines=6)

    ex_factory_date = ""
    for line in lines:
        if re.match(r"^\d{4}-\d{2}-\d{2}$", line):
            ex_factory_date = line
            break

    items = parse_order_lines(lines)
    style_code = items[0]["article_name"] if items else ""
    total_quantity = sum(int_qty(x.get("quantity", 0)) for x in items)

    po_line_ref = ""
    m = re.search(r"##PO#([A-Z0-9\-]+)", raw_text)
    if m:
        po_line_ref = m.group(1)

    dest_key = destination_key_from_internal_type(internal_type)
    pdf_customer_address = extract_customer_address(lines)

    order = {
        "id": re.sub(r"[^A-Za-z0-9]+", "_", po_number or pdf_path.stem).strip("_") or pdf_path.stem,
        "source_filename": pdf_path.name,
        "supplier": "PELINTEX BULGARIA LTD",
        "supplier_address": "PELINTEX BULGARIA LTD\n5 HRISTO BOTEV BLVD\nRUSE\n7000\nBulgaria",
        "customer": "Acne Studios AB",
        "destination_key": dest_key,
        "destination": DESTINATION_PRESETS[dest_key]["display"],
        "pdf_customer_address": pdf_customer_address,
        "ship_to": default_destination_address(dest_key),
        "po_number": po_number,
        "po_line_ref": po_line_ref,
        "season": season,
        "internal_order_type": internal_type,
        "product_name": product_name,
        "style_code": style_code,
        "origin": "BULGARIA",
        "currency": "EUR",
        "ex_factory_date": ex_factory_date,
        "shipment_date": ex_factory_date,
        "total_quantity": total_quantity,
        "items": items,
        "draft_carton_strategy": "one_carton_per_color",
        "used_ocr": extract["used_ocr"],
        "warnings": extract["warnings"],
        "raw_text_preview": raw_text[:4000],
    }
    order["cartons"] = build_cartons(order)
    return order


def natural_key(value: Any):
    s = str(value)
    return [int(t) if t.isdigit() else t.lower() for t in re.split(r"(\d+)", s)]


def build_cartons(order: Dict[str, Any]) -> List[Dict[str, Any]]:
    strategy = order.get("draft_carton_strategy", "one_carton_per_color")
    if strategy == "from_existing_packing" and order.get("cartons"):
        # Existing packing lists already contain the exact carton split. Keep it unless
        # the user manually switches to a generated carton strategy.
        return order.get("cartons", [])
    items = order.get("items", [])
    cartons: List[Dict[str, Any]] = []
    if strategy == "one_carton_per_line":
        for idx, item in enumerate(items, start=1):
            cartons.append({
                "carton_no": idx,
                "color": item.get("color", ""),
                "item_numbers": item.get("item_number", ""),
                "skus": [{
                    "size": item.get("size", ""),
                    "units": int_qty(item.get("quantity", 0)),
                    "item_number": item.get("item_number", ""),
                    "color": item.get("color", ""),
                    "color_code": item.get("color_code", infer_color_code(item.get("item_number", ""), item.get("size", ""))),
                    "article_name": item.get("article_name", order.get("style_code", "")),
                }]
            })
        return cartons

    grouped: Dict[str, List[Dict[str, Any]]] = {}
    for item in items:
        grouped.setdefault(item.get("color", ""), []).append(item)
    for idx, (color, group) in enumerate(grouped.items(), start=1):
        skus = []
        for item in sorted(group, key=lambda x: natural_key(x.get("size", ""))):
            skus.append({
                "size": item.get("size", ""),
                "units": int_qty(item.get("quantity", 0)),
                "item_number": item.get("item_number", ""),
                "color": color,
                "color_code": item.get("color_code", infer_color_code(item.get("item_number", ""), item.get("size", ""))),
                "article_name": item.get("article_name", order.get("style_code", "")),
            })
        item_nums = ", ".join(sorted({x.get("item_number", "") for x in group if x.get("item_number")}, key=natural_key))
        cartons.append({"carton_no": idx, "color": color, "item_numbers": item_nums, "skus": skus})
    return cartons


def recalc_order(order: Dict[str, Any]) -> Dict[str, Any]:
    # Normalize destination, edited quantities and rebuild carton draft.
    ensure_destination_defaults(order, fill_address_if_empty=True)
    if not clean_line(order.get("shipment_date", "")):
        order["shipment_date"] = order.get("ex_factory_date", "")
    for idx, item in enumerate(order.get("items", []), start=1):
        item["line_nr"] = item.get("line_nr") or idx
        item["quantity"] = int_qty(item.get("quantity", 0))
        if not item.get("article_name"):
            item["article_name"] = order.get("style_code", "")
        if not item.get("color_code"):
            item["color_code"] = infer_color_code(item.get("item_number", ""), item.get("size", ""))
    order["total_quantity"] = sum(int_qty(x.get("quantity", 0)) for x in order.get("items", []))
    order["cartons"] = build_cartons(order)
    return order


def safe_name(value: str) -> str:
    value = re.sub(r"[^A-Za-z0-9._-]+", "_", value or "export").strip("_")
    return value[:120] or "export"


def order_export_stem(order: Dict[str, Any], index: Optional[int] = None) -> str:
    """Build a stable, unique, human-readable filename stem for exports.

    Older builds sometimes produced several exported Excel files with the same
    visible name when multiple orders/styles were packed into one ZIP.  This
    stem deliberately includes the order number, destination and style/article
    information, plus an optional order index for Export All, so the packing
    and label workbooks never overwrite each other when unzipped or copied into
    one folder.
    """
    parts = []
    if index is not None:
        parts.append(f"{index:02d}")
    po = str(order.get("po_number") or order.get("po_reference") or "ORDER").strip()
    dest = str(order.get("destination") or order.get("destination_key") or "DEST").strip()
    style = str(order.get("style_code") or order.get("article_name") or order.get("style_name") or "STYLE").strip()
    source = str(order.get("source_filename") or order.get("id") or "").strip()
    parts.extend([po, dest, style])
    if (not po or po == "ORDER") and source:
        parts.append(Path(source).stem)
    return safe_name("_".join([x for x in parts if x]))


def packing_filename(order: Dict[str, Any], index: Optional[int] = None) -> str:
    return safe_name(f"PACKING_LIST_{order_export_stem(order, index)}.xlsx")


def labels_filename(order: Dict[str, Any], index: Optional[int] = None) -> str:
    return safe_name(f"CARTON_LABELS_{order_export_stem(order, index)}.xlsx")


def zip_filename(order: Dict[str, Any]) -> str:
    return safe_name(f"ACNE_PACKING_AND_LABELS_{order_export_stem(order)}.zip")


def setup_print(ws, orientation: str, paper_size: int = 9):
    ws.page_setup.paperSize = paper_size  # A4
    ws.page_setup.orientation = orientation
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins = PageMargins(left=0.25, right=0.25, top=0.35, bottom=0.35, header=0.1, footer=0.1)


def style_range(ws, cell_range: str, border=BORDER_THIN, fill=None, font=None, align=None):
    for row in ws[cell_range]:
        for cell in row:
            cell.border = border
            if fill:
                cell.fill = fill
            if font:
                cell.font = font
            if align:
                cell.alignment = align


ACNE_SIZE_SCALE_ROWS = [
    (27, ["One Size"], "ACCESSORIES"),
    (28, ["40-42", "43-46"], "SOCKS"),
    (29, ["XS/S", "M/L"], "KNITS"),
    (30, ["2", "3", "4", "6", "8", "10"], "MINI"),
    (31, ["XXS", "XS", "S", "M", "L", "XL"], "KNITS, JERSEY, FLEECE"),
    (32, ["32", "34", "36", "38", "40", "42"], "MULTIPLE FEMALE CATEGORIES"),
    (33, ["44", "46", "48", "50", "52", "54"], "MULTIPLE MALE CATEGORIES"),
    (34, ["35", "36", "37", "38", "39", "40", "41", "42", "43", "44", "45", "46"], "SHOES"),
    (35, ["24", "25", "26", "27", "28", "29", "30", "31", "32", "33", "34", "36"], "DENIM"),
    (36, ["111", "112", "113", "114", "115", "116", "117", "121", "122", "123", "124", "125", "126", "127"], "DENIM BOTTOMS (REGULAR & PETIT)"),
    (37, ["140", "141", "142", "143", "144", "145", "146", "147", "148", "149", "150", "151", "152", "153", "154", "155"], "UNISEX DENIM (REGULAR & PETIT)"),
]


def destination_key_from_internal_type(internal_type: str) -> str:
    t = (internal_type or "").upper()
    if "CN3PL" in t:
        return "CN"
    if "US3PL" in t:
        return "US"
    if "JP3PL" in t:
        return "JP"
    if "KR3PL" in t:
        return "KR"
    # The plain PurchaseOrder destination in the uploaded set is WH SE Spånga.
    return "SE"


def destination_key_from_order(order: Dict[str, Any]) -> str:
    key = clean_line(order.get("destination_key", "")).upper()
    if key in DESTINATION_PRESETS:
        return key
    return destination_key_from_internal_type(order.get("internal_order_type", ""))


def destination_from_internal_type(internal_type: str) -> str:
    key = destination_key_from_internal_type(internal_type)
    return DESTINATION_PRESETS[key]["summary"]


def destination_display(order: Dict[str, Any]) -> str:
    key = destination_key_from_order(order)
    return DESTINATION_PRESETS[key]["display"]


def destination_summary(order: Dict[str, Any]) -> str:
    key = destination_key_from_order(order)
    return DESTINATION_PRESETS[key]["summary"]


def default_destination_address(key: str) -> str:
    preset = DESTINATION_PRESETS.get((key or "SE").upper(), DESTINATION_PRESETS["SE"])
    return "\n".join(preset["label_address"])


def ensure_destination_defaults(order: Dict[str, Any], fill_address_if_empty: bool = True) -> Dict[str, Any]:
    key = destination_key_from_order(order)
    order["destination_key"] = key
    order["destination"] = DESTINATION_PRESETS[key]["display"]
    if fill_address_if_empty and not clean_line(order.get("ship_to", "")):
        order["ship_to"] = default_destination_address(key)
    return order


def choose_size_scale(sizes: List[str]) -> tuple:
    clean_sizes = [str(s).replace(".0", "").strip() for s in sizes if str(s).strip()]
    size_set = set(clean_sizes)
    for row_no, scale, category in ACNE_SIZE_SCALE_ROWS:
        if size_set and size_set.issubset(set(scale)):
            return row_no, scale, category
    # Fallback: use the actual detected sizes in a custom dynamic row.
    return 32, clean_sizes[:16], "DYNAMIC SIZE SCALE"


def carton_color_code(carton: Dict[str, Any]) -> str:
    codes = []
    for sku in carton.get("skus", []):
        code = sku.get("color_code") or infer_color_code(sku.get("item_number", ""), sku.get("size", ""))
        if code and code not in codes:
            codes.append(code)
    return ", ".join(codes)


def apply_acne_grid_style(ws, cell_range: str, fill=None, font=None, medium=False):
    border = BORDER_MEDIUM if medium else BORDER_THIN
    for row in ws[cell_range]:
        for cell in row:
            cell.border = border
            if fill is not None:
                cell.fill = fill
            if font is not None:
                cell.font = font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def write_address_block(ws, start_col: int, start_row: int, title: str, lines: List[str]):
    ws.merge_cells(start_row=start_row, start_column=start_col, end_row=start_row, end_column=start_col + 3)
    h = ws.cell(start_row, start_col, title)
    h.font = Font(bold=True, color="FFFFFF")
    h.fill = BLACK_FILL
    h.alignment = Alignment(horizontal="center", vertical="center")
    apply_acne_grid_style(ws, f"{get_column_letter(start_col)}{start_row}:{get_column_letter(start_col+3)}{start_row}", BLACK_FILL, Font(bold=True, color="FFFFFF"), medium=True)
    ws.row_dimensions[start_row].height = 18
    for idx in range(4):
        r = start_row + 2 + idx
        txt = lines[idx] if idx < len(lines) else ""
        ws.merge_cells(start_row=r, start_column=start_col, end_row=r, end_column=start_col + 3)
        c = ws.cell(r, start_col, txt)
        c.font = Font(bold=True, size=9)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True, shrink_to_fit=True)
        ws.row_dimensions[r].height = 21
        apply_acne_grid_style(ws, f"{get_column_letter(start_col)}{r}:{get_column_letter(start_col+3)}{r}", WHITE_FILL, Font(bold=True, size=9), medium=False)


def write_acne_size_scale(ws, top_row: int, active_row: int, active_sizes: List[str], include_categories: bool = True):
    # Header row.
    ws.merge_cells(start_row=top_row, start_column=6, end_row=top_row, end_column=21)
    ws.cell(top_row, 6, "SIZE SCALE")
    ws.cell(top_row, 6).font = Font(bold=True, color="FFFFFF")
    ws.cell(top_row, 6).fill = BLACK_FILL
    ws.cell(top_row, 6).alignment = Alignment(horizontal="center", vertical="center")
    apply_acne_grid_style(ws, f"F{top_row}:U{top_row}", BLACK_FILL, Font(bold=True, color="FFFFFF"), medium=True)
    if include_categories:
        ws.merge_cells(start_row=top_row, start_column=22, end_row=top_row, end_column=23)
        ws.cell(top_row, 22, "PRODUCT CATEGORIES")
        apply_acne_grid_style(ws, f"V{top_row}:W{top_row}", BLACK_FILL, Font(bold=True, color="FFFFFF"), medium=True)

    row_map = {r: (scale[:], category) for r, scale, category in ACNE_SIZE_SCALE_ROWS}
    # If the detected scale is custom, overwrite the selected display row with the actual sizes.
    if active_row in row_map and not set(active_sizes).issubset(set(row_map[active_row][0])):
        row_map[active_row] = (active_sizes[:16], "DYNAMIC SIZE SCALE")

    for r, scale, category in ACNE_SIZE_SCALE_ROWS:
        rr = top_row + (r - 26)
        vals, cat = row_map.get(r, (scale, category))
        for i in range(16):
            c = ws.cell(rr, 6 + i, vals[i] if i < len(vals) else "")
            c.font = Font(bold=True)
            c.alignment = Alignment(horizontal="center", vertical="center")
            c.border = BORDER_THIN
            if r == active_row:
                c.fill = PatternFill("solid", fgColor="FCE4D6")
            else:
                c.fill = WHITE_FILL
        if include_categories:
            ws.merge_cells(start_row=rr, start_column=22, end_row=rr, end_column=23)
            c = ws.cell(rr, 22, cat)
            c.font = Font(bold=True)
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            c.border = BORDER_THIN
            c.fill = WHITE_FILL


def write_carton_table(ws, start_row: int, order: Dict[str, Any], active_sizes: List[str], show_style_once: bool = True) -> int:
    instruction = "CHOOSE APPLICABLE SIZE SCALE FROM ROWS ABOVE (DO NOT MIX ROWS!)"
    headers = {
        2: "Hanging",
        3: "STYLE NAME",
        4: "COLOR NAME",
        5: "COLOR CODE",
        22: "TTL UNITS",
        23: "Boxes",
    }
    for col, text in headers.items():
        ws.cell(start_row, col, text)
        ws.cell(start_row, col).font = Font(bold=True, color="FFFFFF")
        ws.cell(start_row, col).fill = BLACK_FILL
        ws.cell(start_row, col).alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.cell(start_row, col).border = BORDER_MEDIUM
    ws.merge_cells(start_row=start_row, start_column=6, end_row=start_row, end_column=21)
    ws.cell(start_row, 6, instruction)
    apply_acne_grid_style(ws, f"F{start_row}:U{start_row}", BLACK_FILL, Font(bold=True, color="FFFFFF"), medium=True)

    size_to_offset = {str(s): i for i, s in enumerate(active_sizes[:16])}
    row = start_row + 1
    total_units = 0
    product = order.get("product_name", "") or order.get("style_code", "")
    for idx, carton in enumerate(order.get("cartons", [])):
        qty_by_size = {str(s): 0 for s in active_sizes[:16]}
        for sku in carton.get("skus", []):
            s = str(sku.get("size", "")).replace(".0", "")
            qty_by_size[s] = qty_by_size.get(s, 0) + int_qty(sku.get("units", 0))
        row_total = sum(qty_by_size.values())
        total_units += row_total
        values = {
            2: "",
            3: product if (idx == 0 or not show_style_once) else "",
            4: carton.get("color", ""),
            5: carton_color_code(carton),
            22: row_total if row_total else "",
            23: order.get("box_size", "60/40/40"),
        }
        for col, val in values.items():
            c = ws.cell(row, col, val)
            c.border = BORDER_THIN
            c.fill = WHITE_FILL
            c.font = Font(bold=True if col in (3, 4, 5) else False)
            c.alignment = Alignment(horizontal="center" if col != 3 else "left", vertical="center", wrap_text=True)
        for size, qty in qty_by_size.items():
            if size in size_to_offset and qty:
                c = ws.cell(row, 6 + size_to_offset[size], qty)
                c.border = BORDER_THIN
                c.fill = WHITE_FILL
                c.alignment = Alignment(horizontal="center", vertical="center")
                c.font = Font(bold=True)
        # Ensure empty size cells have borders too.
        for col in range(6, 22):
            c = ws.cell(row, col)
            c.border = BORDER_THIN
            if c.value in (None, ""):
                c.fill = WHITE_FILL
        ws.row_dimensions[row].height = 18
        row += 1

    # TTL PO row.
    ttl_row = row + 1
    ws.merge_cells(start_row=ttl_row, start_column=4, end_row=ttl_row, end_column=21)
    ws.cell(ttl_row, 4, "TTL PO:")
    ws.cell(ttl_row, 4).font = Font(bold=True, color="FFFFFF")
    ws.cell(ttl_row, 4).fill = BLACK_FILL
    ws.cell(ttl_row, 4).alignment = Alignment(horizontal="right", vertical="center")
    ws.cell(ttl_row, 22, total_units)
    ws.cell(ttl_row, 22).font = Font(bold=True, color="FFFFFF")
    ws.cell(ttl_row, 22).fill = BLACK_FILL
    ws.cell(ttl_row, 22).alignment = Alignment(horizontal="center", vertical="center")
    apply_acne_grid_style(ws, f"D{ttl_row}:V{ttl_row}", BLACK_FILL, Font(bold=True, color="FFFFFF"), medium=True)
    return ttl_row




def write_acne_studios_logo(ws):
    """Large text logo/header matching the top area of the Acne packing template."""
    # The uploaded Acne template has the logo image anchored across the upper middle
    # of the Packing List sheet.  A clean Excel text version is more reliable than
    # embedding EMF/PNG files on different PCs.
    try:
        ws.merge_cells(start_row=2, start_column=5, end_row=5, end_column=14)  # E2:N5
    except Exception:
        pass
    logo = ws.cell(2, 5, "ACNE STUDIOS")
    logo.font = Font(name="Arial", size=32, bold=True, color="000000")
    logo.alignment = Alignment(horizontal="center", vertical="center")
    for r in range(1, 7):
        ws.row_dimensions[r].height = 20
    ws.row_dimensions[2].height = 28
    ws.row_dimensions[3].height = 24
    ws.row_dimensions[4].height = 24
    ws.row_dimensions[5].height = 20
    # Keep the area clean/white and without grid-style borders.
    for row in ws.iter_rows(min_row=1, max_row=6, min_col=2, max_col=23):
        for cell in row:
            if cell.coordinate != logo.coordinate:
                cell.fill = WHITE_FILL


def estimated_text_lines(text: Any, chars_per_line: int) -> int:
    parts = []
    for line in str(text or "").splitlines() or [""]:
        line = line.strip()
        parts.append(max(1, (len(line) + chars_per_line - 1) // chars_per_line))
    return max(1, sum(parts))


def label_alignment(horizontal: str = "center", shrink: bool = False):
    return Alignment(horizontal=horizontal, vertical="center", wrap_text=True, shrink_to_fit=shrink)

def _clear_range_values(ws, min_row: int, max_row: int, min_col: int, max_col: int):
    for row in ws.iter_rows(min_row=min_row, max_row=max_row, min_col=min_col, max_col=max_col):
        for cell in row:
            cell.value = None


def _safe_set(ws, row: int, col: int, value: Any):
    """Set only real cells; top-left of merged ranges should be targeted by caller."""
    try:
        ws.cell(row, col).value = value
    except Exception:
        pass


def _find_template_row(ws, predicate, start_row: int = 1, end_row: Optional[int] = None) -> int:
    end_row = end_row or ws.max_row
    for r in range(start_row, end_row + 1):
        values = [clean_line(ws.cell(r, c).value) for c in range(1, ws.max_column + 1)]
        if predicate(values):
            return r
    return -1


def _template_layout(ws) -> Dict[str, int]:
    main_header = _find_template_row(ws, lambda v: "CARTON #" in v and "STYLE NAME" in v, 1, 80)
    if main_header == -1:
        main_header = 38
    size_headers = []
    for r in range(1, ws.max_row + 1):
        if clean_line(ws.cell(r, 6).value).upper() == "SIZE SCALE":
            size_headers.append(r)
    top_size_header = size_headers[0] if size_headers else 26
    bottom_size_header = next((r for r in size_headers if r > main_header + 20), 151)
    bottom_header = _find_template_row(
        ws,
        lambda v: "COLOR NAME" in v and any("CHOOSE APPLICABLE" in x for x in v),
        bottom_size_header + 1,
        min(ws.max_row, bottom_size_header + 40),
    )
    if bottom_header == -1:
        bottom_header = bottom_size_header + 12
    bottom_total = _find_template_row(
        ws,
        lambda v: any("TTL PO" in x for x in v),
        bottom_header + 1,
        min(ws.max_row, bottom_header + 40),
    )
    if bottom_total == -1:
        bottom_total = bottom_header + 11
    return {
        "top_size_header": top_size_header,
        "main_header": main_header,
        "main_data_start": main_header + 1,
        "main_data_end": bottom_size_header - 1,
        "bottom_size_header": bottom_size_header,
        "bottom_header": bottom_header,
        "bottom_data_start": bottom_header + 1,
        "bottom_data_end": bottom_total - 1,
        "bottom_total": bottom_total,
    }


def _is_merged_child(cell) -> bool:
    return cell.__class__.__name__ == "MergedCell"


def _clear_real_cells(ws, min_row: int, max_row: int, min_col: int, max_col: int):
    for row in ws.iter_rows(min_row=min_row, max_row=max_row, min_col=min_col, max_col=max_col):
        for cell in row:
            if not _is_merged_child(cell):
                cell.value = None


def _write_template_size_scale_selection(ws, active_row_template: int, active_sizes: List[str]):
    """Keep template scale rows, but highlight/fill the applicable size row.

    The destination templates have different lower block positions, so the rows are
    detected instead of hardcoded.
    """
    layout = _template_layout(ws)
    red_fill = PatternFill("solid", fgColor="FF0000")
    white_fill = WHITE_FILL
    active_offset = max(0, active_row_template - 27)

    for header_row, max_col in ((layout["top_size_header"], 23), (layout["bottom_size_header"], 21)):
        for rr in range(header_row + 1, header_row + 12):
            for c in range(6, max_col + 1):
                if not _is_merged_child(ws.cell(rr, c)):
                    ws.cell(rr, c).fill = white_fill
        active_excel_row = header_row + 1 + active_offset
        for c in range(6, max_col + 1):
            if not _is_merged_child(ws.cell(active_excel_row, c)):
                ws.cell(active_excel_row, c).fill = red_fill

    # If a dynamic/custom scale is needed, overwrite the active row values.
    standard = []
    for row_no, scale, _cat in ACNE_SIZE_SCALE_ROWS:
        if row_no == active_row_template:
            standard = [str(x) for x in scale]
            break
    if active_sizes and not set(map(str, active_sizes)).issubset(set(standard)):
        for header_row in (layout["top_size_header"], layout["bottom_size_header"]):
            r = header_row + 1 + active_offset
            for i in range(16):
                cell = ws.cell(r, 6 + i)
                if not _is_merged_child(cell):
                    cell.value = active_sizes[i] if i < len(active_sizes) else ""


def _template_write_cartons(ws, order: Dict[str, Any], active_sizes: List[str]) -> int:
    """Populate the Acne template's fixed carton blocks while preserving its formatting."""
    layout = _template_layout(ws)
    size_to_offset = {str(s).replace(".0", ""): i for i, s in enumerate(active_sizes[:16])}
    cartons = order.get("cartons", [])
    product = order.get("product_name", "") or order.get("style_code", "")

    # Main carton block.
    _clear_real_cells(ws, layout["main_data_start"], layout["main_data_end"], 2, 23)
    for r in range(layout["main_data_start"], layout["main_data_end"] + 1):
        ws.cell(r, 22).value = f"=SUM(F{r}:U{r})"
        ws.cell(r, 22).number_format = "0;-0;;@"

    max_main_rows = max(0, layout["main_data_end"] - layout["main_data_start"] + 1)
    for idx, carton in enumerate(cartons[:max_main_rows]):
        r = layout["main_data_start"] + idx
        ws.cell(r, 2).value = carton.get("carton_no", idx + 1)
        ws.cell(r, 3).value = product if idx == 0 else ""
        ws.cell(r, 4).value = carton.get("color", "")
        ws.cell(r, 5).value = carton_color_code(carton)
        ws.cell(r, 23).value = order.get("box_size", "60/40/40")
        for sku in carton.get("skus", []):
            sz = str(sku.get("size", "")).replace(".0", "")
            q = int_qty(sku.get("units", 0))
            if sz in size_to_offset:
                cell = ws.cell(r, 6 + size_to_offset[sz])
                if not _is_merged_child(cell):
                    cell.value = (cell.value or 0) + q

    # Bottom compact block.
    _clear_real_cells(ws, layout["bottom_data_start"], layout["bottom_data_end"], 4, 22)
    for r in range(layout["bottom_data_start"], layout["bottom_data_end"] + 1):
        ws.cell(r, 22).value = f"=SUM(F{r}:U{r})"
        ws.cell(r, 22).number_format = "0;-0;;@"

    max_bottom_rows = max(0, layout["bottom_data_end"] - layout["bottom_data_start"] + 1)
    for idx, carton in enumerate(cartons[:max_bottom_rows]):
        r = layout["bottom_data_start"] + idx
        ws.cell(r, 4).value = carton.get("color", "")
        ws.cell(r, 5).value = carton_color_code(carton)
        for sku in carton.get("skus", []):
            sz = str(sku.get("size", "")).replace(".0", "")
            q = int_qty(sku.get("units", 0))
            if sz in size_to_offset:
                cell = ws.cell(r, 6 + size_to_offset[sz])
                if not _is_merged_child(cell):
                    cell.value = (cell.value or 0) + q

    ws.cell(layout["bottom_total"], 4).value = "TTL PO:"
    ws.cell(layout["bottom_total"], 22).value = f"=SUM(V{layout['bottom_data_start']}:V{layout['bottom_data_end']})"
    ws.cell(layout["bottom_total"], 22).number_format = "0;-0;;@"
    return layout["bottom_total"]


def _recreate_summary_from_template(wb: Workbook, order: Dict[str, Any], active_sizes: List[str]):
    if "summery" not in wb.sheetnames:
        ws = wb.create_sheet("summery")
    else:
        ws = wb["summery"]
    ws.freeze_panes = None
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].hidden = True
    ws.column_dimensions["A"].width = 0.1
    try:
        # Unmerge before clearing because merged child cells are read-only in openpyxl.
        for rng in list(ws.merged_cells.ranges):
            ws.unmerge_cells(str(rng))
    except Exception:
        pass
    # Preserve the template's compact area and rewrite the values.
    for row in ws.iter_rows(min_row=1, max_row=20, min_col=2, max_col=20):
        for cell in row:
            cell.value = None
    last_col = 5 + len(active_sizes[:16])
    if last_col < 11:
        last_col = 11
    ws.merge_cells(start_row=1, start_column=2, end_row=1, end_column=last_col)
    ws.merge_cells(start_row=11, start_column=2, end_row=11, end_column=3)
    style_name_short = order.get("style_code", "") or order.get("po_number", "")
    ws.cell(1, 2).value = f"NO ORDER/ N поръчка: {style_name_short}"
    apply_acne_grid_style(ws, f"B1:{get_column_letter(last_col)}1", BLACK_FILL, Font(bold=True, color="FFFFFF"), medium=True)
    headers = ["STYLE CODE", "COLOR / ЦВЯТ", "DESTINATION", *active_sizes[:16], "TOTAL"]
    for idx, h in enumerate(headers, start=2):
        c = ws.cell(2, idx, h)
        c.value = h
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = BLACK_FILL
        c.border = BORDER_MEDIUM
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    groups: Dict[tuple, Dict[str, Any]] = {}
    for carton in order.get("cartons", []):
        key = (carton.get("color", ""), carton_color_code(carton))
        groups.setdefault(key, {"sizes": {s: 0 for s in active_sizes[:16]}, "total": 0})
        for sku in carton.get("skus", []):
            sz = str(sku.get("size", "")).replace(".0", "")
            q = int_qty(sku.get("units", 0))
            if sz in groups[key]["sizes"]:
                groups[key]["sizes"][sz] += q
            groups[key]["total"] += q
    dest = destination_summary(order)
    row = 3
    for (color, code), info in groups.items():
        ws.cell(row, 2).value = order.get("style_code", "")
        ws.cell(row, 3).value = color
        ws.cell(row, 4).value = dest
        for i, sz in enumerate(active_sizes[:16], start=5):
            ws.cell(row, i).value = info["sizes"].get(sz, 0) or ""
        ws.cell(row, last_col).value = f"=SUM(E{row}:{get_column_letter(last_col-1)}{row})"
        apply_acne_grid_style(ws, f"B{row}:{get_column_letter(last_col)}{row}", WHITE_FILL, Font(bold=False), medium=False)
        row += 1
    total_row = max(row + 1, 11)
    ws.cell(total_row, 2).value = "TOTAL per size \nВСИЧКО по размер"
    ws.cell(total_row, 4).value = dest
    for i, sz in enumerate(active_sizes[:16], start=5):
        ws.cell(total_row, i).value = f"=SUM({get_column_letter(i)}3:{get_column_letter(i)}{row-1})"
    ws.cell(total_row, last_col).value = f"=SUM(E{total_row}:{get_column_letter(last_col-1)}{total_row})"
    apply_acne_grid_style(ws, f"B{total_row}:{get_column_letter(last_col)}{total_row}", LIGHT_GREEN_FILL, Font(bold=True), medium=True)
    ws.print_area = f"B1:{get_column_letter(last_col)}{total_row}"


def _set_template_delivery_row(ws, row_no: int, text: str):
    """Write a delivery-address line into the row's merged J:O region.

    Some Acne templates use J-row merges, while the KR sample has K11:N11.
    This helper writes only to the top-left cell of the merge so openpyxl does
    not fail on merged child cells.
    """
    target_col = 10
    for rng in ws.merged_cells.ranges:
        if rng.min_row == row_no and rng.max_row == row_no and rng.max_col >= 10 and rng.min_col <= 15:
            target_col = rng.min_col
            break
    cell = ws.cell(row_no, target_col)
    cell.value = text
    cell.font = Font(bold=False, size=10)
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True, shrink_to_fit=True)


def _write_template_delivery_address(ws, order: Dict[str, Any]):
    key = destination_key_from_order(order)
    preset_rows = DESTINATION_PRESETS[key]["packing_rows"]
    manual_lines = [clean_line(x) for x in str(order.get("ship_to", "")).splitlines() if clean_line(x)]

    # Clear the writable merged-cell anchors in the whole delivery-address area.
    for row_no in range(10, 17):
        _set_template_delivery_row(ws, row_no, "")

    # If the current address exactly matches the preset, keep the template-specific
    # row placement.  If the user edited the address manually, place it from row 10
    # downward so no text is silently dropped.
    preset_address = default_destination_address(key)
    if clean_line(str(order.get("ship_to", "")).replace("\n", " ")) == clean_line(preset_address.replace("\n", " ")):
        for row_no, text in preset_rows.items():
            _set_template_delivery_row(ws, row_no, text)
    else:
        for row_no, text in zip(range(10, 17), manual_lines[:7]):
            _set_template_delivery_row(ws, row_no, text)


def create_packing_workbook(order: Dict[str, Any]) -> bytes:
    """Create a Packing List by cloning the uploaded Acne template and filling data.

    This keeps the original column widths, row heights, borders, red size-scale row,
    formulas, lower summary block, and overall print geometry much closer than
    rebuilding the workbook from scratch.
    """
    order = recalc_order(order)
    template_path = PACKING_TEMPLATES.get(destination_key_from_order(order), PACKING_TEMPLATE_XLSX)
    if not template_path.exists():
        template_path = PACKING_TEMPLATE_XLSX
    if template_path.exists():
        wb = load_workbook(template_path)
        ws = wb["Packing List"] if "Packing List" in wb.sheetnames else wb.active
    else:
        # Fallback for users who remove the template file.
        wb = Workbook()
        ws = wb.active
        ws.title = "Packing List"
    setup_print(ws, "landscape")
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = None
    ws.column_dimensions["A"].hidden = True
    ws.column_dimensions["A"].width = 0.1

    sizes = sorted({str(i.get("size", "")).replace(".0", "") for i in order.get("items", []) if str(i.get("size", "")).strip()}, key=natural_key)
    active_row, active_sizes, _product_category = choose_size_scale(sizes)
    total_cartons = max(1, len(order.get("cartons", [])))
    current_date = datetime.now().date()

    # Restore a clean text logo in the same blank top area used by the template image.
    # The original template contains a WMF image; openpyxl cannot preserve WMF reliably,
    # so we replace it with a centered Excel text logo that survives saving.
    try:
        for rng in list(ws.merged_cells.ranges):
            if str(rng) in ("B5:V5", "B2:V6"):
                ws.unmerge_cells(str(rng))
        ws.merge_cells("B2:V6")
    except Exception:
        pass
    ws.cell(2, 2).value = "ACNE STUDIOS"
    ws.cell(2, 2).font = Font(name="Arial", size=30, bold=True, color="000000")
    ws.cell(2, 2).alignment = Alignment(horizontal="center", vertical="center")
    for rr in range(2, 7):
        ws.row_dimensions[rr].height = 18

    ws.cell(8, 2).value = "PACKING LIST"
    ws.cell(8, 3).value = datetime.now().strftime("%Y%m%d") + "-1"
    ws.cell(10, 3).value = order.get("po_number", "")
    ws.cell(11, 3).value = order.get("product_name", "") or order.get("style_code", "")
    ws.cell(12, 3).value = order.get("vendor_invoice", "")
    ws.cell(14, 3).value = total_cartons
    ws.cell(15, 3).value = "=(0.6*0.4*0.4)*C14"
    ws.cell(16, 3).value = "=0.32*V166"
    ws.cell(17, 3).value = "=0.32*V166"
    # Keep the template's original formula logic, but make the packing surcharge follow carton count.
    ws.cell(18, 3).value = f"=C17+{round(total_cartons * 0.56, 2)}"
    ws.cell(20, 3).value = current_date
    ws.cell(20, 3).number_format = "yyyy-mm-dd"
    ws.cell(21, 3).value = order.get("ex_factory_date", "")
    ws.cell(23, 3).value = order.get("awb", "")
    ws.cell(24, 3).value = order.get("tracking", "")

    shipper_lines = ["PELINTEX BULGARIA LTD", "5 HRISTO BOTEV BLVD", "RUSE 7000", "BULGARIA"]
    for idx, txt in enumerate(shipper_lines[:4], start=11):
        ws.cell(idx, 5).value = txt
    _write_template_delivery_address(ws, order)

    # Partial Shipment block: match the visual template supplied by the user.
    # Keep it large/readable and avoid Excel shrink-to-fit.  The red X is rich text
    # inside the same cell so the original template geometry stays intact.
    ws.cell(13, 22).value = "Partial Shipment:  ☑"
    try:
        ws.cell(13, 23).value = CellRichText(
            TextBlock(InlineFont(rFont="Arial", sz=20, color="000000"), "Partial Shipment #: "),
            TextBlock(InlineFont(rFont="Arial", sz=20, color="FF0000"), "X"),
        )
    except Exception:
        ws.cell(13, 23).value = "Partial Shipment #: X"
        ws.cell(13, 23).font = Font(name="Arial", size=20, color="000000")
    ws.cell(13, 22).font = Font(name="Arial", size=20, color="000000")
    ws.cell(13, 23).font = Font(name="Arial", size=20, color="000000")
    ws.cell(13, 22).alignment = Alignment(horizontal="right", vertical="center", wrap_text=False, shrink_to_fit=False)
    ws.cell(13, 23).alignment = Alignment(horizontal="left", vertical="center", wrap_text=False, shrink_to_fit=False)
    ws.row_dimensions[13].height = 31

    _write_template_size_scale_selection(ws, active_row, active_sizes)
    bottom_total_row = _template_write_cartons(ws, order, active_sizes)

    # Keep the sample-like fixed print range.  It includes the logo/header, the large upper carton grid,
    # the lower compact grid, and the TTL row just like the provided workbook.
    ws.print_area = f"B1:W{bottom_total_row}"
    ws.freeze_panes = None

    _recreate_summary_from_template(wb, order, active_sizes)
    for sh in wb.worksheets:
        sh.freeze_panes = None

    bio = io.BytesIO()
    wb.save(bio)
    return bio.getvalue()


def create_summary_sheet(wb: Workbook, order: Dict[str, Any], active_sizes: List[str]):
    ws = wb.create_sheet("summery")
    setup_print(ws, "landscape")
    ws.page_setup.fitToHeight = 1
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = None
    ws.column_dimensions["A"].hidden = True
    ws.column_dimensions["A"].width = 0.1
    destination = destination_summary(order)
    for col, width in {"B": 22, "C": 28, "D": 28, "K": 12}.items():
        ws.column_dimensions[col].width = width
    for col in range(5, 5 + len(active_sizes[:16]) + 1):
        ws.column_dimensions[get_column_letter(col)].width = 10
    last_col = 5 + len(active_sizes[:16])

    ws.merge_cells(start_row=1, start_column=2, end_row=1, end_column=last_col)
    title = f"NO ORDER/ N поръчка: {order.get('po_number','')}"
    ws.cell(1, 2, title)
    apply_acne_grid_style(ws, f"B1:{get_column_letter(last_col)}1", BLACK_FILL, Font(bold=True, color="FFFFFF"), medium=True)

    headers = ["STYLE CODE", "COLOR / ЦВЯТ", "DESTINATION", *active_sizes[:16], "TOTAL"]
    ws.get_active_cell = None
    for idx, h in enumerate(headers, start=2):
        c = ws.cell(2, idx, h)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = BLACK_FILL
        c.border = BORDER_MEDIUM
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    # Group by color/color code to make the sheet closer to the uploaded summary.
    groups: Dict[tuple, Dict[str, Any]] = {}
    for carton in order.get("cartons", []):
        key = (carton.get("color", ""), carton_color_code(carton))
        if key not in groups:
            groups[key] = {"sizes": {s: 0 for s in active_sizes[:16]}, "total": 0}
        for sku in carton.get("skus", []):
            sz = str(sku.get("size", "")).replace(".0", "")
            q = int_qty(sku.get("units", 0))
            if sz in groups[key]["sizes"]:
                groups[key]["sizes"][sz] += q
            groups[key]["total"] += q
    row = 3
    style_code = order.get("style_code", "")
    for (color, code), info in groups.items():
        ws.cell(row, 2, style_code)
        ws.cell(row, 3, color)
        ws.cell(row, 4, destination)
        for i, sz in enumerate(active_sizes[:16], start=5):
            ws.cell(row, i, info["sizes"].get(sz, 0) or "")
        ws.cell(row, last_col, info["total"])
        apply_acne_grid_style(ws, f"B{row}:{get_column_letter(last_col)}{row}", WHITE_FILL, Font(bold=False), medium=False)
        row += 1
    total_row = row + 1
    ws.merge_cells(start_row=total_row, start_column=2, end_row=total_row, end_column=4)
    ws.cell(total_row, 2, "TOTAL per size \nВСИЧКО по размер")
    for i, sz in enumerate(active_sizes[:16], start=5):
        total = sum(info["sizes"].get(sz, 0) for info in groups.values())
        ws.cell(total_row, i, total or "")
    ws.cell(total_row, last_col, sum(info["total"] for info in groups.values()))
    apply_acne_grid_style(ws, f"B{total_row}:{get_column_letter(last_col)}{total_row}", LIGHT_GREEN_FILL, Font(bold=True), medium=True)
    ws.print_area = f"B1:{get_column_letter(last_col)}{total_row}"

def label_fields(order: Dict[str, Any], carton: Dict[str, Any], sku: Dict[str, Any], total_cartons: int) -> List[tuple]:
    return [
        ("Ship To:", order.get("ship_to", "")),
        ("PO #:", order.get("po_number", "")),
        ("Style Name:", order.get("product_name", "")),
        ("Color Name:", sku.get("color") or carton.get("color", "")),
        ("Size:", str(sku.get("size", "")).replace(".0", "")),
        ("Style #:", order.get("style_code", sku.get("article_name", ""))),
        ("Number of Units:", int_qty(sku.get("units", 0))),
        ("Carton #:", f"{carton.get('carton_no', '')} of {total_cartons}"),
        ("Shipment #:", "1 of 1"),
        ("Shipment Date:", order.get("shipment_date") or order.get("ex_factory_date", "")),
        ("Country of Origin:", order.get("origin", "BULGARIA")),
    ]


def multi_label_fields(order: Dict[str, Any], carton: Dict[str, Any], sku: Dict[str, Any], total_cartons: int, block_index: int) -> List[tuple]:
    """Multi-label layout matching the original ACNE HTML tool.

    In the 2x2 multi-SKU sheet:
    - first label gets Ship To address as the top tall row;
    - second label gets Shipment Date as the top tall row;
    - third/fourth labels start directly with PO #.
    This keeps the old carton-sticker behavior and avoids repeating the address
    on every multi label.
    """
    shipment_date = order.get("shipment_date") or order.get("ex_factory_date", "")
    common = [
        ("PO #:", order.get("po_number", "")),
        ("Style Name:", order.get("product_name", "")),
        ("Color Name:", sku.get("color") or carton.get("color", "")),
        ("Size:", str(sku.get("size", "")).replace(".0", "")),
        ("Style #:", order.get("style_code", sku.get("article_name", ""))),
        ("Number of Units:", int_qty(sku.get("units", 0))),
        ("Carton #:", f"{carton.get('carton_no', '')} of {total_cartons}"),
        ("Shipment #:", "1 of 1"),
        ("Country of Origin:", order.get("origin", "BULGARIA")),
    ]
    if block_index == 0:
        return [("Ship To:", order.get("ship_to", ""))] + common
    if block_index == 1:
        return [("Shipment Date:", shipment_date)] + common
    return common

def multi_label_group_key(order: Dict[str, Any], carton: Dict[str, Any], sku: Dict[str, Any]) -> str:
    """Group multi-SKU labels by Acne item family / color code, not by size SKU.

    Acne PDF item numbers often include the size at the end:
    AF0617-EA1032, AF0617-EA1034, AF0617-EA1036.
    Those must print together on one landscape label sheet because they share
    the same item family/color code EA10. The same applies to J83032/34/36.
    """
    color_code = clean_line(sku.get("color_code") or infer_color_code(sku.get("item_number", ""), sku.get("size", "")))
    if color_code:
        return color_code

    item_number = clean_line(sku.get("item_number") or carton.get("item_numbers") or "")
    size = clean_line(str(sku.get("size", "")).replace(".0", ""))
    if item_number and size and item_number.upper().endswith(size.upper()):
        family = item_number[:-len(size)].rstrip("-_/ ")
        if family:
            return family

    return item_number or clean_line(sku.get("color") or carton.get("color") or order.get("style_code") or "No Item")


def draw_single_label(ws, order: Dict[str, Any], carton: Dict[str, Any], sku: Dict[str, Any], total_cartons: int):
    setup_print(ws, "portrait")
    ws.page_margins = PageMargins(left=0.18, right=0.18, top=0.22, bottom=0.22, header=0.05, footer=0.05)
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 58
    row = 2
    for label, value in label_fields(order, carton, sku, total_cartons):
        is_ship_to = label == "Ship To:"
        value_font_size = 12 if is_ship_to else 15
        label_font_size = 15 if is_ship_to else 16
        ws.cell(row, 1, label)
        ws.cell(row, 2, value)
        ws.cell(row, 1).font = Font(bold=True, size=label_font_size)
        ws.cell(row, 2).font = Font(bold=True, size=value_font_size)
        ws.cell(row, 1).alignment = label_alignment("right", shrink=True)
        ws.cell(row, 2).alignment = label_alignment("center", shrink=True)
        ws.cell(row, 1).border = BORDER_MEDIUM
        ws.cell(row, 2).border = BORDER_MEDIUM
        if is_ship_to:
            # Extra height prevents multi-line delivery addresses from being cut when printed.
            ws.row_dimensions[row].height = max(78, min(118, estimated_text_lines(value, 30) * 18))
        else:
            ws.row_dimensions[row].height = 39
        row += 1
    ws.merge_cells(start_row=row+1, start_column=1, end_row=row+2, end_column=2)
    bc = ws.cell(row+1, 1, "PLACE THE BARCODE STICKER HERE")
    bc.font = Font(color="FF0000", bold=True, size=14)
    bc.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[row+1].height = 28
    ws.row_dimensions[row+2].height = 28
    for r in range(row+1, row+3):
        for c in range(1, 3):
            ws.cell(r, c).border = Border(left=MEDIUM, right=MEDIUM, top=MEDIUM, bottom=MEDIUM)
    ws.print_area = f"A1:B{row+3}"

def draw_multi_block(ws, start_row: int, start_col: int, order: Dict[str, Any], carton: Dict[str, Any], sku: Dict[str, Any], total_cartons: int, block_index: int):
    # Each block spans 2 columns x 11 rows. V10 uses more of the A4 landscape page
    # while leaving a real printable top margin so the upper border is not clipped
    # by physical printer margins.
    label_col = start_col
    value_col = start_col + 1
    ws.column_dimensions[get_column_letter(label_col)].width = 16.6
    ws.column_dimensions[get_column_letter(value_col)].width = 41.8
    fields = multi_label_fields(order, carton, sku, total_cartons, block_index)
    for idx, (label, value) in enumerate(fields):
        r = start_row + idx
        is_top_special = label in ("Ship To:", "Shipment Date:")
        is_ship_to = label == "Ship To:"
        ws.cell(r, label_col, label)
        ws.cell(r, value_col, value)
        ws.cell(r, label_col).font = Font(bold=True, size=9.8 if is_top_special else 10.4)
        ws.cell(r, value_col).font = Font(bold=True, size=8.4 if is_ship_to else (10.2 if label == "Shipment Date:" else 9.8))
        ws.cell(r, label_col).alignment = label_alignment("right", shrink=True)
        ws.cell(r, value_col).alignment = label_alignment("center", shrink=True)
        ws.cell(r, label_col).border = BORDER_MEDIUM
        ws.cell(r, value_col).border = BORDER_MEDIUM
        if is_ship_to:
            needed = estimated_text_lines(value, 34) * 11.2
            ws.row_dimensions[r].height = max(ws.row_dimensions[r].height or 0, max(68, min(80, needed)))
        elif label == "Shipment Date:":
            # The second top block on multi SKU labels shows shipment date instead of address.
            ws.row_dimensions[r].height = max(ws.row_dimensions[r].height or 0, 68)
        else:
            ws.row_dimensions[r].height = max(ws.row_dimensions[r].height or 0, 19.6)


def safe_sheet_name(base: str, existing: set) -> str:
    """Excel-safe unique sheet name, max 31 chars."""
    base = re.sub(r"[\\/*?:\[\]]+", "_", str(base or "Sheet")).strip()
    base = re.sub(r"\s+", " ", base) or "Sheet"
    base = base[:31]
    name = base
    n = 2
    while name in existing:
        suffix = f" {n}"
        name = base[:31 - len(suffix)] + suffix
        n += 1
    existing.add(name)
    return name

def create_labels_workbook(order: Dict[str, Any]) -> bytes:
    """Create printable label workbook using the same single/multi logic as the original HTML tool.

    Original ACNE logic:
    - evaluate each carton independently;
    - if that carton has exactly 1 SKU, export it as a Single label sheet;
    - if that carton has 2+ SKUs, export those SKUs as Multi labels for that carton;
    - multi cartons are split into A4 landscape sheets in chunks of max 4 labels.

    This intentionally avoids cross-carton regrouping during label export.  The draft
    carton builder already groups PDF order lines into carton-like rows (default: one
    carton per color), so two color/item families with three sizes each become two
    landscape sheets with three labels each, matching the older Excel/HTML behavior.
    """
    order = recalc_order(order)
    wb = Workbook()
    wb.remove(wb.active)
    total_cartons = max(1, len(order.get("cartons", [])))
    existing_names = set()

    single_index = 1
    multi_index = 1

    for carton in order.get("cartons", []):
        skus = carton.get("skus", [])
        if not skus:
            continue

        # Old logic: a carton with one SKU is Single, no matter which item/color it is.
        if len(skus) == 1:
            sku = skus[0]
            sheet_name = safe_sheet_name(f"Single {single_index:03d}", existing_names)
            ws = wb.create_sheet(sheet_name)
            draw_single_label(ws, order, carton, sku, total_cartons)
            ws.freeze_panes = None
            single_index += 1
            continue

        # Old logic: a carton with multiple SKUs becomes Multi, split only when it
        # exceeds the 4-label landscape grid.  Do not mix labels from another carton.
        carton_no = clean_line(str(carton.get("carton_no", multi_index)))
        code = clean_line(carton_color_code(carton) or carton.get("color") or "")
        short_code = re.sub(r"[^A-Za-z0-9]+", "_", code).strip("_")
        base_name = f"Multi CTN{carton_no}" + (f" {short_code}" if short_code else "")
        page_no = 1

        for chunk_start in range(0, len(skus), 4):
            ws = wb.create_sheet(safe_sheet_name(f"{base_name} {page_no:02d}", existing_names))
            setup_print(ws, "landscape")
            ws.freeze_panes = None
            # Keep V10 label sizing/print margins: 2x2 grid, larger labels, safe top margin.
            ws.page_margins = PageMargins(left=0.18, right=0.18, top=0.30, bottom=0.20, header=0.04, footer=0.04)
            ws.sheet_view.showGridLines = False

            chunk = skus[chunk_start:chunk_start + 4]
            positions = [(2, 1), (2, 4), (16, 1), (16, 4)]
            for idx, sku in enumerate(chunk):
                r, c = positions[idx]
                draw_multi_block(ws, r, c, order, carton, sku, total_cartons, idx)

            ws.row_dimensions[1].height = 8
            ws.column_dimensions["C"].width = 3.0
            ws.column_dimensions["F"].width = 1.0
            for gap_row in (13, 14, 15):
                ws.row_dimensions[gap_row].height = 6
            for r in range(1, 28):
                if ws.row_dimensions[r].height is None:
                    ws.row_dimensions[r].height = 19.6
            ws.print_area = "A1:E27"
            page_no += 1
            multi_index += 1

    if not wb.sheetnames:
        ws = wb.create_sheet("No Labels")
        ws.freeze_panes = None
        ws["A1"] = "No label data found."

    for ws in wb.worksheets:
        ws.freeze_panes = None

    bio = io.BytesIO()
    wb.save(bio)
    return bio.getvalue()

def make_zip(files: List[tuple]) -> bytes:
    """Create a ZIP and guarantee unique internal filenames.

    If two generated orders accidentally produce the same base name, Windows
    unzip/copy operations would ask to overwrite.  We guard against that here
    by appending _02, _03, ... to duplicate archive names.
    """
    bio = io.BytesIO()
    used = set()
    with zipfile.ZipFile(bio, "w", zipfile.ZIP_DEFLATED) as zf:
        for name, data in files:
            arcname = str(name).replace("\\", "/")
            if arcname in used:
                folder, fname = arcname.rsplit("/", 1) if "/" in arcname else ("", arcname)
                stem, ext = os.path.splitext(fname)
                n = 2
                candidate = arcname
                while candidate in used:
                    new_fname = f"{stem}_{n:02d}{ext}"
                    candidate = f"{folder}/{new_fname}" if folder else new_fname
                    n += 1
                arcname = candidate
            used.add(arcname)
            zf.writestr(arcname, data)
    return bio.getvalue()



# -----------------------------
# Minimal local web server (no Flask dependency)
# -----------------------------
import mimetypes
import webbrowser
import threading
from email.parser import BytesParser
from email.policy import default as email_default_policy
from http.server import ThreadingHTTPServer, BaseHTTPRequestHandler
from urllib.parse import urlparse


def secure_filename_local(filename: str) -> str:
    filename = os.path.basename(filename or "upload.pdf")
    filename = re.sub(r"[^A-Za-z0-9._ -]+", "_", filename).strip()
    return filename or "upload.pdf"


def response_bytes(data: bytes, status: int = 200, content_type: str = "application/octet-stream", headers: Optional[Dict[str, str]] = None):
    return status, data, content_type, headers or {}


class Handler(BaseHTTPRequestHandler):
    server_version = "ACNEPDFApp/1.1"

    def log_message(self, fmt, *args):
        print("[%s] %s" % (self.log_date_time_string(), fmt % args))

    def send_payload(self, status: int, data: bytes, content_type: str, extra_headers: Optional[Dict[str, str]] = None):
        self.send_response(status)
        self.send_header("Content-Type", content_type)
        self.send_header("Content-Length", str(len(data)))
        self.send_header("Cache-Control", "no-store")
        for key, value in (extra_headers or {}).items():
            self.send_header(key, value)
        self.end_headers()
        self.wfile.write(data)

    def send_json(self, obj: Any, status: int = 200):
        data = json.dumps(obj, ensure_ascii=False).encode("utf-8")
        self.send_payload(status, data, "application/json; charset=utf-8")

    def read_json(self) -> Any:
        length = int(self.headers.get("Content-Length", "0"))
        raw = self.rfile.read(length)
        return json.loads(raw.decode("utf-8") or "{}")

    def do_GET(self):
        parsed = urlparse(self.path)
        if parsed.path == "/" or parsed.path == "/index.html":
            data = (APP_DIR / "templates" / "index.html").read_bytes()
            self.send_payload(200, data, "text/html; charset=utf-8")
            return
        if parsed.path == "/health":
            self.send_json({"ok": True, "ocr_available_python_libs": OCR_AVAILABLE})
            return
        self.send_json({"error": "Not found"}, status=404)

    def do_POST(self):
        parsed = urlparse(self.path)
        try:
            if parsed.path in ("/api/shutdown", "/api/window-closed"):
                self.send_json({"ok": True, "message": "App closing"})
                threading.Timer(0.25, lambda: os._exit(0)).start()
                return
            if parsed.path == "/api/upload":
                self.handle_upload()
                return
            payload = self.read_json()
            if parsed.path == "/api/recalc":
                self.send_json(recalc_order(payload))
                return
            if parsed.path == "/api/export/packing":
                order = recalc_order(payload)
                data = create_packing_workbook(order)
                name = packing_filename(order)
                self.send_payload(200, data, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", {"Content-Disposition": f'attachment; filename="{name}"'})
                return
            if parsed.path == "/api/export/labels":
                order = recalc_order(payload)
                data = create_labels_workbook(order)
                name = labels_filename(order)
                self.send_payload(200, data, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", {"Content-Disposition": f'attachment; filename="{name}"'})
                return
            if parsed.path == "/api/export/both":
                order = recalc_order(payload)
                base = order_export_stem(order)
                zdata = make_zip([
                    (packing_filename(order), create_packing_workbook(order)),
                    (labels_filename(order), create_labels_workbook(order)),
                ])
                self.send_payload(200, zdata, "application/zip", {"Content-Disposition": f'attachment; filename="{zip_filename(order)}"'})
                return
            if parsed.path == "/api/export/all":
                orders = [recalc_order(o) for o in payload.get("orders", [])]
                files = []
                for idx, order in enumerate(orders, 1):
                    base = order_export_stem(order, idx)
                    # Include the file type and the unique order stem in each filename.
                    # Even if the user copies all files out of the ZIP folders into one
                    # Windows folder, packing and labels will not overwrite each other.
                    files.append((f"{base}/{packing_filename(order, idx)}", create_packing_workbook(order)))
                    files.append((f"{base}/{labels_filename(order, idx)}", create_labels_workbook(order)))
                zdata = make_zip(files)
                self.send_payload(200, zdata, "application/zip", {"Content-Disposition": 'attachment; filename="ACNE_ALL_PACKING_LABELS.zip"'})
                return
            self.send_json({"error": "Not found"}, status=404)
        except Exception as exc:
            import traceback
            traceback.print_exc()
            self.send_json({"error": str(exc)}, status=500)

    def handle_upload(self):
        content_type = self.headers.get("Content-Type", "")
        if not content_type.lower().startswith("multipart/form-data"):
            self.send_json({"orders": [], "error": "Expected multipart/form-data"}, status=400)
            return

        length = int(self.headers.get("Content-Length", "0"))
        body = self.rfile.read(length)
        raw_message = (
            f"Content-Type: {content_type}\r\n"
            "MIME-Version: 1.0\r\n\r\n"
        ).encode("utf-8") + body
        msg = BytesParser(policy=email_default_policy).parsebytes(raw_message)

        orders = []
        for part in msg.iter_parts():
            disposition = part.get("Content-Disposition", "")
            if "form-data" not in disposition:
                continue
            name = part.get_param("name", header="content-disposition")
            filename = part.get_filename()
            if name != "files" or not filename:
                continue
            ext = Path(filename).suffix.lower()
            if ext not in (".pdf", ".xlsx", ".xlsm", ".xls"):
                continue
            safe = secure_filename_local(filename)
            temp_path = UPLOAD_DIR / f"{datetime.now().strftime('%Y%m%d%H%M%S%f')}_{safe}"
            data = part.get_payload(decode=True) or b""
            temp_path.write_bytes(data)
            try:
                if ext == ".pdf":
                    order = parse_pdf_order(temp_path)
                    order["source_filename"] = filename
                elif ext in (".xlsx", ".xlsm"):
                    order = parse_packing_workbook(temp_path, filename)
                else:
                    raise ValueError("Legacy .xls files are not supported by the local parser. Please save the packing list as .xlsx and upload it again.")
                orders.append(order)
            except Exception as exc:
                orders.append({
                    "id": safe_name(Path(filename).stem),
                    "source_filename": filename,
                    "po_number": "",
                    "style_code": "",
                    "product_name": "",
                    "destination_key": "SE",
                    "destination": DESTINATION_PRESETS["SE"]["display"],
                    "ship_to": default_destination_address("SE"),
                    "shipment_date": "",
                    "ex_factory_date": "",
                    "items": [],
                    "cartons": [],
                    "draft_carton_strategy": "one_carton_per_color",
                    "warnings": [f"Parse failed: {exc}"],
                })
            finally:
                try:
                    temp_path.unlink(missing_ok=True)
                except Exception:
                    pass
        self.send_json({"orders": orders})


def main():
    url = "http://127.0.0.1:5057"
    print("ACNE PDF Packing + Labels app")
    print("Open:", url)
    httpd = ThreadingHTTPServer(("127.0.0.1", 5057), Handler)
    # Open exactly one browser tab from Python after the server is ready.
    try:
        threading.Timer(0.6, lambda: webbrowser.open_new_tab(url)).start()
    except Exception:
        pass
    print("Close the browser tab to stop the app automatically, or press Ctrl+C.")
    httpd.serve_forever()


if __name__ == "__main__":
    main()
